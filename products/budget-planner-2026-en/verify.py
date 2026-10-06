#!/usr/bin/env python3
"""Verifica delle formule di 2026 Budget Planner (versione inglese).

Uso:  python3 build.py && python3 verify.py

1. Controlla che build.py sia deterministico (stessi byte a ogni esecuzione).
2. Cerca funzioni non compatibili con Excel 2016 / Google Sheets / LibreOffice.
3. Ricalcola i file con LibreOffice headless (soffice --convert-to xlsx in una cartella temporanea)
   e cerca errori (#REF!, #NAME?, #VALUE!, #DIV/0!, #N/A, #NUM!).
4. Confronta i valori del file di esempio con quelli calcolati qui in Python dagli stessi dati:
   somme per categoria e mese, budget, riepiloghi, dashboard, classifica, obiettivi, fondi, debiti
   (simulazione mese per mese di minimi, snowball e avalanche, piu' NPER in forma chiusa).
5. Prova i controlli e i selettori su una copia con errori voluti (date, categorie, fondi, debiti,
   mese scelto a mano, metodo snowball, valuta).
6. Scrive images/src/data.js con i numeri dell'esempio, usati dalle immagini dell'inserzione.
Esce con codice 1 se qualcosa non torna.
"""
import datetime as dt
import hashlib
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

sys.dont_write_bytecode = True  # niente __pycache__ nella cartella del prodotto
import build as B  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ERR_RE = re.compile(r'^(#REF!|#NAME\?|#VALUE!|#DIV/0!|#N/A|#NUM!|#NULL!|Err:\d+)')
FORBIDDEN = re.compile(r'\b(LET|LAMBDA|XLOOKUP|XMATCH|FILTER|SORT|SORTBY|UNIQUE|SEQUENCE|TEXTJOIN|IFS|SWITCH|'
                       r'MAXIFS|MINIFS|CONCAT|INDIRECT|OFFSET|TEXT|CHOOSECOLS|TAKE|DROP|VSTACK|HSTACK)\(')
problems = []
checks = 0


def check(cond, msg):
    global checks
    checks += 1
    if not cond:
        problems.append(msg)


def close(a, b, tol=0.006):
    try:
        return abs(float(a) - float(b)) <= tol
    except (TypeError, ValueError):
        return False


def recalc(paths, outdir):
    profile = tempfile.mkdtemp(prefix='lo-profile-')
    try:
        subprocess.run(['soffice', f'-env:UserInstallation=file://{profile}', '--headless', '--convert-to', 'xlsx',
                        '--outdir', outdir, *paths], check=True, capture_output=True, timeout=300)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = [os.path.join(outdir, os.path.basename(p)) for p in paths]
    for p in out:
        if not os.path.exists(p):
            problems.append(f'LibreOffice non ha prodotto {p}')
    return out


def scan_errors(path, label):
    wb = load_workbook(path, data_only=True)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and ERR_RE.match(c.value):
                    n += 1
                    if n <= 20:
                        problems.append(f'{label}: error {c.value} in {ws.title}!{c.coordinate}')
    check(n == 0, f'{label}: {n} formula errors')
    return wb


def scan_formulas(path, label):
    wb = load_workbook(path)
    count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    count += 1
                    m = FORBIDDEN.search(c.value.upper())
                    if m:
                        problems.append(f'{label}: forbidden function {m.group(1)} in {ws.title}!{c.coordinate}')
    check(not wb.vba_archive if hasattr(wb, 'vba_archive') else True, f'{label}: contains macros')
    return count


# ---------------------------------------------------------------- valori attesi (Python puro)
def months_between(d, asof):
    return max(0, (d.year - asof.year) * 12 + d.month - asof.month)


def simulate(debts, order, extra, rollover):
    """Same rules as Debt Schedule. debts: list of (name, bal, apr, min) in the given order."""
    ds = [debts[i] for i in order]
    bal = [float(d[1]) for d in ds]
    rate = [float(d[2]) for d in ds]
    mins = [float(d[3]) for d in ds]
    pay = sum(mins) + (extra if rollover else 0)
    hist = []
    total_int = 0.0
    for _ in range(B.HORIZON):
        prev = bal[:]
        interest = sum(prev[k] * rate[k] for k in range(len(ds))) / 12
        need = [0.0 if prev[k] < 0.005 else max(0.0, prev[k] * (1 + rate[k] / 12) - mins[k]) for k in range(len(ds))]
        if rollover:
            pool = pay - (sum(prev) + interest - sum(need))
            new, cum = [], 0.0
            for k in range(len(ds)):
                new.append(max(0.0, need[k] - max(0.0, pool - cum)))
                cum += need[k]
            bal = new
        else:
            bal = need
        total_int += interest
        hist.append(bal[:])

    def months(series, start):
        if start < 0.005:
            return None
        n = sum(1 for v in series if v > 0.005)
        return 'Over 30 years' if n >= B.HORIZON else n + 1

    per = [months([h[k] for h in hist], float(ds[k][1])) for k in range(len(ds))]
    tot = months([sum(h) for h in hist], sum(float(d[1]) for d in ds))
    totals = [sum(float(d[1]) for d in ds)] + [sum(h) for h in hist]
    return {'names': [d[0] for d in ds], 'months': per, 'debt_free': tot, 'interest': total_int, 'pay': pay, 'totals': totals}


def expected_example():
    tx = B.example_transactions()
    cats = {c[0]: c[1] for c in B.CATEGORIES}
    act = {c: [0.0] * 12 for c in cats}
    from_funds = [0.0] * 12
    for d, desc, cat, amt, acc, fund, note in tx:
        act[cat][d.month - 1] += amt
        if fund and cats[cat] in (B.FIX, B.VAR):
            from_funds[d.month - 1] += amt
    plan = {}
    for c in cats:
        typical, ov = B.EXAMPLE_BUDGET.get(c, (0, {}))
        plan[c] = [float(ov.get(m, typical)) for m in range(1, 13)]
    plan_ff = [float(B.EXAMPLE_FROM_FUNDS.get(m, 0)) for m in range(1, 13)]

    def by_type(src, typ):
        return [sum(src[c][m] for c in cats if cats[c] == typ) for m in range(12)]

    rows = {}
    for label, src, ff in (('act', act, from_funds), ('plan', plan, plan_ff)):
        inc, fix, var, sav = (by_type(src, t) for t in B.TYPES)
        exp = [fix[m] + var[m] for m in range(12)]
        left = [inc[m] - exp[m] - sav[m] + ff[m] for m in range(12)]
        rows[label] = {6: inc, 7: fix, 8: var, 9: exp, 10: sav, 11: ff, 12: left}

    asof = B.EXAMPLE_SETTINGS['asof']
    goals = []
    for name, target, date, start, plan_m in B.EXAMPLE_GOALS:
        added = sum(t[3] for t in tx if t[5] == name and cats[t[2]] == B.SAV) - \
            sum(t[3] for t in tx if t[5] == name and cats[t[2]] in (B.FIX, B.VAR))
        saved = start + added
        needed = max(0, target - saved)
        ml = months_between(date, asof)
        per_month = 0 if needed <= 0 else needed / max(1, ml)
        status = 'Reached' if needed <= 0 else ('Target date reached' if ml == 0 else
                                                 ('On track' if plan_m >= per_month - 0.005 else 'Behind'))
        goals.append({'name': name, 'target': target, 'date': date.isoformat(), 'start': start, 'added': added, 'saved': saved,
                      'needed': needed, 'pct': min(1, saved / target), 'months': ml, 'per_month': per_month, 'plan': plan_m,
                      'status': status})
    funds = []
    for name, target, date, start in B.EXAMPLE_FUNDS:
        added = sum(t[3] for t in tx if t[5] == name and cats[t[2]] == B.SAV)
        spent = sum(t[3] for t in tx if t[5] == name and cats[t[2]] in (B.FIX, B.VAR))
        balance = start + added - spent
        needed = max(0, target - balance)
        ml = months_between(date, asof)
        funds.append({'name': name, 'target': target, 'date': date.isoformat(), 'start': start, 'added': added, 'spent': spent,
                      'balance': balance, 'needed': needed, 'pct': min(1, max(0, balance / target)), 'months': ml,
                      'per_month': 0 if needed <= 0 else needed / max(1, ml)})

    debts = B.EXAMPLE_DEBTS
    n = len(debts)
    aval = sorted(range(n), key=lambda i: (-debts[i][2], debts[i][1], i))
    snow = sorted(range(n), key=lambda i: (debts[i][1], -debts[i][2], i))
    extra = B.EXAMPLE_DEBT_OPTIONS['extra']
    sims = {'min': simulate(debts, list(range(n)), extra, False), 'snow': simulate(debts, snow, extra, True),
            'aval': simulate(debts, aval, extra, True)}
    return {'tx': tx, 'act': act, 'plan': plan, 'rows': rows, 'goals': goals, 'funds': funds,
            'aval': aval, 'snow': snow, 'sims': sims}


def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return dt.datetime(d.year + y, m + 1, 1)


# ---------------------------------------------------------------- controlli sul file di esempio
def verify_example(wb, E):
    asof = B.EXAMPLE_SETTINGS['asof']
    st = wb['Settings']
    check(st['C12'].value == 12, f'Settings!C12 month used = {st["C12"].value}, expected 12')
    check(st['C14'].value == ' ($)', f'Settings!C14 = {st["C14"].value!r}')
    mb, ms, ad = wb['Monthly Budget'], wb['Monthly Summary'], wb['Annual Dashboard']
    cats = [c[0] for c in B.CATEGORIES]
    for i, c in enumerate(cats):
        r = B.G0 + i
        check(mb[f'A{r}'].value == c and ms[f'A{r}'].value == c, f'category row {r}: {ms[f"A{r}"].value} != {c}')
        for m in range(12):
            pv = mb.cell(r, 4 + m).value
            av = ms.cell(r, B.MS_MONTH0 + m).value
            check(close(pv, E['plan'][c][m]), f'Monthly Budget {c} month {m + 1}: {pv} != {E["plan"][c][m]}')
            check(close(av, E['act'][c][m]), f'Monthly Summary {c} month {m + 1}: {av} != {E["act"][c][m]}')
        check(close(ms[f'X{r}'].value, sum(E['act'][c])), f'actual year {c}')
        check(close(ms[f'Y{r}'].value, sum(E['plan'][c])), f'planned year {c}')
        check(close(ms[f'C{r}'].value, E['plan'][c][11]) and close(ms[f'D{r}'].value, E['act'][c][11]), f'December block {c}')
        check(close(ms[f'I{r}'].value, sum(E['act'][c])), f'YTD actual {c}')
    for r in range(6, 13):
        for m in range(12):
            check(close(ms.cell(r, B.MS_MONTH0 + m).value, E['rows']['act'][r][m]),
                  f'Monthly Summary row {r} month {m + 1}: {ms.cell(r, B.MS_MONTH0 + m).value} != {E["rows"]["act"][r][m]}')
            check(close(mb.cell(r, 4 + m).value or 0, E['rows']['plan'][r][m]),   # riga 11: celle vuote = 0
                  f'Monthly Budget row {r} month {m + 1}: {mb.cell(r, 4 + m).value} != {E["rows"]["plan"][r][m]}')
        check(close(ms[f'X{r}'].value, sum(E['rows']['act'][r])), f'Monthly Summary row {r} year')
        check(close(ms[f'Y{r}'].value, sum(E['rows']['plan'][r])), f'Monthly Summary row {r} planned year')
        check(close(mb[f'P{r}'].value, sum(E['rows']['plan'][r])), f'Monthly Budget row {r} year')
    A = {k: sum(v) for k, v in E['rows']['act'].items()}
    P = {k: sum(v) for k, v in E['rows']['plan'].items()}
    check(close(ad['B5'].value, A[6]), 'dashboard income')
    check(close(ad['D5'].value, A[9]), 'dashboard expenses')
    check(close(ad['F5'].value, A[10]), 'dashboard savings & debt')
    check(close(ad['H5'].value, (A[6] - A[9]) / A[6], 1e-9), 'dashboard savings rate')
    check(close(ad['J5'].value, A[12]), 'dashboard left over')
    check(close(ad['L5'].value, sum(d[1] for d in B.EXAMPLE_DEBTS)), 'dashboard total debt')
    check(ad['B6'].value == f'{round(A[6] / P[6] * 100)}% of the yearly plan', f'dashboard income note {ad["B6"].value}')
    for m in range(12):
        r = 14 + m
        check(close(ad[f'C{r}'].value, E['rows']['act'][6][m]) and close(ad[f'D{r}'].value, E['rows']['act'][9][m])
              and close(ad[f'F{r}'].value, E['rows']['act'][12][m]), f'dashboard month row {r}')
    # top 10
    spend = sorted(((sum(E['act'][c]), i, c) for i, c in enumerate(cats)
                    if dict((x[0], x[1]) for x in B.CATEGORIES)[c] in (B.FIX, B.VAR)), reverse=True)[:10]
    for k, (tot, _, c) in enumerate(spend):
        r = 41 + k
        check(ad[f'C{r}'].value == c and close(ad[f'F{r}'].value, tot), f'top 10 #{k + 1}: {ad[f"C{r}"].value} != {c}')
        check(close(ad[f'G{r}'].value, tot / A[9], 1e-9), f'top 10 share #{k + 1}')
    for r in (30, 31, 32, 33):
        check(ad[f'M{r}'].value == 0, f'dashboard check row {r} = {ad[f"M{r}"].value}')
    check(ad['M29'].value == len(E['tx']), f'transactions logged {ad["M29"].value} != {len(E["tx"])}')
    a3 = wb['Transactions']['A3'].value
    check(a3 == f'Rows used: {len(E["tx"])} of 2,000 · Total income {B.YEAR}: {A[6]:,.2f} · Total expenses {B.YEAR}: {A[9]:,.2f}',
          f'Transactions A3 {a3!r}')

    # goals
    sg = wb['Savings Goals']
    for i, g in enumerate(E['goals']):
        r = B.GOAL0 + i
        check(close(sg[f'F{r}'].value, g['added']) and close(sg[f'G{r}'].value, g['saved']), f'goal {g["name"]} saved')
        check(close(sg[f'H{r}'].value, g['needed']) and sg[f'K{r}'].value == g['months'], f'goal {g["name"]} months/needed')
        check(close(sg[f'L{r}'].value, g['per_month']), f'goal {g["name"]} per month {sg[f"L{r}"].value} != {g["per_month"]}')
        check(sg[f'O{r}'].value == g['status'], f'goal {g["name"]} status {sg[f"O{r}"].value} != {g["status"]}')
        if g['needed'] > 0 and g['plan'] > 0:
            exp_date = add_months(asof, math.ceil(g['needed'] / g['plan']))
            check(sg[f'N{r}'].value == exp_date, f'goal {g["name"]} reached by {sg[f"N{r}"].value} != {exp_date}')
    t = B.GOAL1 + 1
    check(close(sg[f'G{t}'].value, sum(g['saved'] for g in E['goals'])), 'goals total saved')
    # funds
    sf = wb['Sinking Funds']
    for i, f in enumerate(E['funds']):
        r = B.FUND0 + i
        for col, key in (('F', 'added'), ('G', 'spent'), ('H', 'balance'), ('I', 'needed'), ('M', 'per_month')):
            check(close(sf[f'{col}{r}'].value, f[key]), f'fund {f["name"]} {key}: {sf[f"{col}{r}"].value} != {f[key]}')
        check(sf[f'L{r}'].value == f['months'], f'fund {f["name"]} months')
    t = B.FUND1 + 1
    check(close(sf[f'M{t}'].value, sum(f['per_month'] for f in E['funds'])), 'funds total per month')
    check(close(sf[f'M{t + 2}'].value, B.EXAMPLE_BUDGET['Sinking funds'][0]), 'funds: planned in budget')

    # debts
    dp = wb['Debt Payoff']
    debts = B.EXAMPLE_DEBTS
    for k, i in enumerate(E['aval']):
        check(dp[f'I{B.DEBT0 + i}'].value == k + 1, f'avalanche order of {debts[i][0]}')
        check(dp[f'B{33 + k}'].value == debts[i][0], f'avalanche list #{k + 1}')
    for k, i in enumerate(E['snow']):
        check(dp[f'J{B.DEBT0 + i}'].value == k + 1, f'snowball order of {debts[i][0]}')
        check(dp[f'B{46 + k}'].value == debts[i][0], f'snowball list #{k + 1}')
    sims = E['sims']
    for i, d in enumerate(debts):
        r = B.DEBT0 + i
        mo = sims['min']['months'][i]
        check(dp[f'G{r}'].value == mo, f'{d[0]} months (minimums) {dp[f"G{r}"].value} != {mo}')
        # forma chiusa (NPER) per il solo minimo
        rr = d[2] / 12
        n = d[1] / d[3] if rr == 0 else -math.log(1 - rr * d[1] / d[3]) / math.log(1 + rr)
        check(mo == math.ceil(n - 1e-9), f'{d[0]}: simulated {mo} months vs NPER {n:.3f}')
        k = E['aval'].index(i)
        check(dp[f'M{r}'].value == sims['aval']['months'][k], f'{d[0]} months (avalanche) {dp[f"M{r}"].value} != {sims["aval"]["months"][k]}')
        check(dp[f'K{r}'].value == add_months(asof, sims['aval']['months'][k]), f'{d[0]} paid off date')
    for col, key in (('C', 'min'), ('D', 'snow'), ('E', 'aval')):
        s = sims[key]
        check(close(dp[f'{col}23'].value, s['pay']), f'compare {key}: paid each month')
        check(dp[f'{col}24'].value == s['debt_free'], f'compare {key}: months {dp[f"{col}24"].value} != {s["debt_free"]}')
        check(close(dp[f'{col}26'].value, s['interest'], 0.01), f'compare {key}: interest {dp[f"{col}26"].value} != {s["interest"]}')
    check(sims['aval']['interest'] < sims['snow']['interest'] < sims['min']['interest'], 'avalanche < snowball < minimums interest')
    check(close(dp['E27'].value, sims['min']['interest'] - sims['aval']['interest'], 0.01), 'interest saved (avalanche)')
    step = max(1, math.ceil(max(sims['snow']['debt_free'], sims['aval']['debt_free']) / 24))
    check(dp['I59'].value == step, f'chart step {dp["I59"].value} != {step}')
    for i in range(25):
        mth = min(B.HORIZON, i * step)
        check(close(dp[f'E{60 + i}'].value, sims['aval']['totals'][mth], 0.01), f'chart avalanche point {i}')
        check(close(dp[f'C{60 + i}'].value, sims['min']['totals'][mth], 0.01), f'chart minimums point {i}')
    ad_debtfree = wb['Annual Dashboard']['J9'].value
    check(ad_debtfree == add_months(asof, sims['aval']['debt_free']), f'dashboard debt-free {ad_debtfree}')
    tr = wb['Transactions']
    bad = [tr[f'I{r}'].value for r in range(B.TX0, B.TX0 + len(E['tx'])) if tr[f'I{r}'].value]
    check(not bad, f'example transactions with check messages: {bad[:5]}')


# ---------------------------------------------------------------- copia con errori voluti
def broken_copy(path):
    wb = B.build(True)
    tr = wb['Transactions']
    n = len(B.example_transactions())
    r0 = B.TX0 + n
    rows = [
        (None, 'No date', 'Groceries', 10, 'Date missing'),
        ('15/03/2026', 'Text date', 'Groceries', 10, 'Date not recognized'),
        (dt.date(2025, 12, 31), 'Last year', 'Groceries', 10, 'Not in 2026'),
        (dt.date(2026, 3, 1), 'No category', None, 10, 'Category missing'),
        (dt.date(2026, 3, 1), 'Wrong category', 'Food', 10, 'Category not in list'),
        (dt.date(2026, 3, 1), 'No amount', 'Groceries', None, 'Amount missing'),
        (dt.date(2026, 3, 1), 'Text amount', 'Groceries', '12.50', 'Amount is not a number'),
        (dt.date(2026, 3, 1), 'Untyped category', 'Untyped', 10, 'Category has no type'),
        (dt.date(2026, 3, 1), 'Wrong fund', 'Groceries', 10, 'Goal / fund not in list'),
        (dt.date(2026, 3, 1), 'Fund on income', 'Side income', 10, 'Fund tag ignored on income'),
        (dt.date(2026, 3, 5), 'Refund', 'Groceries', -25.5, ''),
    ]
    for i, (d, desc, cat, amt, _) in enumerate(rows):
        r = r0 + i
        tr[f'A{r}'] = d
        tr[f'B{r}'] = desc
        tr[f'C{r}'] = cat
        tr[f'D{r}'] = amt
        tr[f'F{r}'] = {'Wrong fund': 'Boat', 'Fund on income': 'Holiday gifts'}.get(desc)
    ms = wb['Monthly Summary']
    ms['B3'] = 'March'
    dp = wb['Debt Payoff']
    dp['C4'] = 'Snowball'
    dp['B14'] = 'Low payment card'
    dp['C14'] = 1000
    dp['D14'] = 0.24
    dp['E14'] = 15            # sotto gli interessi del mese (20)
    dp['B15'] = None
    dp['C15'] = 500           # importo senza nome
    cat = wb['Categories']
    cat[f'A{B.CAT0 + 40}'] = 'Groceries'          # nome doppio
    cat[f'B{B.CAT0 + 40}'] = B.VAR
    cat[f'A{B.CAT0 + 41}'] = 'Untyped'            # categoria senza tipo
    wb['Monthly Budget'][f'C{B.G0 + 50}'] = 99    # importo su una riga senza categoria: non deve contare
    wb['Sinking Funds'][f'B{B.FUND0 + 6}'] = 'Pet fund'   # fondo dopo due righe vuote: il menu non deve avere buchi
    wb['Settings']['C5'] = '£'
    B.save_deterministic(wb, path)
    return [r[4] for r in rows], r0


def verify_broken(wb, expected_msgs, r0, E):
    tr = wb['Transactions']
    for i, msg in enumerate(expected_msgs):
        got = tr[f'I{r0 + i}'].value or ''
        check(got == msg, f'broken copy row {r0 + i}: check {got!r} != {msg!r}')
    ad = wb['Annual Dashboard']
    check(ad['M30'].value == sum(1 for m in expected_msgs if m), f'dashboard rows to check {ad["M30"].value}')
    check(ad['M31'].value == 3, f'dashboard categories to check {ad["M31"].value}')
    check(ad['M33'].value == 2, f'dashboard debts to check {ad["M33"].value}')
    dp = wb['Debt Payoff']
    check(dp['L14'].value == 'Minimum does not cover interest', f'debt check {dp["L14"].value}')
    check(dp['L15'].value == 'Enter a name', f'debt check {dp["L15"].value}')
    check(dp['G14'].value == 'Over 30 years', f'low payment card minimums only: {dp["G14"].value}')
    check(dp['C24'].value == 'Over 30 years', f'minimums only debt-free: {dp["C24"].value}')
    check(isinstance(dp['D24'].value, int) or isinstance(dp['D24'].value, float), f'snowball debt-free: {dp["D24"].value}')
    check(dp['K18'].value == dp['D25'].value, 'snowball selected: K18 = D25')
    ms = wb['Monthly Summary']
    st = wb['Settings']
    check(st['C12'].value == 3, f'month picked by hand: {st["C12"].value}')
    check(ms['C4'].value == 'March 2026 (£)', f'heading with currency: {ms["C4"].value!r}')
    g = B.G0 + [c[0] for c in B.CATEGORIES].index('Groceries')
    # contano solo le righe con data valida di marzo 2026: "Wrong fund" (+10) e "Refund" (-25.50)
    act_march = E['act']['Groceries'][2] + 10 - 25.5
    check(close(ms[f'D{g}'].value, act_march), f'March groceries actual {ms[f"D{g}"].value} != {act_march}')
    check(close(ms[f'I{g}'].value, sum(E['act']['Groceries'][:3]) + 10 - 25.5), 'YTD March groceries')
    check(close(ms[f'H{g}'].value, sum(E['plan']['Groceries'][:3])), 'YTD March planned groceries')
    check(wb['Transactions']['D5'].value == 'Amount (£)', 'Transactions heading currency')
    check(ms['H4'].value == 'Year to date (Jan–Mar) (£)', f'YTD heading {ms["H4"].value!r}')
    # riga senza nome in Monthly Budget: nessun importo pianificato nel riepilogo
    r = B.G0 + 50
    check(ms[f'Y{r}'].value == 0 and ms[f'H{r}'].value == 0 and ms[f'C{r}'].value == 0, f'unnamed budget row counted: {ms[f"Y{r}"].value}')
    check(close(ms['Y8'].value, sum(E['rows']['plan'][8])), 'unnamed budget row changed planned variable expenses')
    # Transactions A3: totali solo dell'anno del piano (la riga del 2025 e le date di testo non contano)
    a3 = wb['Transactions']['A3'].value
    inc = sum(E['rows']['act'][6]) + 10            # "Fund on income"
    exp = sum(E['rows']['act'][9]) + 10 - 25.5     # "Wrong fund" e "Refund"
    check(f'Total income {B.YEAR}: {inc:,.2f}' in a3 and f'Total expenses {B.YEAR}: {exp:,.2f}' in a3,
          f'Transactions A3 {a3!r}: expected income {inc:,.2f}, expenses {exp:,.2f}')
    # menu obiettivi/fondi senza righe vuote in mezzo
    lst = [wb['Lists'][f'G{2 + i}'].value or '' for i in range(B.N_GOAL + B.N_FUND)]
    want = [g[0] for g in B.EXAMPLE_GOALS] + [f[0] for f in B.EXAMPLE_FUNDS] + ['Pet fund']
    check(lst == want + [''] * (len(lst) - len(want)), f'goal/fund drop-down list {lst}')


# ---------------------------------------------------------------- dati per le immagini
def r2(v):
    return round(float(v), 6) if isinstance(v, (int, float)) else v


def iso(v):
    return v.strftime('%Y-%m-%d') if isinstance(v, (dt.date, dt.datetime)) else v


def write_data_js(wb, E):
    ms, ad, mb = wb['Monthly Summary'], wb['Annual Dashboard'], wb['Monthly Budget']
    tr, sg, sf, dp = wb['Transactions'], wb['Savings Goals'], wb['Sinking Funds'], wb['Debt Payoff']
    cats = []
    for i, (name, typ, _) in enumerate(B.CATEGORIES):
        r = B.G0 + i
        cats.append({'name': name, 'type': typ, 'typical': r2(mb[f'C{r}'].value or 0),
                     'plan': [r2(mb.cell(r, 4 + m).value) for m in range(12)],
                     'act': [r2(ms.cell(r, B.MS_MONTH0 + m).value) for m in range(12)],
                     'planYear': r2(ms[f'Y{r}'].value), 'actYear': r2(ms[f'X{r}'].value)})
    rows = {}
    for r, key in ((6, 'income'), (7, 'fixed'), (8, 'variable'), (9, 'expenses'), (10, 'savings'), (11, 'fromFunds'), (12, 'left')):
        rows[key] = {'plan': [r2(mb.cell(r, 4 + m).value) for m in range(12)],
                     'act': [r2(ms.cell(r, B.MS_MONTH0 + m).value) for m in range(12)],
                     'planYear': r2(ms[f'Y{r}'].value), 'actYear': r2(ms[f'X{r}'].value)}
    txs = []
    for r in range(B.TX0, B.TX0 + len(E['tx'])):
        txs.append([iso(tr[f'A{r}'].value), tr[f'B{r}'].value, tr[f'C{r}'].value, r2(tr[f'D{r}'].value), tr[f'E{r}'].value,
                    tr[f'F{r}'].value or '', tr[f'H{r}'].value])
    goals = [{k: (iso(v) if k == 'reachedBy' else r2(v)) for k, v in zip(
        ('name', 'target', 'date', 'start', 'added', 'saved', 'needed', 'pct', 'bar', 'months', 'perMonth', 'plan', 'reachedBy', 'status'),
        [iso(sg.cell(r, c).value) for c in range(2, 16)])} for r in range(B.GOAL0, B.GOAL0 + len(B.EXAMPLE_GOALS))]
    funds = [{k: r2(v) for k, v in zip(
        ('name', 'target', 'date', 'start', 'added', 'spent', 'balance', 'needed', 'pct', 'bar', 'months', 'perMonth', 'status'),
        [iso(sf.cell(r, c).value) for c in range(2, 15)])} for r in range(B.FUND0, B.FUND0 + len(B.EXAMPLE_FUNDS))]
    debts = [{k: (iso(v) if k in ('paidMin', 'paidPlan') else r2(v)) for k, v in zip(
        ('name', 'balance', 'apr', 'min', 'interest', 'monthsMin', 'paidMin', 'aval', 'snow', 'paidPlan', 'check', 'monthsPlan'),
        [dp.cell(r, c).value for c in range(2, 14)])} for r in range(B.DEBT0, B.DEBT0 + len(B.EXAMPLE_DEBTS))]
    compare = {col: {'pay': r2(dp[f'{col}23'].value), 'months': dp[f'{col}24'].value, 'date': iso(dp[f'{col}25'].value),
                     'interest': r2(dp[f'{col}26'].value), 'saved': r2(dp[f'{col}27'].value)} for col in 'CDE'}
    chart = [{'label': dp[f'B{r}'].value, 'min': r2(dp[f'C{r}'].value), 'snow': r2(dp[f'D{r}'].value), 'aval': r2(dp[f'E{r}'].value)}
             for r in range(60, 85)]
    top = [{'name': ad[f'C{r}'].value, 'value': r2(ad[f'F{r}'].value), 'share': r2(ad[f'G{r}'].value)} for r in range(41, 51)]
    data = {
        'year': B.YEAR, 'name': B.EXAMPLE_SETTINGS['name'], 'currency': B.EXAMPLE_SETTINGS['currency'],
        'file': f'{B.BASE}-example.xlsx', 'months': B.SHORT, 'monthNames': B.MONTHS,
        'kpi': {'income': r2(ad['B5'].value), 'expenses': r2(ad['D5'].value), 'savings': r2(ad['F5'].value),
                'rate': r2(ad['H5'].value), 'left': r2(ad['J5'].value), 'debt': r2(ad['L5'].value),
                'notes': [ad[f'{c}6'].value for c in 'BDFHJL'],
                'goalsSaved': r2(ad['B9'].value), 'goalsNote': ad['B10'].value, 'fundsBalance': r2(ad['F9'].value),
                'fundsNote': ad['F10'].value, 'debtFree': iso(ad['J9'].value), 'debtNote': ad['J10'].value},
        'categories': cats, 'rows': rows, 'transactions': txs, 'top': top, 'goals': goals, 'funds': funds,
        'goalsTotal': {'target': r2(sg[f'C{B.GOAL1 + 1}'].value), 'saved': r2(sg[f'G{B.GOAL1 + 1}'].value),
                       'plan': r2(sg[f'M{B.GOAL1 + 1}'].value), 'needed': r2(sg[f'L{B.GOAL1 + 1}'].value)},
        'fundsTotal': {'target': r2(sf[f'C{B.FUND1 + 1}'].value), 'balance': r2(sf[f'H{B.FUND1 + 1}'].value),
                       'perMonth': r2(sf[f'M{B.FUND1 + 1}'].value)},
        'debts': debts, 'method': dp['C4'].value, 'extra': r2(dp['C5'].value), 'compare': compare, 'chart': chart,
        'debtTotal': {'balance': r2(dp[f'C{B.DEBT1 + 1}'].value), 'min': r2(dp[f'E{B.DEBT1 + 1}'].value)},
        'avalOrder': [dp[f'B{r}'].value for r in range(33, 33 + len(B.EXAMPLE_DEBTS))],
        'snowOrder': [dp[f'B{r}'].value for r in range(46, 46 + len(B.EXAMPLE_DEBTS))],
    }
    out = os.path.join(HERE, 'images', 'src', 'data.js')
    with open(out, 'w', encoding='utf-8') as fh:
        fh.write('// Generato da verify.py dal file di esempio ricalcolato con LibreOffice. Non modificare a mano.\n')
        fh.write('window.DATA = ' + json.dumps(data, ensure_ascii=False, indent=1) + ';\n')
    print(f'written {os.path.relpath(out, HERE)}')


def main():
    files = [os.path.join(B.OUT, f'{B.BASE}.xlsx'), os.path.join(B.OUT, f'{B.BASE}-example.xlsx')]
    tmp = tempfile.mkdtemp(prefix='bp-verify-')
    try:
        # 1. determinismo
        for example, path in zip((False, True), files):
            buf = os.path.join(tmp, 'again-' + os.path.basename(path))
            B.save_deterministic(B.build(example), buf)
            same = hashlib.sha256(open(buf, 'rb').read()).hexdigest() == hashlib.sha256(open(path, 'rb').read()).hexdigest()
            check(same, f'{os.path.basename(path)} differs from a fresh build (run build.py, or build is not deterministic)')
        # 2. funzioni
        total = sum(scan_formulas(p, os.path.basename(p)) for p in files)
        # 3-5. ricalcolo
        broken = os.path.join(tmp, 'broken.xlsx')
        msgs, r0 = broken_copy(broken)
        outdir = os.path.join(tmp, 'out')
        os.makedirs(outdir)
        blank_r, example_r, broken_r = recalc(files + [broken], outdir)
        scan_errors(blank_r, 'blank')
        wb_ex = scan_errors(example_r, 'example')
        wb_br = scan_errors(broken_r, 'broken copy')
        E = expected_example()
        verify_example(wb_ex, E)
        verify_broken(wb_br, msgs, r0, E)
        wb_blank = load_workbook(blank_r, data_only=True)
        check(wb_blank['Annual Dashboard']['B5'].value == 0, 'blank: income should be 0')
        check(wb_blank['Debt Payoff']['C25'].value == '–', 'blank: debt-free date should be a dash')
        check(wb_blank['Debt Payoff']['B28'].value == 'Your plan: enter your debts above.', 'blank: debt plan sentence')
        check(wb_blank['Start Here']['B1'].value == f'{B.YEAR} Budget Planner', 'blank: Start Here title')
        write_data_js(wb_ex, E)
        print(f'formulas: {total} in 2 files; checks: {checks}')
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    if problems:
        print(f'{len(problems)} PROBLEMS:')
        for p in problems[:80]:
            print(' -', p)
        sys.exit(1)
    print('OK: no formula errors, values match the hand calculation')


if __name__ == '__main__':
    main()
