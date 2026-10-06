#!/usr/bin/env python3
"""2026 Budget Planner (NettoChiaro, versione inglese): rigenera i file consegnati all'acquirente.

Uso:  python3 build.py

Scrive in files/:
  budget-planner-2026.xlsx           modello vuoto da compilare
  budget-planner-2026-example.xlsx   stesso modello con un anno di dati di una famiglia di fantasia

Deterministico: stessi dati, stessi byte (le date dei metadati e dello zip sono fissate).
Solo funzioni comuni a Excel 2016+, Google Sheets e LibreOffice (niente LET, XLOOKUP, FILTER,
OFFSET, INDIRECT, macro). Verifica delle formule: python3 verify.py (ricalcola con LibreOffice).
"""
import calendar
import datetime as dt
import io
import os
import random
import re
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.drawing.line import LineProperties
from openpyxl.worksheet.pagebreak import Break
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

YEAR = 2026
VERSION = '2026.1'
SLUG = 'budget-planner-2026-en'
BASE = 'budget-planner-2026'
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'files')
FIXED = dt.datetime(2026, 1, 1, 0, 0, 0)

# ---------------------------------------------------------------- dimensioni (riferimenti fissi)
CAT0, N_CAT = 6, 60                  # Categories!A6:A65
CAT1 = CAT0 + N_CAT - 1
TX0, N_TX = 6, 2000                  # Transactions!A6:I2005
TX1 = TX0 + N_TX - 1
G0 = 14                              # prima riga categorie in Monthly Budget / Monthly Summary
G1 = G0 + N_CAT - 1                  # 73 (riga categoria = riga Categories + 8)
ACC0, N_ACC = 18, 8                  # Settings!B18:B25
ACC1 = ACC0 + N_ACC - 1
GOAL0, N_GOAL = 6, 12                # Savings Goals!B6:B17
GOAL1 = GOAL0 + N_GOAL - 1
FUND0, N_FUND = 6, 12                # Sinking Funds!B6:B17
FUND1 = FUND0 + N_FUND - 1
DEBT0, N_DEBT = 8, 10                # Debt Payoff!B8:B17
DEBT1 = DEBT0 + N_DEBT - 1
HORIZON = 360                        # mesi simulati in Debt Schedule (30 anni)
S0 = 10                              # riga del mese 0 in Debt Schedule
S1 = S0 + HORIZON                    # 370

MB, MS, AD = "'Monthly Budget'", "'Monthly Summary'", "'Annual Dashboard'"
SG, SF, DP, DS = "'Savings Goals'", "'Sinking Funds'", "'Debt Payoff'", "'Debt Schedule'"

YEAR_REF = 'Settings!$C$4'
CUR_SYM = 'Settings!$C$5'
NAME_REF = 'Settings!$C$6'
ASOF = 'Settings!$C$11'
MON = 'Settings!$C$12'
MON_NAME = 'Settings!$C$13'
CUR = 'Settings!$C$14'               # " ($)" oppure "" se il simbolo e' vuoto
MON_SHORT = 'Settings!$C$15'

L_MONTHS = 'Lists!$A$2:$A$13'
L_CHOICES = 'Lists!$C$2:$C$14'
L_TYPES = 'Lists!$D$2:$D$5'
L_METHODS = 'Lists!$E$2:$E$3'
L_CURR = 'Lists!$F$2:$F$13'
L_FUNDS = 'Lists!$G$2:$G$25'
L_MNUM = 'Lists!$I$2:$T$2'           # numeri dei mesi 1..12 in orizzontale (criteri SUMIFS)

T_DATE = f'Transactions!$A${TX0}:$A${TX1}'
T_CAT = f'Transactions!$C${TX0}:$C${TX1}'
T_AMT = f'Transactions!$D${TX0}:$D${TX1}'
T_FUND = f'Transactions!$F${TX0}:$F${TX1}'
T_TYPE = f'Transactions!$H${TX0}:$H${TX1}'
T_CHK = f'Transactions!$I${TX0}:$I${TX1}'
C_NAMES = f'Categories!$A${CAT0}:$A${CAT1}'
C_TYPES = f'Categories!$B${CAT0}:$B${CAT1}'

MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September',
          'October', 'November', 'December']
SHORT = [m[:3] for m in MONTHS]
INC, FIX, VAR, SAV = 'Income', 'Fixed bills', 'Variable expenses', 'Savings & debt'
TYPES = [INC, FIX, VAR, SAV]
CURRENCIES = ['$', '€', '£', 'CA$', 'A$', 'NZ$', 'CHF', 'kr', 'zł', 'R', '₹', '¥']

CATEGORIES = [
    ('Paycheck 1', INC, 'Take-home pay after tax. Rename it, for example "Alex pay".'),
    ('Paycheck 2', INC, 'Second earner, if any.'),
    ('Side income', INC, 'Freelance work, selling things, tips.'),
    ('Other income', INC, 'Tax refunds, gifts, interest, benefits.'),
    ('Rent / Mortgage', FIX, ''),
    ('Utilities', FIX, 'Electricity, gas, water, trash.'),
    ('Internet & Phone', FIX, ''),
    ('Insurance', FIX, 'Car, home or renters, life, health.'),
    ('Subscriptions', FIX, 'Streaming, apps, gym, memberships.'),
    ('Car payment', FIX, 'Car loan or lease.'),
    ('Groceries', VAR, 'Food and household supplies.'),
    ('Dining out', VAR, 'Restaurants, takeout, coffee.'),
    ('Transport / Gas', VAR, 'Fuel, public transport, parking, tolls, car repairs.'),
    ('Health', VAR, 'Pharmacy, doctor and dentist visits, copays.'),
    ('Kids', VAR, 'Childcare, school, activities, clothes.'),
    ('Personal care', VAR, 'Haircuts, toiletries, cosmetics.'),
    ('Entertainment', VAR, 'Movies, events, hobbies, games.'),
    ('Gifts', VAR, 'Birthdays, holidays, donations.'),
    ('Travel', VAR, 'Trips and vacations.'),
    ('Clothing', VAR, ''),
    ('Home & repairs', VAR, 'Furniture, tools, repairs, appliances.'),
    ('Pets', VAR, 'Food, vet, grooming.'),
    ('Misc', VAR, 'Anything that does not fit elsewhere.'),
    ('Emergency fund', SAV, 'Money moved to your emergency savings.'),
    ('Savings goals', SAV, 'Money moved to the goals in Savings Goals.'),
    ('Sinking funds', SAV, 'Money set aside for the funds in Sinking Funds.'),
    ('Retirement & investing', SAV, 'Pension or retirement account, investments.'),
    ('Debt payments', SAV, 'Credit cards, student and personal loans (see Debt Payoff).'),
]
ACCOUNTS = ['Checking', 'Savings', 'Credit card', 'Debit card', 'Cash']

# ---------------------------------------------------------------- stile
ARIAL = 'Arial'
TEAL, TEAL_D, TEAL_S = '0F766E', '0B5C56', 'E0F2EF'
TEXT, MUTED, LINE, SURF2 = '17201C', '5B6862', 'D9E0DC', 'EEF2EF'
INPUT, INPUT_LINE = 'FFF6D5', 'E3D49B'
WARN, WARN_T = 'FFE4C7', '9A3412'
GOOD, GOOD_S, BAD, BAD_S = '15803D', 'E7F6EC', 'B91C1C', 'FDECEC'
CH_INC, CH_EXP, CH_GREY = '0D9488', 'C2410C', '9AA5A0'

NUM = '#,##0.00'                         # importi da compilare
NUMZ = '#,##0.00;-#,##0.00;"–"'          # importi calcolati (zero = trattino)
PCT = '0.0%'
PCT0 = '0%'
PCT_IN = '0.00%'
DATE_F = 'yyyy-mm-dd'
MONTH_F = 'yyyy-mm'


def font(size=10, bold=False, color=TEXT, italic=False, underline=None):
    return Font(name=ARIAL, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill('solid', fgColor=color)


THIN = Side(style='thin', color=LINE)
THIN_IN = Side(style='thin', color=INPUT_LINE)
B_ROW = Border(bottom=THIN)
B_IN = Border(left=THIN_IN, right=THIN_IN, top=THIN_IN, bottom=THIN_IN)
B_TOT = Border(top=Side(style='thin', color=TEAL), bottom=Side(style='thin', color=TEAL))
UNLOCK = Protection(locked=False)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_C = Alignment(horizontal='left', vertical='center')
LEFT_IN = Alignment(horizontal='left', vertical='center', indent=1)
RIGHT_C = Alignment(horizontal='right', vertical='center')
RIGHT = Alignment(horizontal='right')


def chart_title(text):
    cp = CharacterProperties(sz=1100, b=True, solidFill=TEXT)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def light_grid(axis):
    axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill='E3E8E5', w=6350)))


def put(ws, ref, value=None, f=None, bg=None, fmt=None, al=None, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = f or font()
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if border:
        c.border = border
    return c


def inp(ws, ref, value=None, fmt=None, al=None):
    """Input cell: light yellow, unlocked."""
    c = put(ws, ref, value, bg=INPUT, fmt=fmt, al=al or LEFT_IN, border=B_IN)
    c.protection = UNLOCK
    return c


def title(ws, text, sub=None, col='A'):
    put(ws, f'{col}1', text, f=font(18, True, TEAL))
    ws.row_dimensions[1].height = 30
    if sub:
        put(ws, f'{col}2', sub, f=font(10, color=MUTED))


def header(ws, row, cols_texts, height=30):
    for col, text in cols_texts:
        put(ws, f'{col}{row}', text, f=font(10, True, 'FFFFFF'), bg=TEAL, al=CENTER)
    ws.row_dimensions[row].height = height


def section(ws, ref, text, span_to=None):
    c = put(ws, ref, text, f=font(11, True, TEAL_D), bg=TEAL_S)
    if span_to:
        for col in range(c.column + 1, span_to + 1):
            ws.cell(c.row, col).fill = fill(TEAL_S)
    ws.row_dimensions[c.row].height = 20


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def print_setup(ws, title_rows=None, area=None, portrait=False):
    ws.page_setup.orientation = 'portrait' if portrait else 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows
    if area:
        ws.print_area = area


def protect(ws, sort=False):
    """Protection without a password; formatting rows, columns and cells stays allowed."""
    p = ws.protection
    p.sheet = True
    p.formatColumns = False   # False = allowed
    p.formatRows = False
    p.formatCells = False
    if sort:
        p.sort = False
        p.autoFilter = False


def dv_list(ws, src, ranges, err_title, msg, strict=True):
    dv = DataValidation(type='list', formula1=src, allow_blank=True, showErrorMessage=strict,
                        errorTitle=err_title, error=msg)
    ws.add_data_validation(dv)
    dv.add(ranges)


def dv_date(ws, ranges):
    dv = DataValidation(type='date', operator='between', formula1='36526', formula2='73050',
                        allow_blank=True, showErrorMessage=True, errorTitle='Date not valid',
                        error='Enter a date, for example 2026-03-15 (year-month-day).')
    ws.add_data_validation(dv)
    dv.add(ranges)


def dv_num(ws, ranges, lo='-100000000', hi='100000000', msg='Enter a number, for example 125.50.'):
    dv = DataValidation(type='decimal', operator='between', formula1=lo, formula2=hi,
                        allow_blank=True, showErrorMessage=True, errorTitle='Value not valid', error=msg)
    ws.add_data_validation(dv)
    dv.add(ranges)


def ytd_label():
    """'Jan' for January, otherwise 'Jan–<month>'."""
    return f'IF({MON}=1,"Jan","Jan–"&{MON_SHORT})'


def ym(date_expr):
    """Locale-independent 'yyyy-mm' text (TEXT() format codes change with the user's language)."""
    return f'YEAR({date_expr})&"-"&RIGHT("0"&MONTH({date_expr}),2)'


# ---------------------------------------------------------------- dati di esempio (famiglia di fantasia)
def day(m, d):
    return dt.date(YEAR, m, min(d, calendar.monthrange(YEAR, m)[1]))


EXAMPLE_SETTINGS = {'currency': '$', 'name': 'Alex & Sam (sample data)', 'asof': dt.date(YEAR, 12, 31)}

# (goal, target, target date, starting amount, monthly plan)
EXAMPLE_GOALS = [
    ('Emergency fund', 12000, dt.date(2027, 12, 31), 6500, 300),
    ('House down payment', 30000, dt.date(2029, 12, 31), 12000, 300),
    ('New laptop', 1500, dt.date(2027, 3, 31), 200, 100),
    ('Anniversary trip', 2500, dt.date(2027, 10, 31), 0, 250),
]
# (fund, amount needed, due date, starting balance)
EXAMPLE_FUNDS = [
    ('Car maintenance', 900, dt.date(2027, 9, 30), 350),
    ('Holiday gifts', 700, dt.date(2027, 12, 1), 80),
    ('Summer vacation', 1800, dt.date(2027, 7, 1), 400),
    ('Home repairs', 600, dt.date(2027, 12, 31), 150),
]
# (debt, balance, APR, minimum payment)
EXAMPLE_DEBTS = [
    ('Credit card', 4200, 0.2499, 110),
    ('Medical bill', 480, 0, 40),
    ('Personal loan', 6200, 0.119, 210),
    ('Student loan', 14500, 0.055, 165),
    ('Car loan', 11800, 0.069, 389),
]
EXAMPLE_DEBT_OPTIONS = {'method': 'Avalanche', 'extra': 150}

# category -> (typical month, {month: different amount})
EXAMPLE_BUDGET = {
    'Paycheck 1': (4200, {}), 'Paycheck 2': (2950, {}), 'Side income': (250, {}),
    'Other income': (0, {4: 1200, 12: 300}),
    'Rent / Mortgage': (1850, {}), 'Utilities': (200, {1: 260, 2: 250, 7: 240, 8: 245, 12: 260}),
    'Internet & Phone': (135, {}), 'Insurance': (146, {}), 'Subscriptions': (72, {3: 102}),
    'Car payment': (389, {}),
    'Groceries': (720, {11: 760, 12: 820}), 'Dining out': (240, {}), 'Transport / Gas': (200, {9: 750}),
    'Health': (60, {}), 'Kids': (390, {7: 500, 8: 640}), 'Personal care': (45, {}),
    'Entertainment': (80, {}), 'Gifts': (40, {12: 650}), 'Travel': (0, {5: 200, 7: 1600}),
    'Clothing': (60, {}), 'Home & repairs': (60, {}), 'Pets': (50, {3: 210}), 'Misc': (40, {}),
    'Emergency fund': (300, {}), 'Savings goals': (400, {}), 'Sinking funds': (300, {}),
    'Retirement & investing': (600, {}), 'Debt payments': (675, {}),
}
EXAMPLE_FROM_FUNDS = {7: 1500, 9: 500, 12: 600}   # "Paid from goals & funds" planned by month


def example_transactions():
    """One year of transactions of a fictional two-earner household (amounts in the sheet currency)."""
    rnd = random.Random(2026)
    tx = []

    def add(m, d, desc, cat, amt, acc='Checking', fund='', note=''):
        tx.append((day(m, d), desc, cat, round(amt, 2), acc, fund, note))

    def u(a, b):
        return round(rnd.uniform(a, b), 2)

    side_jobs = ['Freelance design job', 'Sold old bike', 'Weekend tutoring', 'Freelance logo', 'Sold furniture online']
    heat = {1: 1.7, 2: 1.6, 3: 1.2, 4: 0.8, 5: 0.5, 6: 0.3, 7: 0.3, 8: 0.3, 9: 0.4, 10: 0.8, 11: 1.3, 12: 1.7}
    cool = {1: 0.9, 2: 0.9, 3: 0.9, 4: 0.9, 5: 1.0, 6: 1.25, 7: 1.5, 8: 1.55, 9: 1.2, 10: 0.95, 11: 0.9, 12: 1.0}
    for m in range(1, 13):
        # income
        add(m, 1, 'Alex paycheck', 'Paycheck 1', 2100)
        add(m, 15, 'Alex paycheck', 'Paycheck 1', 2100)
        add(m, 28, 'Sam paycheck', 'Paycheck 2', 2950)
        if m in (2, 3, 5, 6, 8, 9, 10, 11):
            add(m, rnd.randint(6, 24), side_jobs[m % len(side_jobs)], 'Side income', u(140, 460))
        # fixed bills
        add(m, 1, 'Rent', 'Rent / Mortgage', 1850)
        add(m, 18, 'Electric bill', 'Utilities', round(92 * cool[m] + rnd.uniform(-5, 5), 2))
        add(m, 20, 'Gas bill', 'Utilities', round(14 + 52 * heat[m] + rnd.uniform(-4, 4), 2))
        add(m, 22, 'Water & trash', 'Utilities', 54.30)
        add(m, 12, 'Internet', 'Internet & Phone', 65)
        add(m, 14, 'Phone plan', 'Internet & Phone', 70, 'Credit card')
        add(m, 5, 'Car insurance', 'Insurance', 128)
        add(m, 5, 'Renters insurance', 'Insurance', 18)
        add(m, 3, 'Video streaming', 'Subscriptions', 15.49, 'Credit card')
        add(m, 7, 'Music streaming', 'Subscriptions', 10.99, 'Credit card')
        add(m, 2, 'Gym membership', 'Subscriptions', 45)
        if m == 3:
            add(m, 9, 'Cloud storage (yearly)', 'Subscriptions', 29.99, 'Credit card')
        add(m, 10, 'Car loan payment', 'Car payment', 389)
        # variable expenses
        for d in (3, 10, 17, 24) + ((30,) if m in (1, 3, 5, 8, 10, 12) else ()):
            base = 168 if m < 11 else 185
            add(m, d, rnd.choice(['Supermarket', 'Supermarket', 'Farmers market', 'Warehouse store']), 'Groceries',
                u(base - 42, base + 32), 'Debit card')
        for _ in range(rnd.randint(5, 8)):
            add(m, rnd.randint(1, 28), rnd.choice(['Pizza night', 'Coffee', 'Brunch', 'Takeout', 'Lunch out', 'Dinner out']),
                'Dining out', u(9, 68), 'Credit card')
        for d in (4, 12, 20, 27):
            add(m, d, 'Gas station', 'Transport / Gas', u(38, 56), 'Credit card')
        if m in (2, 5, 9, 11):
            add(m, rnd.randint(5, 25), 'Parking', 'Transport / Gas', u(8, 24), 'Cash')
        if m == 3:
            add(m, 21, 'Oil change', 'Transport / Gas', 69.99, 'Debit card', 'Car maintenance')
        if m == 9:
            add(m, 16, 'New tires', 'Transport / Gas', 520, 'Debit card', 'Car maintenance')
        for _ in range(rnd.randint(0, 2)):
            add(m, rnd.randint(1, 28), rnd.choice(['Pharmacy', 'Doctor copay', 'Dentist copay']), 'Health', u(15, 120), 'Debit card')
        add(m, 1, 'After-school program' if m not in (7, 8) else 'Summer camp', 'Kids', 320 if m not in (7, 8) else 430)
        add(m, rnd.randint(8, 26), rnd.choice(['Swim lessons', 'School trip', 'Kids clothes', 'Books and toys']), 'Kids', u(30, 95), 'Debit card')
        if m == 8:
            add(m, 20, 'School supplies', 'Kids', 145.60, 'Debit card')
        for _ in range(rnd.randint(1, 2)):
            add(m, rnd.randint(1, 28), rnd.choice(['Haircut', 'Toiletries', 'Haircut']), 'Personal care', u(14, 48), 'Debit card')
        for _ in range(rnd.randint(1, 3)):
            add(m, rnd.randint(1, 28), rnd.choice(['Movie tickets', 'Concert', 'Board game', 'Museum', 'Bowling']),
                'Entertainment', u(12, 58), 'Credit card')
        if m != 12 and rnd.random() < 0.75:
            add(m, rnd.randint(1, 28), rnd.choice(['Birthday present', 'Charity donation', 'Wedding gift']), 'Gifts', u(25, 70), 'Credit card')
        if m == 12:
            for d, desc, amt in ((4, 'Holiday gifts', 214.80), (11, 'Holiday gifts', 188.45), (16, 'Holiday gifts', 162.25),
                                 (19, 'Holiday decorations', 89.50)):
                add(m, d, desc, 'Gifts', amt, 'Credit card', 'Holiday gifts')
        if m == 5:
            add(m, 23, 'Weekend trip', 'Travel', 212.40, 'Credit card')
        if m == 7:
            add(m, 6, 'Flights', 'Travel', 782, 'Credit card', 'Summer vacation')
            add(m, 6, 'Vacation rental', 'Travel', 640, 'Credit card', 'Summer vacation')
            add(m, 18, 'Rental car', 'Travel', 158, 'Credit card', 'Summer vacation')
        if rnd.random() < 0.8:
            add(m, rnd.randint(1, 28), rnd.choice(['Jeans', 'Shoes', 'Jacket', 'T-shirts']), 'Clothing', u(25, 120), 'Credit card')
        if m == 10:
            add(m, 9, 'Washer repair', 'Home & repairs', 185, 'Debit card', 'Home repairs')
        else:
            add(m, rnd.randint(1, 28), rnd.choice(['Hardware store', 'Light bulbs', 'Kitchen items', 'Garden supplies']),
                'Home & repairs', u(18, 95), 'Debit card')
        add(m, 9, 'Pet food', 'Pets', u(42, 52), 'Debit card')
        if m == 3:
            add(m, 14, 'Vet visit', 'Pets', 160, 'Debit card')
        for _ in range(rnd.randint(1, 2)):
            add(m, rnd.randint(1, 28), rnd.choice(['Post office', 'Phone case', 'Bank fee', 'Office supplies']), 'Misc', u(8, 38), 'Cash')
        # other income
        if m == 4:
            add(m, 18, 'Tax refund', 'Other income', 1240)
        if m == 12:
            add(m, 22, 'Holiday gift from family', 'Other income', 300)
        # savings and debt (transfers on the 2nd)
        add(m, 2, 'Transfer to emergency fund', 'Emergency fund', 300, 'Savings', 'Emergency fund')
        add(m, 2, 'Transfer: house fund', 'Savings goals', 300, 'Savings', 'House down payment')
        add(m, 2, 'Transfer: laptop', 'Savings goals', 100, 'Savings', 'New laptop')
        for fund, amt in (('Car maintenance', 60), ('Holiday gifts', 50), ('Summer vacation', 150), ('Home repairs', 40)):
            add(m, 2, f'Set aside: {fund.lower()}', 'Sinking funds', amt, 'Savings', fund)
        add(m, 3, 'Retirement contribution', 'Retirement & investing', 600)
        add(m, 20, 'Credit card payment + extra', 'Debt payments', 260)
        add(m, 20, 'Medical bill installment', 'Debt payments', 40)
        add(m, 21, 'Personal loan payment', 'Debt payments', 210)
        add(m, 25, 'Student loan payment', 'Debt payments', 165)
    tx.sort(key=lambda t: t[0])   # sort stabile: a parita' di data resta l'ordine di inserimento
    return tx


# ---------------------------------------------------------------- fogli
def sheet_lists(wb):
    ws = wb.create_sheet('Lists')
    put(ws, 'A1', 'Months', f=font(10, True))
    put(ws, 'B1', 'Short', f=font(10, True))
    put(ws, 'C1', 'Month to review', f=font(10, True))
    put(ws, 'C2', 'Auto')
    for i, m in enumerate(MONTHS):
        put(ws, f'A{2 + i}', m)
        put(ws, f'B{2 + i}', SHORT[i])
        put(ws, f'C{3 + i}', m)
    put(ws, 'D1', 'Types', f=font(10, True))
    for i, t in enumerate(TYPES):
        put(ws, f'D{2 + i}', t)
    put(ws, 'E1', 'Payoff methods', f=font(10, True))
    put(ws, 'E2', 'Avalanche')
    put(ws, 'E3', 'Snowball')
    put(ws, 'F1', 'Currency symbols', f=font(10, True))
    for i, c in enumerate(CURRENCIES):
        put(ws, f'F{2 + i}', c)
    # V: goal and fund names in sheet order; W: position among the non-empty ones; G: the same names without gaps,
    # so the Transactions drop-down shows funds right after the goals (no blank lines in between).
    n = N_GOAL + N_FUND
    put(ws, 'G1', 'Goals and funds (for the Transactions drop-down)', f=font(10, True))
    put(ws, 'V1', 'Goals and funds (sheet order)', f=font(10, True))
    put(ws, 'W1', 'Position', f=font(10, True))
    for i in range(N_GOAL):
        r = GOAL0 + i
        put(ws, f'V{2 + i}', f'=IF({SG}!$B${r}="","",{SG}!$B${r}&"")')
    for i in range(N_FUND):
        r = FUND0 + i
        put(ws, f'V{2 + N_GOAL + i}', f'=IF({SF}!$B${r}="","",{SF}!$B${r}&"")')
    for i in range(n):
        put(ws, f'W{2 + i}', f'=IF(V{2 + i}="","",COUNTIF($V$2:V{2 + i},"?*"))')
        put(ws, f'G{2 + i}', f'=IFERROR(INDEX($V$2:$V${1 + n},MATCH({i + 1},$W$2:$W${1 + n},0)),"")')
    put(ws, 'I1', 'Month numbers', f=font(10, True))
    for m in range(12):
        put(ws, f'{L(9 + m)}2', m + 1)
    widths(ws, {'A': 12, 'B': 8, 'C': 16, 'D': 18, 'E': 14, 'F': 16, 'G': 30, 'V': 30, 'W': 10})
    ws.sheet_state = 'hidden'
    protect(ws)
    return ws


def sheet_start(wb):
    ws = wb.create_sheet('Start Here')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 3, 'B': 22, 'C': 92})
    put(ws, 'B1', f'{YEAR} Budget Planner', f=font(22, True, TEAL))   # testo fisso: visibile anche in Protected View
    ws.row_dimensions[1].height = 36
    put(ws, 'B2', 'Plan your money month by month, log every transaction and see your whole year at a glance.',
        f=font(11, color=MUTED))
    r = 4

    def sec(text):
        nonlocal r
        r += 1
        section(ws, f'B{r}', text, span_to=3)
        r += 1

    def para(text, bullet='', bold_lead=None, size=10, color=TEXT):
        nonlocal r
        ws.merge_cells(f'B{r}:C{r}')
        lead = (bullet + ' ' if bullet else '') + (bold_lead + ' ' if bold_lead else '')
        c = put(ws, f'B{r}', None, f=font(size, color=color), al=Alignment(wrap_text=True, vertical='top'))
        if bold_lead:
            c.value = CellRichText([TextBlock(InlineFont(rFont=ARIAL, sz=size, b=True, color=TEXT), lead),
                                    TextBlock(InlineFont(rFont=ARIAL, sz=size, color=TEXT), text)])
        else:
            c.value = lead + text
        lines = max(1, -(-len(lead + text) // 128))
        ws.row_dimensions[r].height = 13.5 * lines + 4
        r += 1
        return c

    sec('Start in 6 steps')
    steps = [
        'Settings: check the year and choose your currency symbol ($, £, € ...). It appears in the column headings.',
        'Categories: rename, delete or add categories (up to 60). Give each one a type: Income, Fixed bills, '
        'Variable expenses or Savings & debt.',
        'Monthly Budget: type a typical monthly amount for each category. All 12 months fill in by themselves; '
        'type over a single month for costs that do not repeat every month (annual insurance, holiday gifts).',
        'Transactions: log every payment and every income with date, description, category (from the drop-down list) '
        'and amount. Use positive amounts; for a refund, type a negative amount in the same category.',
        'Monthly Summary and Annual Dashboard: nothing to type, they update by themselves. Pick the month to review at '
        'the top of Monthly Summary.',
        'Optional: track your Savings Goals, Sinking Funds and Debt Payoff plan.',
    ]
    for i, s in enumerate(steps, 1):
        head, rest = s.split(': ', 1)
        para(rest, bullet=f'{i}.', bold_lead=head + ':')

    sec('Color key')
    ws.merge_cells(f'B{r}:C{r}')
    put(ws, f'B{r}', 'Yellow cells: your inputs. Type over them freely.', bg=INPUT, border=B_IN,
        al=Alignment(vertical='center', indent=1))
    ws.row_dimensions[r].height = 20
    r += 1
    ws.merge_cells(f'B{r}:C{r}')
    put(ws, f'B{r}', 'White cells: formulas. The sheets are protected so you do not delete them by mistake.',
        border=B_ROW, al=Alignment(vertical='center', indent=1))
    ws.row_dimensions[r].height = 20
    r += 1
    ws.merge_cells(f'B{r}:C{r}')
    put(ws, f'B{r}', 'Differences in green are better than planned; in red, worse than planned '
        '(spending more, or earning or saving less).', f=font(10, True, GOOD), al=Alignment(vertical='center', indent=1))
    ws.row_dimensions[r].height = 20
    r += 1

    sec('Good to know')
    notes = [
        ('Dates.', 'Dates use the ISO format year-month-day (2026-03-15), which reads the same in every country. '
                   'Type dates that way. To show them differently, select the column and change the date format '
                   '(Format Cells in Excel, Format > Number in Google Sheets).'),
        ('Credit cards.', 'Log each card purchase as an expense when you make it. Do not log the monthly card bill as '
                          'well, or the same spending counts twice. Payments on an old balance you are paying off go '
                          'under Debt payments.'),
        ('Goals and sinking funds.', 'When you move money into a goal or fund, log it in a Savings & debt category '
                                     '(for example Sinking funds) and pick the goal or fund in the "Goal / fund" column. '
                                     'When you spend that money, log the expense in its normal category and pick the '
                                     'same fund: the fund balance goes down.'),
        ('Paid from goals & funds.', 'Spending covered by money you set aside earlier. It is added back in "Left over", '
                                     'so the same money is not counted twice. In Monthly Budget you can plan it '
                                     'for the months when you expect to use a fund.'),
        ('Savings rate.', 'Savings rate = (income − expenses) ÷ income: the share of your income you did not spend. '
                          'Savings, investments and debt payments count as saved; car payments are a fixed bill.'),
        ('Left over.', 'Left over = income − expenses − savings & debt + paid from goals & funds. In a zero-based '
                       'budget, "Left to budget" in Monthly Budget is close to zero.'),
        ('Debt Payoff.', 'Estimates assume that rates and minimum payments stay the same. Your lender\'s figures '
                         'can differ slightly.'),
        ('Renaming a category.', 'If you rename a category after logging transactions, change the old name in '
                                 'Transactions too (Find and Replace). Until then, those rows are flagged in the Check '
                                 'column and left out of the totals.'),
        ('Sorting.', 'Filters work while the sheets are protected. To sort Transactions, unprotect the sheet first '
                     '(no password), sort, then protect it again.'),
        ('Printing.', 'Each sheet prints one page wide. For Transactions, select the rows you need and print the '
                      'selection; otherwise all 2,000 rows print.'),
        ('Next year.', 'The file covers one calendar year. Save a copy, clear Transactions and change the year in Settings.'),
    ]
    for lead, text in notes:
        para(text, bullet='•', bold_lead=lead)

    sec('Works with')
    compat = [
        'Excel 2016 or later (Windows and Mac), Excel on the web, Google Sheets and LibreOffice Calc. '
        'No macros, no add-ins, no internet connection needed.',
        'Excel: if the file opens in Protected View (yellow bar at the top), click Enable Editing. Until then the '
        'formulas are not calculated and many cells look empty.',
        'Google Sheets: upload the file to your Drive, open it and choose File > Save as Google Sheets. '
        'Charts can look slightly different, and sheet protection may not carry over: type only in the yellow cells.',
        'Sheet protection has no password. To change formulas or the layout: Review > Unprotect Sheet in Excel, '
        'Tools > Protect Sheet in LibreOffice, Data > Protect sheets and ranges in Google Sheets.',
    ]
    for t in compat:
        para(t, bullet='•')

    sec('Sheets in this file')
    sheets = [
        ('Settings', 'Year, currency symbol, plan name, accounts and the date used for "months left".'),
        ('Categories', 'Your categories and their types.'),
        ('Transactions', 'One log for the whole year: 2,000 rows with drop-down lists and automatic checks.'),
        ('Monthly Budget', 'Planned amounts per category and month, and "Left to budget".'),
        ('Monthly Summary', 'Actual vs planned for the month you choose and the year to date, plus actuals for every month.'),
        ('Annual Dashboard', 'Income, expenses, savings rate, top spending categories and charts.'),
        ('Savings Goals', 'Targets, progress and how much to save each month.'),
        ('Sinking Funds', 'Money set aside for irregular costs: balance, amount still needed, monthly amount.'),
        ('Debt Payoff', 'Snowball or avalanche order, debt-free date and interest saved.'),
        ('Debt Schedule', 'Month-by-month balances behind Debt Payoff (calculated, nothing to type).'),
    ]
    for name, desc in sheets:
        c = put(ws, f'B{r}', name, f=font(10, True, TEAL, underline='single'), border=B_ROW, al=LEFT_C)
        c.hyperlink = Hyperlink(ref='', location=f"'{name}'!A1")
        c.value = name
        put(ws, f'C{r}', desc, border=B_ROW, al=LEFT_C)
        ws.row_dimensions[r].height = 18
        r += 1

    r += 1
    para('This planner is a personal organization tool. It is not financial, tax or legal advice.', size=9, color=MUTED)
    para(f'Made by NettoChiaro · version {VERSION} · Questions? Send us a message on Etsy.', size=9, color=MUTED)
    print_setup(ws, portrait=True, area=f'A1:C{r}')
    ws.page_setup.fitToHeight = 1
    protect(ws)
    return ws


def sheet_settings(wb, ex):
    ws = wb.create_sheet('Settings')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 3, 'B': 34, 'C': 24, 'D': 70})
    put(ws, 'B1', 'Settings', f=font(18, True, TEAL))
    ws.row_dimensions[1].height = 30
    put(ws, 'B2', 'Set these once. Yellow cells are inputs.', f=font(10, color=MUTED))
    header(ws, 3, [('B', 'Setting'), ('C', 'Your value'), ('D', 'Notes')], height=22)
    rows = [
        (4, 'Budget year', YEAR, '0', 'The calendar year of this planner. Transactions from other years are flagged.'),
        (5, 'Currency symbol', ex['currency'] if ex else '$', None,
         'Shown in the column headings, for example "Amount ($)". Pick one from the list or type your own.'),
        (6, 'Plan name (optional)', ex['name'] if ex else None, None, 'Shown on the dashboard, for example "Smith household".'),
        (7, 'As-of date (optional)', ex['asof'] if ex else None, DATE_F,
         'Leave empty to use today\'s date. Used for "months left" and payoff dates.'),
    ]
    for r, label, val, fmt, note in rows:
        put(ws, f'B{r}', label, f=font(10, True), border=B_ROW, al=LEFT_C)
        inp(ws, f'C{r}', val, fmt=fmt, al=LEFT_C)
        put(ws, f'D{r}', note, f=font(9, color=MUTED), border=B_ROW, al=Alignment(vertical='center', wrap_text=True))
        ws.row_dimensions[r].height = 26
    dv = DataValidation(type='whole', operator='between', formula1='2000', formula2='2100', allow_blank=False,
                        showErrorMessage=True, errorTitle='Year not valid', error='Enter a year between 2000 and 2100.')
    ws.add_data_validation(dv)
    dv.add('C4')
    dv_list(ws, L_CURR, 'C5', 'Currency', 'Pick a symbol or type your own.', strict=False)
    dv_date(ws, 'C7')

    section(ws, 'B9', 'Values used by the formulas (calculated)', span_to=4)
    calc = [
        (10, 'Month to review', f'={MS}!$B$3', None, 'Chosen at the top of Monthly Summary. Auto = the latest month.'),
        (11, 'As-of date used', '=IF(ISNUMBER(C7),C7,TODAY())', DATE_F, ''),
        (12, 'Month used (number)',
         f'=IFERROR(MATCH({MS}!$B$3,{L_MONTHS},0),IF(YEAR(C11)<C4,1,IF(YEAR(C11)>C4,12,MONTH(C11))))', '0',
         'Auto: the month of the as-of date; December for a past year, January for a future year.'),
        (13, 'Month used (name)', f'=INDEX({L_MONTHS},C12)', None, ''),
        (14, 'Currency in headings', '=IF(TRIM(C5)="",""," ("&TRIM(C5)&")")', None, ''),
        (15, 'Month used (short)', '=LEFT(C13,3)', None, ''),
    ]
    for r, label, f, fmt, note in calc:
        put(ws, f'B{r}', label, border=B_ROW, al=LEFT_C)
        put(ws, f'C{r}', f, f=font(10, True), fmt=fmt, border=B_ROW, al=LEFT_C)
        put(ws, f'D{r}', note, f=font(9, color=MUTED), border=B_ROW, al=LEFT_C)

    section(ws, 'B17', 'Accounts and payment methods (drop-down in Transactions)', span_to=4)
    for i in range(N_ACC):
        r = ACC0 + i
        inp(ws, f'B{r}', ACCOUNTS[i] if i < len(ACCOUNTS) else None)
    put(ws, f'C{ACC0}', 'Optional. Rename or add up to 8.', f=font(9, color=MUTED))
    print_setup(ws, area=f'A1:D{ACC1}')
    protect(ws)
    return ws


def sheet_categories(wb, cats):
    ws = wb.create_sheet('Categories')
    widths(ws, {'A': 30, 'B': 20, 'C': 56, 'D': 22})
    title(ws, 'Categories', 'Up to 60 categories. Each one needs a type. Changes show up everywhere, including the drop-down list in Transactions.')
    put(ws, 'A3', 'Types: Income · Fixed bills (same amount every month) · Variable expenses · Savings & debt (savings, '
        'investing and debt payments). Avoid the characters * ? ~ in names.', f=font(9, color=MUTED))
    header(ws, 5, [('A', 'Category'), ('B', 'Type'), ('C', 'Notes'), ('D', 'Check')], height=24)
    for i in range(N_CAT):
        r = CAT0 + i
        name, typ, note = cats[i] if i < len(cats) else (None, None, None)
        inp(ws, f'A{r}', name)
        inp(ws, f'B{r}', typ)
        inp(ws, f'C{r}', note or None)
        put(ws, f'D{r}', f'=IF(A{r}="",IF(B{r}="","","Name missing"),IF(B{r}="","Choose a type",'
                         f'IF(COUNTIF($A${CAT0}:$A${CAT1},A{r})>1,"Name used twice",'
                         f'IF(ISNA(MATCH(B{r},{L_TYPES},0)),"Unknown type",""))))',
            f=font(9, True, WARN_T), border=B_ROW)
    dv_list(ws, L_TYPES, f'B{CAT0}:B{CAT1}', 'Type not valid', 'Choose Income, Fixed bills, Variable expenses or Savings & debt.')
    ws.conditional_formatting.add(f'D{CAT0}:D{CAT1}', FormulaRule(formula=[f'D{CAT0}<>""'], fill=fill(WARN)))
    ws.freeze_panes = f'A{CAT0}'
    print_setup(ws, title_rows='5:5', area=f'A1:D{CAT1}')
    protect(ws)
    return ws


def sheet_transactions(wb, tx):
    ws = wb.create_sheet('Transactions')
    widths(ws, {'A': 12.5, 'B': 32, 'C': 24, 'D': 13, 'E': 15, 'F': 22, 'G': 24, 'H': 17, 'I': 24})
    title(ws, 'Transactions', 'One row per payment or income. Dates as year-month-day (2026-03-15). '
          'Positive amounts; a refund is a negative amount in the same category.')
    yr = f'A{TX0}:A{TX1},">="&DATE({YEAR_REF},1,1),A{TX0}:A{TX1},"<"&DATE({YEAR_REF}+1,1,1)'
    tot = lambda typ: f'SUMIFS(D{TX0}:D{TX1},H{TX0}:H{TX1},"{typ}",{yr})'  # noqa: E731
    put(ws, 'A3', f'="Rows used: "&COUNT(D{TX0}:D{TX1})&" of {N_TX:,} · Total income "&{YEAR_REF}&": "&FIXED({tot(INC)},2)'
        f'&" · Total expenses "&{YEAR_REF}&": "&FIXED({tot(FIX)}+{tot(VAR)},2)', f=font(10, True, TEAL_D))
    header(ws, 5, [('A', 'Date'), ('B', 'Description'), ('C', 'Category'), ('D', f'="Amount"&{CUR}'),
                   ('E', 'Account'), ('F', 'Goal / fund (optional)'), ('G', 'Notes'), ('H', 'Type'), ('I', 'Check')])
    for i in range(N_TX):
        r = TX0 + i
        t = tx[i] if i < len(tx) else (None,) * 7
        inp(ws, f'A{r}', t[0], fmt=DATE_F, al=Alignment(horizontal='center'))
        inp(ws, f'B{r}', t[1])
        inp(ws, f'C{r}', t[2])
        inp(ws, f'D{r}', t[3], fmt=NUM, al=RIGHT)
        inp(ws, f'E{r}', t[4])
        inp(ws, f'F{r}', t[5] or None)
        inp(ws, f'G{r}', t[6] or None)
        put(ws, f'H{r}', f'=IF(C{r}="","",IFERROR(INDEX({C_TYPES},MATCH(C{r},{C_NAMES},0))&"",""))',
            f=font(9, color=MUTED), border=B_ROW)
        put(ws, f'I{r}', f'=IF(AND(A{r}="",C{r}="",D{r}=""),"",IF(A{r}="","Date missing",IF(NOT(ISNUMBER(A{r})),"Date not recognized",'
                         f'IF(YEAR(A{r})<>{YEAR_REF},"Not in "&{YEAR_REF},IF(C{r}="","Category missing",'
                         f'IF(H{r}="",IF(COUNTIF({C_NAMES},C{r})>0,"Category has no type","Category not in list"),'
                         f'IF(D{r}="","Amount missing",IF(NOT(ISNUMBER(D{r})),"Amount is not a number",'
                         f'IF(F{r}="","",IF(ISNA(MATCH(F{r},{L_FUNDS},0)),"Goal / fund not in list",'
                         f'IF(H{r}="Income","Fund tag ignored on income","")))))))))))',
            f=font(9, True, WARN_T), border=B_ROW)
    dv_date(ws, f'A{TX0}:A{TX1}')
    dv_list(ws, C_NAMES, f'C{TX0}:C{TX1}', 'Category not in list', 'Choose a category from the list. You can edit the list in the Categories sheet.')
    dv_num(ws, f'D{TX0}:D{TX1}')
    dv_list(ws, f'Settings!$B${ACC0}:$B${ACC1}', f'E{TX0}:E{TX1}', 'Account', 'Choose an account. You can edit the list in Settings.')
    dv_list(ws, L_FUNDS, f'F{TX0}:F{TX1}', 'Goal or fund not in list', 'Choose a goal or fund. Add them in Savings Goals or Sinking Funds.')
    ws.conditional_formatting.add(f'I{TX0}:I{TX1}', FormulaRule(formula=[f'I{TX0}<>""'], fill=fill(WARN)))
    ws.auto_filter.ref = f'A5:I{TX1}'
    ws.freeze_panes = f'A{TX0}'
    print_setup(ws, title_rows='5:5')
    protect(ws, sort=True)
    return ws


SUMMARY_ROWS = [  # (row, label, tag in column B, type for SUMIFS or None)
    (6, 'Income', INC, INC),
    (7, 'Fixed bills', FIX, FIX),
    (8, 'Variable expenses', VAR, VAR),
    (9, 'Total expenses', 'All expenses', None),
    (10, 'Savings & debt payments', SAV, SAV),
    (11, 'Paid from goals & funds', 'Funds', None),
    (12, 'Left over', 'Result', None),
]


def summary_style(ws, r, label, tag, last_col, skip=()):
    bold = r in (9, 12)
    bg = SURF2 if bold else None
    put(ws, f'A{r}', label, f=font(10, True), bg=bg, border=B_TOT if bold else B_ROW)
    put(ws, f'B{r}', tag, f=font(9, color=MUTED), bg=bg, border=B_TOT if bold else B_ROW)
    for c in range(3, last_col + 1):
        if L(c) in skip:
            continue
        cell = ws.cell(r, c)
        cell.font = font(10, bold)
        cell.border = B_TOT if bold else B_ROW
        if bg:
            cell.fill = fill(bg)


def cat_rows(ws):
    for i in range(N_CAT):
        r = G0 + i
        cr = CAT0 + i
        put(ws, f'A{r}', f'=IF(Categories!$A{cr}="","",Categories!$A{cr})', border=B_ROW)
        put(ws, f'B{r}', f'=IF(A{r}="","",Categories!$B{cr}&"")', f=font(9, color=MUTED), border=B_ROW)


def diff_rules(ws, rng, first):
    """Red = worse than planned, green = better. Expenses: actual above plan is worse; income, savings, result: below."""
    col, row = first
    ref = f'{col}{row}'
    exp = f'OR($B{row}="{FIX}",$B{row}="{VAR}",$B{row}="All expenses")'
    oth = f'OR($B{row}="{INC}",$B{row}="{SAV}",$B{row}="Result")'
    red = f'AND(ISNUMBER({ref}),OR(AND({exp},{ref}>0.004),AND({oth},{ref}<-0.004)))'
    green = f'AND(ISNUMBER({ref}),OR(AND({exp},{ref}<-0.004),AND({oth},{ref}>0.004)))'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[red], font=Font(name=ARIAL, color=BAD, bold=True), fill=fill(BAD_S)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[green], font=Font(name=ARIAL, color=GOOD, bold=True)))


def sheet_budget(wb, cats, budget, from_funds):
    ws = wb.create_sheet('Monthly Budget')
    widths(ws, {'A': 30, 'B': 17, 'C': 13, **{L(c): 11 for c in range(4, 16)}, 'P': 13})
    title(ws, 'Monthly Budget', 'Type a typical monthly amount in column C: the 12 months fill in by themselves. '
          'For a month that is different (annual bills, holidays), type over that month.')
    put(ws, 'A3', 'To go back to the typical amount, copy a month that you did not change. The summary rows at the top are calculated.',
        f=font(9, color=MUTED))
    ws.merge_cells('D4:O4')
    put(ws, 'D4', f'="Planned by month"&{CUR}', f=font(10, True, TEAL_D), bg=TEAL_S, al=CENTER)
    header(ws, 5, [('A', 'Category'), ('B', 'Type'), ('C', f'="Typical month"&{CUR}'), ('P', f'="Year total"&{CUR}')])
    for m in range(12):
        put(ws, f'{L(4 + m)}5', SHORT[m], f=font(10, True, 'FFFFFF'), bg=TEAL, al=CENTER)
    put(ws, 'A13', 'Categories', f=font(10, True, TEAL_D), bg=TEAL_S)
    for c in range(2, 17):
        ws.cell(13, c).fill = fill(TEAL_S)
    cat_rows(ws)
    cols = ['C'] + [L(4 + m) for m in range(12)] + ['P']
    for r, label, tag, typ in SUMMARY_ROWS:
        summary_style(ws, r, label, tag if r != 12 else 'Result', 16)
        for X in cols:
            if typ:
                f = f'=SUMIFS({X}${G0}:{X}${G1},$B${G0}:$B${G1},"{typ}")'
            elif r == 9:
                f = f'={X}7+{X}8'
            elif r == 11:
                if X == 'C':
                    continue
                f = '=SUM(D11:O11)' if X == 'P' else None
            else:
                f = f'={X}6-{X}9-{X}10+N({X}11)'
            if f:
                ws[f'{X}{r}'] = f
            ws[f'{X}{r}'].number_format = NUMZ
    ws['A12'].value = 'Left to budget'
    for m in range(12):
        X = L(4 + m)
        inp(ws, f'{X}11', (from_funds or {}).get(m + 1), fmt=NUM, al=RIGHT)
    put(ws, 'C11', None, bg=SURF2, border=B_ROW)
    names = [c[0] for c in cats]
    for i in range(N_CAT):
        r = G0 + i
        name = names[i] if i < len(names) else None
        typical, override = budget.get(name, (None, {})) if budget else (None, {})
        inp(ws, f'C{r}', typical if typical else (0 if budget and name in budget else None), fmt=NUM, al=RIGHT)
        for m in range(12):
            X = L(4 + m)
            inp(ws, f'{X}{r}', override.get(m + 1, f'=$C{r}'), fmt=NUMZ, al=RIGHT)
        put(ws, f'P{r}', f'=SUM(D{r}:O{r})', f=font(10, True), fmt=NUMZ, border=B_ROW)
    dv_num(ws, f'C{G0}:O{G1}')
    dv_num(ws, 'D11:O11', '0', '100000000', 'Enter a positive amount.')
    ws.conditional_formatting.add('C12:P12', FormulaRule(formula=['C12<-0.004'], font=Font(name=ARIAL, color=BAD, bold=True), fill=fill(BAD_S)))
    put(ws, f'A{G1 + 2}', 'Left to budget = income − expenses − savings & debt + paid from goals & funds. '
        'Below zero means you planned to spend more than you earn.', f=font(9, color=MUTED))
    ws.freeze_panes = 'C6'
    print_setup(ws, title_rows='5:5', area=f'A1:P{G1 + 2}')
    protect(ws)
    return ws


# Monthly Summary: colonne
MS_MONTH0 = 12                       # L = gennaio ... W = dicembre
MS_ACT, MS_PLAN, MS_DIFF, MS_KEY = 'X', 'Y', 'Z', 'AA'


def sheet_summary(wb):
    ws = wb.create_sheet('Monthly Summary')
    widths(ws, {'A': 30, 'B': 17, 'C': 12, 'D': 12, 'E': 12, 'F': 9, 'G': 2, 'H': 12, 'I': 12, 'J': 12, 'K': 2,
                **{L(c): 11 for c in range(12, 24)}, 'X': 12.5, 'Y': 12.5, 'Z': 12.5, 'AA': 8})
    title(ws, 'Monthly Summary', 'Actual amounts from Transactions compared with Monthly Budget. '
          'Difference = actual − planned: green is better than planned, red is worse.')
    put(ws, 'A3', 'Month to review  ▸', f=font(10, True, TEAL_D), al=RIGHT_C)
    inp(ws, 'B3', 'Auto', al=Alignment(horizontal='center', vertical='center'))
    ws['B3'].font = font(10, True)
    ws.merge_cells('C3:J3')
    put(ws, 'C3', f'="Showing "&{MON_NAME}&" "&{YEAR_REF}&". Auto = the latest month (today, or the as-of date in Settings)."',
        f=font(9, color=MUTED), al=LEFT_C)
    ws.row_dimensions[3].height = 22
    dv_list(ws, L_CHOICES, 'B3', 'Month', 'Choose Auto or a month.')

    groups = [('C4:F4', f'={MON_NAME}&" "&{YEAR_REF}&{CUR}'),
              ('H4:J4', f'="Year to date ("&{ytd_label()}&")"&{CUR}'),
              ('L4:W4', f'="Actual by month"&{CUR}'),
              ('X4:Z4', f'="Full year"&{CUR}')]
    for rng, f in groups:
        ws.merge_cells(rng)
        put(ws, rng.split(':')[0], f, f=font(10, True, TEAL_D), bg=TEAL_S, al=CENTER)
    ws.row_dimensions[4].height = 20
    header(ws, 5, [('A', 'Category'), ('B', 'Type'), ('C', 'Planned'), ('D', 'Actual'), ('E', 'Difference'), ('F', '% of plan'),
                   ('H', 'Planned'), ('I', 'Actual'), ('J', 'Difference'),
                   ('X', 'Actual'), ('Y', 'Planned'), ('Z', 'Difference'), ('AA', 'Rank key')])
    for m in range(12):
        put(ws, f'{L(MS_MONTH0 + m)}5', SHORT[m], f=font(10, True, 'FFFFFF'), bg=TEAL, al=CENTER)
    put(ws, 'A13', 'Categories', f=font(10, True, TEAL_D), bg=TEAL_S)
    for c in range(2, 27):
        if L(c) not in ('G', 'K'):
            ws.cell(13, c).fill = fill(TEAL_S)
    cat_rows(ws)
    month_cols = [L(MS_MONTH0 + m) for m in range(12)]

    # summary rows
    for r, label, tag, typ in SUMMARY_ROWS:
        summary_style(ws, r, label, tag, 26, skip=('G', 'K'))
        for X in ['C', 'D', 'H', 'I'] + month_cols + ['X', 'Y']:
            if typ:
                f = f'=SUMIFS({X}${G0}:{X}${G1},$B${G0}:$B${G1},"{typ}")'
            elif r == 9:
                f = f'={X}7+{X}8'
            elif r == 11:
                if X in month_cols:
                    mm = month_cols.index(X) + 1
                    rng = f'{T_DATE},">="&DATE({YEAR_REF},{mm},1),{T_DATE},"<"&DATE({YEAR_REF},{mm + 1},1)'
                    f = (f'=SUMIFS({T_AMT},{T_FUND},"<>",{T_TYPE},"{FIX}",{rng})'
                         f'+SUMIFS({T_AMT},{T_FUND},"<>",{T_TYPE},"{VAR}",{rng})')
                else:
                    f = {'C': f'=INDEX({MB}!$D$11:$O$11,1,{MON})', 'D': f'=INDEX($L$11:$W$11,1,{MON})',
                         'H': f'=SUMIFS({MB}!$D$11:$O$11,{L_MNUM},"<="&{MON})', 'I': f'=SUMIFS($L$11:$W$11,{L_MNUM},"<="&{MON})',
                         'X': '=SUM(L11:W11)', 'Y': f'={MB}!$P$11'}[X]
            else:
                f = f'={X}6-{X}9-{X}10+{X}11'
            ws[f'{X}{r}'] = f
            ws[f'{X}{r}'].number_format = NUMZ
        for X, (a, p) in {'E': ('D', 'C'), 'J': ('I', 'H'), 'Z': ('X', 'Y')}.items():
            ws[f'{X}{r}'] = f'={a}{r}-{p}{r}'
            ws[f'{X}{r}'].number_format = NUMZ
        if r not in (11, 12):
            ws[f'F{r}'] = f'=IF(C{r}=0,"",D{r}/C{r})'
            ws[f'F{r}'].number_format = PCT0
    # category rows
    for i in range(N_CAT):
        r = G0 + i
        for m, X in enumerate(month_cols, 1):
            put(ws, f'{X}{r}', f'=IF($A{r}="",0,SUMIFS({T_AMT},{T_CAT},$A{r},{T_DATE},">="&DATE({YEAR_REF},{m},1),'
                               f'{T_DATE},"<"&DATE({YEAR_REF},{m + 1},1)))', fmt=NUMZ, border=B_ROW)
        cells = {
            'C': f'=IF($A{r}="",0,INDEX({MB}!$D{r}:$O{r},1,{MON}))',
            'D': f'=INDEX($L{r}:$W{r},1,{MON})',
            'E': f'=D{r}-C{r}',
            'F': f'=IF(C{r}=0,"",D{r}/C{r})',
            'H': f'=IF($A{r}="",0,SUMIFS({MB}!$D{r}:$O{r},{L_MNUM},"<="&{MON}))',
            'I': f'=SUMIFS($L{r}:$W{r},{L_MNUM},"<="&{MON})',
            'J': f'=I{r}-H{r}',
            'X': f'=SUM(L{r}:W{r})',
            'Y': f'=IF($A{r}="",0,{MB}!$P{r})',
            'Z': f'=X{r}-Y{r}',
            MS_KEY: f'=IF(AND(X{r}>0,OR(B{r}="{FIX}",B{r}="{VAR}")),X{r}+ROW()/10000000,0)',
        }
        for X, f in cells.items():
            put(ws, f'{X}{r}', f, f=font(10, X in ('D', 'X')), fmt=PCT0 if X == 'F' else ('0.00' if X == MS_KEY else NUMZ),
                border=B_ROW)
        ws[f'{MS_KEY}{r}'].font = font(8, color=MUTED)
    for col in ('E', 'J', 'Z'):
        diff_rules(ws, f'{col}6:{col}{G1}', (col, 6))
    ws.conditional_formatting.add(f'F6:F{G1}', FormulaRule(
        formula=[f'AND(OR($B6="{FIX}",$B6="{VAR}",$B6="All expenses"),ISNUMBER(F6),F6>1.00001)'],
        font=Font(name=ARIAL, color=BAD, bold=True)))
    ws.column_dimensions[MS_KEY].hidden = True
    ws.freeze_panes = 'C6'
    print_setup(ws, title_rows='4:5', area=f'A1:J{G1}')
    ws.print_title_cols = 'A:B'
    ws.print_area = [f'A1:J{G1}', f'L1:Z{G1}']
    protect(ws)
    return ws


def tile(ws, c1, c2, rows, label, value, vfmt, note, note_fmt=None, vsize=18):
    r0, r1, r2 = rows
    for r in rows:
        ws.merge_cells(f'{c1}{r}:{c2}{r}')
        for ci in range(ws[f'{c1}1'].column, ws[f'{c2}1'].column + 1):
            ws.cell(r, ci).fill = fill(TEAL_S)
    for ci in range(ws[f'{c1}1'].column, ws[f'{c2}1'].column + 1):
        ws.cell(r0, ci).border = Border(top=Side(style='thick', color=TEAL))
    put(ws, f'{c1}{r0}', label, f=font(10, True, TEAL_D), bg=TEAL_S, al=Alignment(horizontal='left', vertical='bottom', indent=1))
    put(ws, f'{c1}{r1}', value, f=font(vsize, True), bg=TEAL_S, fmt=vfmt, al=Alignment(horizontal='left', vertical='center', indent=1))
    put(ws, f'{c1}{r2}', note, f=font(9, color=MUTED), bg=TEAL_S, fmt=note_fmt,
        al=Alignment(horizontal='left', vertical='top', indent=1))


def sheet_dashboard(wb):
    ws = wb.create_sheet('Annual Dashboard')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, **{L(c): 12.5 for c in range(2, 14)}, 'N': 2})
    put(ws, 'B1', f'="Budget Planner "&{YEAR_REF}&IF({NAME_REF}="",""," · "&{NAME_REF})', f=font(20, True, TEAL))
    ws.row_dimensions[1].height = 34
    put(ws, 'B2', f'="Actual totals for "&{YEAR_REF}&" from Transactions. Year-to-date comparisons through "&{MON_NAME}&"."&'
                  f'IF({CUR}="",""," Amounts in "&TRIM({CUR_SYM})&".")', f=font(10, color=MUTED))

    pct_plan = lambda a, p: f'=IF({p}=0,"No plan yet",ROUND({a}/{p}*100,0)&"% of the yearly plan")'  # noqa: E731
    inc, exp, sav, left = f'{MS}!$X$6', f'{MS}!$X$9', f'{MS}!$X$10', f'{MS}!$X$12'
    pinc, pexp, psav, pleft = f'{MS}!$Y$6', f'{MS}!$Y$9', f'{MS}!$Y$10', f'{MS}!$Y$12'
    tiles = [
        ('B', 'C', 'Income', f'={inc}', NUM, pct_plan(inc, pinc)),
        ('D', 'E', 'Expenses', f'={exp}', NUM, pct_plan(exp, pexp)),
        ('F', 'G', 'Savings & debt', f'={sav}', NUM, pct_plan(sav, psav)),
        ('H', 'I', 'Savings rate', f'=IF({inc}=0,0,({inc}-{exp})/{inc})', PCT,
         f'="Planned: "&IF({pinc}=0,"–",ROUND(({pinc}-{pexp})/{pinc}*100,1)&"%")'),
        ('J', 'K', 'Left over', f'={left}', NUM, f'="Planned: "&FIXED({pleft},2)'),
        ('L', 'M', 'Total debt', f'={DP}!$C${DEBT1 + 1}', NUM, f'="Minimums: "&FIXED({DP}!$E${DEBT1 + 1},2)&" a month"'),
    ]
    for c1, c2, label, value, fmt, note in tiles:
        tile(ws, c1, c2, (4, 5, 6), label, value, fmt, note)
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 32
    ws.row_dimensions[6].height = 18
    wide = [
        ('B', 'E', 'Savings goals: saved so far', f'={SG}!$G${GOAL1 + 1}', NUM,
         f'=IF({SG}!$C${GOAL1 + 1}=0,"No goals yet","of "&FIXED({SG}!$C${GOAL1 + 1},2)&" in targets ("&ROUND({SG}!$I${GOAL1 + 1}*100,0)&"%)")'),
        ('F', 'I', 'Sinking funds: balance', f'={SF}!$H${FUND1 + 1}', NUM,
         f'=IF({SF}!$C${FUND1 + 1}=0,"No funds yet","Set aside "&FIXED({SF}!$M${FUND1 + 1},2)&" a month to be ready on time")'),
        ('J', 'M', 'Debt-free (your plan)', f'={DP}!$K${DEBT1 + 1}', MONTH_F,
         f'=IF({DP}!$C${DEBT1 + 1}=0,"No debts entered",{DP}!$C$4&" · interest saved: "&IF(ISNUMBER({DP}!$M$28),FIXED({DP}!$M$28,2),"–"))'),
    ]
    for c1, c2, label, value, fmt, note in wide:
        tile(ws, c1, c2, (8, 9, 10), label, value, fmt, note, vsize=16)
    ws['J9'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[8].height = 22
    ws.row_dimensions[9].height = 28
    ws.row_dimensions[10].height = 18

    # month by month
    section(ws, 'B12', 'Month by month', span_to=7)
    header(ws, 13, [('B', 'Month'), ('C', f'="Income"&{CUR}'), ('D', f'="Expenses"&{CUR}'), ('E', f'="Savings & debt"&{CUR}'),
                    ('F', f'="Left over"&{CUR}'), ('G', 'Savings rate')], height=30)
    for m in range(12):
        r = 14 + m
        X = L(MS_MONTH0 + m)
        put(ws, f'B{r}', SHORT[m], border=B_ROW)
        for col, src in (('C', 6), ('D', 9), ('E', 10), ('F', 12)):
            put(ws, f'{col}{r}', f'={MS}!{X}${src}', fmt=NUMZ, border=B_ROW)
        put(ws, f'G{r}', f'=IF(C{r}=0,"",(C{r}-D{r})/C{r})', fmt=PCT, border=B_ROW)
    put(ws, 'B26', 'Year', f=font(10, True), bg=SURF2, border=B_TOT)
    for col in 'CDEF':
        put(ws, f'{col}26', f'=SUM({col}14:{col}25)', f=font(10, True), bg=SURF2, fmt=NUMZ, border=B_TOT)
    put(ws, 'G26', '=IF(C26=0,"",(C26-D26)/C26)', f=font(10, True), bg=SURF2, fmt=PCT, border=B_TOT)
    ws.conditional_formatting.add('F14:F26', FormulaRule(formula=['F14<-0.004'], font=Font(name=ARIAL, color=BAD, bold=True)))

    ch = BarChart()
    ch.type, ch.grouping = 'col', 'clustered'
    ch.title = chart_title('Income and expenses by month')
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=13, max_row=25), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=14, max_row=25))
    for s, color in zip(ch.series, (CH_INC, CH_EXP)):
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    ch.gapWidth = 60
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    light_grid(ch.y_axis)
    ch.legend.position = 'b'
    ch.width, ch.height = 12.4, 8.0
    ws.add_chart(ch, 'I12')

    # year to date
    section(ws, 'B28', f'="Year to date: planned vs actual ("&{ytd_label()}&")"', span_to=7)
    ws.merge_cells('B29:C29')
    header(ws, 29, [('B', 'Line'), ('D', 'Planned'), ('E', 'Actual'), ('F', 'Difference'), ('G', '% of plan')], height=22)
    ytd = [(30, 'Income', 6), (31, 'Fixed bills', 7), (32, 'Variable expenses', 8), (33, 'Total expenses', 9),
           (34, 'Savings & debt', 10), (35, 'Paid from funds', 11), (36, 'Left over', 12)]
    tags = {6: INC, 7: FIX, 8: VAR, 9: 'All expenses', 10: SAV, 11: 'Funds', 12: 'Result'}
    for r, label, src in ytd:
        ws.merge_cells(f'B{r}:C{r}')
        bold = src in (9, 12)
        bg = SURF2 if bold else None
        put(ws, f'B{r}', label, f=font(10, bold), bg=bg, border=B_ROW)
        ws[f'C{r}'].border = B_ROW
        put(ws, f'D{r}', f'={MS}!$H${src}', f=font(10, bold), bg=bg, fmt=NUMZ, border=B_ROW)
        put(ws, f'E{r}', f'={MS}!$I${src}', f=font(10, bold), bg=bg, fmt=NUMZ, border=B_ROW)
        put(ws, f'F{r}', f'=E{r}-D{r}', f=font(10, bold), bg=bg, fmt=NUMZ, border=B_ROW)
        put(ws, f'G{r}', f'=IF(OR(D{r}=0,{src}>10),"",E{r}/D{r})', f=font(10, bold), bg=bg, fmt=PCT0, border=B_ROW)
        put(ws, f'N{r}', tags[src], f=font(8, color='FFFFFF'))   # tipo per i colori (testo bianco, fuori stampa)
    exp = 'OR($N30="Fixed bills",$N30="Variable expenses",$N30="All expenses")'
    oth = 'OR($N30="Income",$N30="Savings & debt",$N30="Result")'
    ws.conditional_formatting.add('F30:F36', FormulaRule(formula=[f'OR(AND({exp},F30>0.004),AND({oth},F30<-0.004))'],
                                                         font=Font(name=ARIAL, color=BAD, bold=True)))
    ws.conditional_formatting.add('F30:F36', FormulaRule(formula=[f'OR(AND({exp},F30<-0.004),AND({oth},F30>0.004))'],
                                                         font=Font(name=ARIAL, color=GOOD, bold=True)))
    put(ws, 'B37', 'Difference = actual − planned. Green is better than planned, red is worse.', f=font(9, color=MUTED))

    # checks
    section(ws, 'I28', 'Checks', span_to=13)
    checks = [
        (29, 'Transactions logged', f'=COUNT({T_AMT})'),
        (30, 'Transaction rows to check', f'=COUNTIF({T_CHK},"?*")'),
        (31, 'Categories to check', f'=COUNTIF(Categories!$D${CAT0}:$D${CAT1},"?*")'),
        (32, 'Goals and funds to check', f'=COUNTIF({SG}!$O${GOAL0}:$O${GOAL1},"Name*")+COUNTIF({SF}!$N${FUND0}:$N${FUND1},"Name*")'),
        (33, 'Debts to check', f'=COUNTIF({DP}!$L${DEBT0}:$L${DEBT1},"?*")'),
    ]
    for r, label, f in checks:
        ws.merge_cells(f'I{r}:L{r}')
        put(ws, f'I{r}', label, border=B_ROW)
        for cc in 'JKL':
            ws[f'{cc}{r}'].border = B_ROW
        put(ws, f'M{r}', f, f=font(10, True), fmt='0', al=RIGHT_C, border=B_ROW)
    ws.conditional_formatting.add('M30:M33', FormulaRule(formula=['M30>0'], fill=fill(WARN)))
    ws.merge_cells('I34:M35')
    put(ws, 'I34', '=IF(SUM(M30:M33)=0,"No problems found.","Some rows need a look: see the Check column in '
                   'Transactions, Categories, Savings Goals, Sinking Funds or Debt Payoff.")', f=font(9, color=MUTED), al=WRAP)

    # top 10
    section(ws, 'B39', 'Top 10 spending categories (year)', span_to=7)
    ws.merge_cells('C40:E40')
    header(ws, 40, [('B', '#'), ('C', 'Category'), ('F', f'="Actual"&{CUR}'), ('G', '% of expenses')], height=30)
    key = f'{MS}!${MS_KEY}${G0}:${MS_KEY}${G1}'
    for k in range(1, 11):
        r = 40 + k
        ws.merge_cells(f'C{r}:E{r}')
        put(ws, f'B{r}', k, f=font(10, True, TEAL), al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'C{r}', f'=IF(LARGE({key},B{r})<=0,"",INDEX({MS}!$A${G0}:$A${G1},MATCH(LARGE({key},B{r}),{key},0)))', border=B_ROW)
        for cc in 'DE':
            ws[f'{cc}{r}'].border = B_ROW
        put(ws, f'F{r}', f'=IF(C{r}="",0,INDEX({MS}!$X${G0}:$X${G1},MATCH(LARGE({key},B{r}),{key},0)))', fmt=NUMZ, border=B_ROW)
        put(ws, f'G{r}', f'=IF(OR(F{r}=0,{exp_ref()}=0),"",F{r}/{exp_ref()})', fmt=PCT, border=B_ROW)
    ch2 = BarChart()
    ch2.type = 'bar'
    ch2.title = chart_title('Top spending categories')
    ch2.add_data(Reference(ws, min_col=6, min_row=40, max_row=50), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=3, min_row=41, max_row=50))
    ch2.series[0].graphicalProperties.solidFill = TEAL
    ch2.series[0].graphicalProperties.line.solidFill = TEAL
    ch2.legend = None
    ch2.gapWidth = 50
    ch2.x_axis.scaling.orientation = 'maxMin'
    ch2.x_axis.delete = False
    ch2.y_axis.delete = False
    ch2.y_axis.numFmt = '#,##0'
    ch2.y_axis.crosses = 'max'
    light_grid(ch2.y_axis)
    ch2.width, ch2.height = 12.4, 6.8
    ws.add_chart(ch2, 'I39')
    print_setup(ws, area='A1:M51')
    ws.page_setup.fitToHeight = 1
    protect(ws)
    return ws


def exp_ref():
    return f'{MS}!$X$9'


def months_left(date_ref):
    return f'MAX(0,(YEAR({date_ref})-YEAR({ASOF}))*12+MONTH({date_ref})-MONTH({ASOF}))'


def bar(pct_ref):
    n = f'ROUND(MIN(1,MAX(0,N({pct_ref})))*10,0)'
    return f'REPT("■",{n})&REPT("□",10-{n})'


def fund_sum(name_ref, typ):
    return f'SUMIFS({T_AMT},{T_FUND},{name_ref},{T_TYPE},"{typ}")'


def sheet_goals(wb, goals):
    ws = wb.create_sheet('Savings Goals')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, 'B': 26, 'C': 13, 'D': 12, 'E': 13, 'F': 14, 'G': 13, 'H': 13, 'I': 9, 'J': 14,
                'K': 8, 'L': 13, 'M': 13, 'N': 12, 'O': 20})
    title(ws, 'Savings Goals', 'Long-term goals: emergency fund, a home, a car. To add money, log a Savings & debt '
          'transaction and pick the goal in the "Goal / fund" column of Transactions.', col='B')
    put(ws, 'B3', f'="Months left are counted from "&{ym(ASOF)}&" (as-of date in Settings)."', f=font(9, color=MUTED))
    header(ws, 5, [('B', 'Goal'), ('C', f'="Target"&{CUR}'), ('D', 'Target date'), ('E', f'="Starting amount"&{CUR}'),
                   ('F', f'="Added in Transactions"&{CUR}'), ('G', f'="Saved so far"&{CUR}'), ('H', f'="Still needed"&{CUR}'),
                   ('I', 'Progress'), ('J', ''), ('K', 'Months left'), ('L', f'="Needed per month"&{CUR}'),
                   ('M', f'="Your monthly plan"&{CUR}'), ('N', 'Reached by (your plan)'), ('O', 'Status')], height=44)
    ws.merge_cells('I5:J5')
    for i in range(N_GOAL):
        r = GOAL0 + i
        g = goals[i] if goals and i < len(goals) else (None,) * 5
        inp(ws, f'B{r}', g[0])
        inp(ws, f'C{r}', g[1], fmt=NUM, al=RIGHT)
        inp(ws, f'D{r}', g[2], fmt=DATE_F, al=Alignment(horizontal='center'))
        inp(ws, f'E{r}', g[3], fmt=NUM, al=RIGHT)
        put(ws, f'F{r}', f'=IF(B{r}="",0,{fund_sum(f"B{r}", SAV)}-{fund_sum(f"B{r}", FIX)}-{fund_sum(f"B{r}", VAR)})', fmt=NUMZ, border=B_ROW)
        put(ws, f'G{r}', f'=IF(B{r}="",0,N(E{r})+F{r})', f=font(10, True), fmt=NUMZ, border=B_ROW)
        put(ws, f'H{r}', f'=IF(B{r}="",0,MAX(0,N(C{r})-G{r}))', fmt=NUMZ, border=B_ROW)
        put(ws, f'I{r}', f'=IF(N(C{r})<=0,"",MIN(1,MAX(0,G{r}/C{r})))', fmt=PCT0, border=B_ROW, al=RIGHT)
        put(ws, f'J{r}', f'=IF(I{r}="","",{bar(f"I{r}")})', f=font(9, color=TEAL), border=B_ROW)
        put(ws, f'K{r}', f'=IF(OR(B{r}="",NOT(ISNUMBER(D{r}))),"",{months_left(f"D{r}")})', fmt='0',
            al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'L{r}', f'=IF(OR(B{r}="",K{r}=""),"",IF(H{r}<=0,0,H{r}/MAX(1,K{r})))', fmt=NUMZ, border=B_ROW)
        inp(ws, f'M{r}', g[4], fmt=NUM, al=RIGHT)
        put(ws, f'N{r}', f'=IF(OR(B{r}="",N(C{r})<=0),"",IF(H{r}<=0,"Reached",IF(N(M{r})<=0,"",'
                         f'DATE(YEAR({ASOF}),MONTH({ASOF})+ROUNDUP(H{r}/M{r},0),1))))', fmt=MONTH_F,
            al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'O{r}', f'=IF(B{r}="",IF(OR(N(C{r})<>0,N(E{r})<>0),"Name missing",""),IF(COUNTIF({L_FUNDS},B{r})>1,"Name used twice",'
                         f'IF(N(C{r})<=0,"Enter a target",IF(H{r}<=0,"Reached",IF(K{r}="","Enter a target date",'
                         f'IF(K{r}=0,"Target date reached",IF(N(M{r})<=0,"Enter a monthly plan",'
                         f'IF(M{r}>=L{r}-0.005,"On track","Behind"))))))))',
            f=font(10, True), border=B_ROW)
    t = GOAL1 + 1
    put(ws, f'B{t}', 'Total', f=font(10, True), bg=SURF2, border=B_TOT)
    for col, f, fmt in [('C', f'=SUM(C{GOAL0}:C{GOAL1})', NUMZ), ('D', None, None), ('E', f'=SUM(E{GOAL0}:E{GOAL1})', NUMZ),
                        ('F', f'=SUM(F{GOAL0}:F{GOAL1})', NUMZ), ('G', f'=SUM(G{GOAL0}:G{GOAL1})', NUMZ),
                        ('H', f'=SUM(H{GOAL0}:H{GOAL1})', NUMZ), ('I', f'=IF(C{t}=0,0,MIN(1,G{t}/C{t}))', PCT0),
                        ('J', f'=IF(C{t}=0,"",{bar(f"I{t}")})', None), ('K', None, None),
                        ('L', f'=SUM(L{GOAL0}:L{GOAL1})', NUMZ), ('M', f'=SUM(M{GOAL0}:M{GOAL1})', NUMZ), ('N', None, None), ('O', None, None)]:
        put(ws, f'{col}{t}', f, f=font(10, True, TEAL if col == 'J' else TEXT), bg=SURF2, fmt=fmt, border=B_TOT)
    put(ws, f'B{t + 2}', 'Saved so far = starting amount + Savings & debt transactions tagged with the goal − expenses tagged with it. '
        'Needed per month = still needed ÷ months left.', f=font(9, color=MUTED))
    put(ws, f'B{t + 3}', 'Starting amount: what the goal already had before you started logging it here. Type the goal name exactly '
        'as in Transactions (the drop-down list uses these names).', f=font(9, color=MUTED))
    dv_num(ws, f'C{GOAL0}:C{GOAL1}', '0', '100000000', 'Enter a positive amount.')
    dv_date(ws, f'D{GOAL0}:D{GOAL1}')
    dv_num(ws, f'E{GOAL0}:E{GOAL1}', '-100000000', '100000000')
    dv_num(ws, f'M{GOAL0}:M{GOAL1}', '0', '100000000', 'Enter a positive amount.')
    rng = f'O{GOAL0}:O{GOAL1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(O{GOAL0}="Reached",O{GOAL0}="On track")'], font=Font(name=ARIAL, color=GOOD, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'O{GOAL0}="Behind"'], font=Font(name=ARIAL, color=BAD, bold=True), fill=fill(BAD_S)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT(O{GOAL0},4)="Name",LEFT(O{GOAL0},5)="Enter",LEFT(O{GOAL0},6)="Target")'],
                                                   font=Font(name=ARIAL, color=WARN_T, bold=True), fill=fill(WARN)))
    ws.freeze_panes = f'C{GOAL0}'
    print_setup(ws, area=f'A1:O{t + 3}')
    protect(ws)
    return ws


def sheet_funds(wb, funds):
    ws = wb.create_sheet('Sinking Funds')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, 'B': 26, 'C': 13, 'D': 12, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 9, 'K': 14,
                'L': 8, 'M': 14, 'N': 18})
    title(ws, 'Sinking Funds', 'Money you set aside every month for costs you know are coming: car repairs, gifts, '
          'vacations, annual bills.', col='B')
    put(ws, 'B3', 'Put money in: a Sinking funds transaction tagged with the fund. Spend it: log the expense in its own '
        'category and tag the same fund.', f=font(9, color=MUTED))
    header(ws, 5, [('B', 'Fund'), ('C', f'="Amount needed"&{CUR}'), ('D', 'Due date'), ('E', f'="Starting balance"&{CUR}'),
                   ('F', f'="Added"&{CUR}'), ('G', f'="Spent"&{CUR}'), ('H', f'="Balance"&{CUR}'),
                   ('I', f'="Still needed"&{CUR}'), ('J', 'Progress'), ('K', ''), ('L', 'Months left'),
                   ('M', f'="Set aside per month"&{CUR}'), ('N', 'Status')], height=44)
    ws.merge_cells('J5:K5')
    for i in range(N_FUND):
        r = FUND0 + i
        g = funds[i] if funds and i < len(funds) else (None,) * 4
        inp(ws, f'B{r}', g[0])
        inp(ws, f'C{r}', g[1], fmt=NUM, al=RIGHT)
        inp(ws, f'D{r}', g[2], fmt=DATE_F, al=Alignment(horizontal='center'))
        inp(ws, f'E{r}', g[3], fmt=NUM, al=RIGHT)
        put(ws, f'F{r}', f'=IF(B{r}="",0,{fund_sum(f"B{r}", SAV)})', fmt=NUMZ, border=B_ROW)
        put(ws, f'G{r}', f'=IF(B{r}="",0,{fund_sum(f"B{r}", FIX)}+{fund_sum(f"B{r}", VAR)})', fmt=NUMZ, border=B_ROW)
        put(ws, f'H{r}', f'=IF(B{r}="",0,N(E{r})+F{r}-G{r})', f=font(10, True), fmt=NUMZ, border=B_ROW)
        put(ws, f'I{r}', f'=IF(B{r}="",0,MAX(0,N(C{r})-H{r}))', fmt=NUMZ, border=B_ROW)
        put(ws, f'J{r}', f'=IF(N(C{r})<=0,"",MIN(1,MAX(0,H{r}/C{r})))', fmt=PCT0, border=B_ROW, al=RIGHT)
        put(ws, f'K{r}', f'=IF(J{r}="","",{bar(f"J{r}")})', f=font(9, color=TEAL), border=B_ROW)
        put(ws, f'L{r}', f'=IF(OR(B{r}="",NOT(ISNUMBER(D{r}))),"",{months_left(f"D{r}")})', fmt='0',
            al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'M{r}', f'=IF(OR(B{r}="",L{r}=""),"",IF(I{r}<=0,0,I{r}/MAX(1,L{r})))', f=font(10, True), fmt=NUMZ, border=B_ROW)
        put(ws, f'N{r}', f'=IF(B{r}="",IF(OR(N(C{r})<>0,N(E{r})<>0),"Name missing",""),IF(COUNTIF({L_FUNDS},B{r})>1,"Name used twice",'
                         f'IF(N(C{r})<=0,"Enter an amount",IF(H{r}<-0.004,"Overspent",IF(I{r}<=0.004,"Ready",'
                         f'IF(L{r}="","Enter a due date",IF(L{r}=0,"Due now","Saving")))))))',
            f=font(10, True), border=B_ROW)
    t = FUND1 + 1
    put(ws, f'B{t}', 'Total', f=font(10, True), bg=SURF2, border=B_TOT)
    for col, f, fmt in [('C', f'=SUM(C{FUND0}:C{FUND1})', NUMZ), ('D', None, None), ('E', f'=SUM(E{FUND0}:E{FUND1})', NUMZ),
                        ('F', f'=SUM(F{FUND0}:F{FUND1})', NUMZ), ('G', f'=SUM(G{FUND0}:G{FUND1})', NUMZ),
                        ('H', f'=SUM(H{FUND0}:H{FUND1})', NUMZ), ('I', f'=SUM(I{FUND0}:I{FUND1})', NUMZ),
                        ('J', f'=IF(C{t}=0,0,MIN(1,MAX(0,H{t}/C{t})))', PCT0), ('K', f'=IF(C{t}=0,"",{bar(f"J{t}")})', None),
                        ('L', None, None), ('M', f'=SUM(M{FUND0}:M{FUND1})', NUMZ), ('N', None, None)]:
        put(ws, f'{col}{t}', f, f=font(10, True, TEAL if col == 'K' else TEXT), bg=SURF2, fmt=fmt, border=B_TOT)
    b = t + 2
    put(ws, f'B{b}', '="Planned in Monthly Budget for the Sinking funds category in "&' + MON_NAME + '&":"', border=B_ROW)
    for cc in 'CDEFGHIJKL':
        ws[f'{cc}{b}'].border = B_ROW
    put(ws, f'M{b}', f'=IFERROR(INDEX({MB}!$D${G0}:$O${G1},MATCH("Sinking funds",{MB}!$A${G0}:$A${G1},0),{MON}),"")',
        f=font(10, True), fmt=NUMZ, border=B_ROW)
    put(ws, f'N{b}', f'=IF(OR(M{t}<=0,M{b}=""),"",IF(M{b}>=M{t}-0.005,"Enough","Plan more"))', f=font(10, True), border=B_ROW)
    put(ws, f'B{b + 2}', 'Balance = starting balance + money added − money spent (transactions tagged with the fund). '
        'Set aside per month = still needed ÷ months left.', f=font(9, color=MUTED))
    put(ws, f'B{b + 3}', 'After the due date, change it to the next one (for example next year): the fund starts filling up again.',
        f=font(9, color=MUTED))
    dv_num(ws, f'C{FUND0}:C{FUND1}', '0', '100000000', 'Enter a positive amount.')
    dv_date(ws, f'D{FUND0}:D{FUND1}')
    dv_num(ws, f'E{FUND0}:E{FUND1}', '-100000000', '100000000')
    rng = f'N{FUND0}:N{FUND1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(N{FUND0}="Ready",N{FUND0}="Saving")'], font=Font(name=ARIAL, color=GOOD, bold=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(N{FUND0}="Overspent",N{FUND0}="Due now")'],
                                                   font=Font(name=ARIAL, color=BAD, bold=True), fill=fill(BAD_S)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT(N{FUND0},4)="Name",LEFT(N{FUND0},5)="Enter")'],
                                                   font=Font(name=ARIAL, color=WARN_T, bold=True), fill=fill(WARN)))
    ws.conditional_formatting.add(f'N{b}', FormulaRule(formula=[f'N{b}="Plan more"'], font=Font(name=ARIAL, color=BAD, bold=True)))
    ws.conditional_formatting.add(f'N{b}', FormulaRule(formula=[f'N{b}="Enough"'], font=Font(name=ARIAL, color=GOOD, bold=True)))
    ws.freeze_panes = f'C{FUND0}'
    print_setup(ws, area=f'A1:N{b + 3}')
    protect(ws)
    return ws


# Debt Schedule: colonne dei tre blocchi (minimi, snowball, avalanche)
MIN_BAL = [L(3 + k) for k in range(N_DEBT)]          # C..L
MIN_INT, MIN_TOT = 'M', 'N'
SNOW = {'need': [L(16 + k) for k in range(N_DEBT)], 'bal': [L(26 + k) for k in range(N_DEBT)],
        'sumb': 'AJ', 'pool': 'AK', 'int': 'AL', 'tot': 'AM', 'order': 'J'}
AVAL = {'need': [L(41 + k) for k in range(N_DEBT)], 'bal': [L(51 + k) for k in range(N_DEBT)],
        'sumb': 'BI', 'pool': 'BJ', 'int': 'BK', 'tot': 'BL', 'order': 'I'}


def months_formula(col):
    rng = f'{col}${S0 + 1}:{col}${S1}'
    return (f'=IF({col}${S0}<0.005,"",IF(COUNTIF({rng},">0.005")>={HORIZON},"Over 30 years",'
            f'COUNTIF({rng},">0.005")+1))')


def sheet_schedule(wb):
    ws = wb.create_sheet('Debt Schedule')
    ws.sheet_properties.tabColor = 'B8C2BD'
    widths(ws, {'A': 8, 'B': 10, **{c: 11 for c in MIN_BAL}, MIN_INT: 10, MIN_TOT: 12, 'O': 2, 'AN': 2})
    put(ws, 'A1', 'Debt Schedule (calculated)', f=font(16, True, TEAL))
    ws.row_dimensions[1].height = 26
    put(ws, 'A2', 'Month-by-month balances behind Debt Payoff. Nothing to type here. Each month: interest is added, '
        'every debt gets its minimum payment, the rest of your monthly total goes to the first debt in the payoff order.',
        f=font(9, color=MUTED))
    put(ws, 'A3', f'="Month 0 = "&{ym(ASOF)}&" (as-of date in Settings). Payments start the month after."', f=font(9, color=MUTED))
    labels = {5: 'Debt', 6: 'APR', 7: 'Minimum', 8: 'Months to payoff'}
    for r, lab in labels.items():
        put(ws, f'B{r}', lab, f=font(9, True, MUTED), al=RIGHT)
    header(ws, 9, [('A', 'Month'), ('B', 'Date')], height=30)

    def block_title(c1, c2, text):
        ws.merge_cells(f'{c1}4:{c2}4')
        put(ws, f'{c1}4', text, f=font(10, True, TEAL_D), bg=TEAL_S, al=CENTER)

    block_title('C', MIN_TOT, 'Minimum payments only (debts in the order you typed them)')
    block_title(SNOW['bal'][0], SNOW['tot'], 'Snowball: smallest balance first')
    block_title(AVAL['bal'][0], AVAL['tot'], 'Avalanche: highest APR first')

    # mese 0 e intestazioni
    for t in range(HORIZON + 1):
        r = S0 + t
        ws.cell(r, 1, t).font = font(9, color=MUTED)
        c = ws.cell(r, 2, f'=DATE(YEAR({ASOF}),MONTH({ASOF})+A{r},1)')
        c.number_format = MONTH_F
        c.font = font(9, color=MUTED)

    # blocco minimi (ordine originale)
    for k, col in enumerate(MIN_BAL):
        dr = DEBT0 + k
        valid = f'AND({DP}!$B${dr}<>"",N({DP}!$C${dr})>0,ISNUMBER({DP}!$D${dr}))'
        put(ws, f'{col}5', f'=IF({valid},{DP}!$B${dr},"")', f=font(9, True), al=CENTER)
        put(ws, f'{col}6', f'=IF({valid},{DP}!$D${dr},0)', f=font(9), fmt=PCT_IN, al=RIGHT)
        put(ws, f'{col}7', f'=IF({valid},N({DP}!$E${dr}),0)', f=font(9), fmt=NUM, al=RIGHT)
        put(ws, f'{col}8', months_formula(col), f=font(9, True), al=RIGHT)
        put(ws, f'{col}{S0}', f'=IF({valid},{DP}!$C${dr},0)', f=font(9), fmt=NUM)
        ws[f'{col}5'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[5].height = 26
    header(ws, 9, [(c, 'Balance') for c in MIN_BAL] + [(MIN_INT, 'Interest'), (MIN_TOT, 'Total balance')], height=30)
    put(ws, f'{MIN_TOT}{S0}', f'=SUM(C{S0}:L{S0})', f=font(9, True), fmt=NUM)
    put(ws, f'{MIN_TOT}8', months_formula(MIN_TOT).replace('""', '0', 1), f=font(9, True), al=RIGHT)

    def ordered_block(B):
        order = f'{DP}!${B["order"]}${DEBT0}:${B["order"]}${DEBT1}'
        for k, col in enumerate(B['bal']):
            m = f'MATCH({k + 1},{order},0)'
            put(ws, f'{col}5', f'=IFERROR(INDEX({DP}!$B${DEBT0}:$B${DEBT1},{m}),"")', f=font(9, True),
                al=Alignment(horizontal='center', vertical='center', wrap_text=True))
            put(ws, f'{col}6', f'=IFERROR(INDEX({DP}!$D${DEBT0}:$D${DEBT1},{m})+0,0)', f=font(9), fmt=PCT_IN, al=RIGHT)
            put(ws, f'{col}7', f'=IFERROR(INDEX({DP}!$E${DEBT0}:$E${DEBT1},{m})+0,0)', f=font(9), fmt=NUM, al=RIGHT)
            put(ws, f'{col}8', months_formula(col), f=font(9, True), al=RIGHT)
            put(ws, f'{col}{S0}', f'=IFERROR(INDEX({DP}!$C${DEBT0}:$C${DEBT1},{m})+0,0)', f=font(9), fmt=NUM)
        b0, b9 = B['bal'][0], B['bal'][-1]
        n0, n9 = B['need'][0], B['need'][-1]
        header(ws, 9, [(c, 'Balance') for c in B['bal']] + [(c, 'After minimum') for c in B['need']] +
               [(B['sumb'], 'Before payment'), (B['pool'], 'Left after minimums'), (B['int'], 'Interest'), (B['tot'], 'Total balance')],
               height=30)
        put(ws, f'{B["pool"]}6', 'Paid each month', f=font(9, True, MUTED), al=RIGHT)
        put(ws, f'{B["pool"]}7', f'=SUM({b0}7:{b9}7)+N({DP}!$C$5)', f=font(9, True), fmt=NUM)
        put(ws, f'{B["tot"]}{S0}', f'=SUM({b0}{S0}:{b9}{S0})', f=font(9, True), fmt=NUM)
        put(ws, f'{B["tot"]}8', months_formula(B['tot']).replace('""', '0', 1), f=font(9, True), al=RIGHT)
        pool_ref = f'${B["pool"]}$7'
        for t in range(1, HORIZON + 1):
            r = S0 + t
            p = r - 1
            for k in range(N_DEBT):
                bc, nc = B['bal'][k], B['need'][k]
                ws[f'{nc}{r}'] = f'=IF({bc}{p}<0.005,0,MAX(0,{bc}{p}*(1+{bc}$6/12)-{bc}$7))'
            ws[f'{B["int"]}{r}'] = f'=SUMPRODUCT({b0}{p}:{b9}{p},${b0}$6:${b9}$6)/12'
            ws[f'{B["sumb"]}{r}'] = f'=SUM({b0}{p}:{b9}{p})+{B["int"]}{r}'
            ws[f'{B["pool"]}{r}'] = f'={pool_ref}-({B["sumb"]}{r}-SUM({n0}{r}:{n9}{r}))'
            for k in range(N_DEBT):
                bc, nc = B['bal'][k], B['need'][k]
                before = '0' if k == 0 else f'SUM({n0}{r}:{B["need"][k - 1]}{r})'
                ws[f'{bc}{r}'] = f'=MAX(0,{nc}{r}-MAX(0,{B["pool"]}{r}-{before}))'
            ws[f'{B["tot"]}{r}'] = f'=SUM({b0}{r}:{b9}{r})'
            for c in B['bal'] + [B['int'], B['tot']]:
                ws[f'{c}{r}'].number_format = NUM
                ws[f'{c}{r}'].font = font(9, c == B['tot'])
            for c in B['need'] + [B['sumb'], B['pool']]:
                ws[f'{c}{r}'].number_format = NUM
                ws[f'{c}{r}'].font = font(9, color=MUTED)
        for c in B['need'] + [B['sumb']]:
            ws.column_dimensions[c].hidden = True
        for c in B['bal']:
            ws.column_dimensions[c].width = 11
        for c in (B['pool'], B['int'], B['tot']):
            ws.column_dimensions[c].width = 12

    ordered_block(SNOW)
    ordered_block(AVAL)
    for t in range(1, HORIZON + 1):
        r = S0 + t
        p = r - 1
        for col in MIN_BAL:
            ws[f'{col}{r}'] = f'=IF({col}{p}<0.005,0,MAX(0,{col}{p}*(1+{col}$6/12)-{col}$7))'
            ws[f'{col}{r}'].number_format = NUM
            ws[f'{col}{r}'].font = font(9)
        ws[f'{MIN_INT}{r}'] = f'=SUMPRODUCT(C{p}:L{p},$C$6:$L$6)/12'
        ws[f'{MIN_TOT}{r}'] = f'=SUM(C{r}:L{r})'
        for c in (MIN_INT, MIN_TOT):
            ws[f'{c}{r}'].number_format = NUM
            ws[f'{c}{r}'].font = font(9, c == MIN_TOT)
    put(ws, f'{MIN_INT}7', 'Total interest', f=font(9, True, MUTED), al=RIGHT)
    put(ws, f'{MIN_INT}8', f'=SUM({MIN_INT}{S0 + 1}:{MIN_INT}{S1})', f=font(9, True), fmt=NUM)
    for B in (SNOW, AVAL):
        put(ws, f'{B["int"]}7', 'Total interest', f=font(9, True, MUTED), al=RIGHT)
        put(ws, f'{B["int"]}8', f'=SUM({B["int"]}{S0 + 1}:{B["int"]}{S1})', f=font(9, True), fmt=NUM)
    ws.freeze_panes = f'C{S0}'
    print_setup(ws, title_rows='9:9')
    protect(ws)
    return ws


def sheet_debts(wb, debts, opts):
    ws = wb.create_sheet('Debt Payoff')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 4, 'B': 30, 'C': 14, 'D': 14, 'E': 14, 'F': 13, 'G': 11, 'H': 12, 'I': 11, 'J': 11, 'K': 13, 'L': 30, 'M': 8})
    title(ws, 'Debt Payoff', 'List your debts. Pay the minimum on every debt and put anything extra on the first one in the '
          'order. When a debt is paid off, its payment rolls over to the next.', col='B')
    put(ws, 'B4', 'Payoff method', f=font(10, True), al=LEFT_C, border=B_ROW)
    inp(ws, 'C4', (opts or {}).get('method', 'Avalanche'), al=Alignment(horizontal='center', vertical='center'))
    ws['C4'].font = font(10, True)
    ws.merge_cells('D4:L4')
    put(ws, 'D4', 'Avalanche: highest APR first (usually the least interest). Snowball: smallest balance first (quick wins).',
        f=font(9, color=MUTED), al=LEFT_C)
    put(ws, 'B5', f'="Extra payment each month"&{CUR}', f=font(10, True), al=LEFT_C, border=B_ROW)
    inp(ws, 'C5', (opts or {}).get('extra', 0), fmt=NUM, al=RIGHT_C)
    ws.merge_cells('D5:L5')
    put(ws, 'D5', 'On top of all the minimum payments. The total you pay each month stays the same until you are debt-free.',
        f=font(9, color=MUTED), al=LEFT_C)
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20
    dv_list(ws, L_METHODS, 'C4', 'Method', 'Choose Avalanche or Snowball.')
    dv_num(ws, 'C5', '0', '100000000', 'Enter a positive amount.')

    header(ws, 7, [('B', 'Debt'), ('C', f'="Balance"&{CUR}'), ('D', 'APR (yearly rate)'), ('E', f'="Minimum payment"&{CUR}'),
                   ('F', f'="Interest this month"&{CUR}'), ('G', 'Months (minimums only)'), ('H', 'Paid off (minimums only)'),
                   ('I', 'Avalanche order'), ('J', 'Snowball order'), ('K', 'Paid off (your plan)'), ('L', 'Check'),
                   ('M', 'Months (your plan)')], height=44)
    valid = f'$B${DEBT0}:$B${DEBT1},"<>",$C${DEBT0}:$C${DEBT1},">0",$D${DEBT0}:$D${DEBT1},">=0"'
    for i in range(N_DEBT):
        r = DEBT0 + i
        d = debts[i] if debts and i < len(debts) else (None,) * 4
        inp(ws, f'B{r}', d[0])
        inp(ws, f'C{r}', d[1], fmt=NUM, al=RIGHT)
        inp(ws, f'D{r}', d[2], fmt=PCT_IN, al=RIGHT)
        inp(ws, f'E{r}', d[3], fmt=NUM, al=RIGHT)
        guard = f'OR(B{r}="",N(C{r})<=0,NOT(ISNUMBER(D{r})))'
        put(ws, f'F{r}', f'=IF({guard},"",C{r}*D{r}/12)', fmt=NUMZ, border=B_ROW)
        put(ws, f'G{r}', f'=IF(I{r}="","",{DS}!{MIN_BAL[i]}$8)', fmt='0', al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'H{r}', f'=IF(G{r}="","",IF(ISNUMBER(G{r}),DATE(YEAR({ASOF}),MONTH({ASOF})+G{r},1),G{r}))', fmt=MONTH_F,
            al=Alignment(horizontal='center'), border=B_ROW)
        prev = f'$B${DEBT0 - 1}:B{r - 1},"<>",$C${DEBT0 - 1}:C{r - 1},C{r},$D${DEBT0 - 1}:D{r - 1},D{r}'
        put(ws, f'I{r}', f'=IF({guard},"",1+COUNTIFS({valid},$D${DEBT0}:$D${DEBT1},">"&D{r})'
                         f'+COUNTIFS({valid},$D${DEBT0}:$D${DEBT1},D{r},$C${DEBT0}:$C${DEBT1},"<"&C{r})+COUNTIFS({prev}))',
            f=font(10, True, TEAL), al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'J{r}', f'=IF({guard},"",1+COUNTIFS({valid},$C${DEBT0}:$C${DEBT1},"<"&C{r})'
                         f'+COUNTIFS({valid},$C${DEBT0}:$C${DEBT1},C{r},$D${DEBT0}:$D${DEBT1},">"&D{r})+COUNTIFS({prev}))',
            f=font(10, True, TEAL), al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'M{r}', f'=IF(OR(I{r}="",J{r}=""),"",IF($C$4="Snowball",INDEX({DS}!${SNOW["bal"][0]}$8:${SNOW["bal"][-1]}$8,1,J{r}),'
                         f'INDEX({DS}!${AVAL["bal"][0]}$8:${AVAL["bal"][-1]}$8,1,I{r})))', f=font(9, color=MUTED), border=B_ROW)
        put(ws, f'K{r}', f'=IF(M{r}="","",IF(ISNUMBER(M{r}),DATE(YEAR({ASOF}),MONTH({ASOF})+M{r},1),M{r}))', f=font(10, True),
            fmt=MONTH_F, al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'L{r}', f'=IF(B{r}="",IF(OR(C{r}<>"",D{r}<>"",E{r}<>""),"Enter a name",""),IF(N(C{r})<=0,"Enter the balance",'
                         f'IF(NOT(ISNUMBER(D{r})),"Enter the APR (0% if none)",IF(N(E{r})<=0,"Enter the minimum payment",'
                         f'IF(E{r}<=C{r}*D{r}/12,"Minimum does not cover interest","")))))', f=font(9, True, WARN_T), border=B_ROW)
    t = DEBT1 + 1
    put(ws, f'B{t}', 'Total', f=font(10, True), bg=SURF2, border=B_TOT)
    tot = {
        'C': (f'=SUM(C{DEBT0}:C{DEBT1})', NUMZ), 'D': (f'=IF(C{t}=0,"",SUMPRODUCT(C{DEBT0}:C{DEBT1},D{DEBT0}:D{DEBT1})/C{t})', PCT_IN),
        'E': (f'=SUM(E{DEBT0}:E{DEBT1})', NUMZ), 'F': (f'=SUM(F{DEBT0}:F{DEBT1})', NUMZ),
        'G': (f'=IF(C{t}=0,"",C24)', '0'), 'H': (f'=IF(C{t}=0,"",C25)', MONTH_F), 'I': (None, None), 'J': (None, None),
        'K': (f'=IF(C{t}=0,"",IF($C$4="Snowball",D25,E25))', MONTH_F), 'L': (None, None), 'M': (None, None),
    }
    for col, (f, fmt) in tot.items():
        put(ws, f'{col}{t}', f, f=font(10, True), bg=SURF2, fmt=fmt, border=B_TOT,
            al=Alignment(horizontal='center') if col in 'GHK' else None)
    put(ws, f'B{t + 1}', 'APR = the yearly interest rate on your statement (type 0% for an interest-free loan). '
        'The APR in the total row is the average weighted by balance.', f=font(9, color=MUTED))
    dv_num(ws, f'C{DEBT0}:C{DEBT1}', '0', '100000000', 'Enter a positive amount.')
    dv_num(ws, f'D{DEBT0}:D{DEBT1}', '0', '1', 'Type the rate with the % sign, for example 22.99%.')
    dv_num(ws, f'E{DEBT0}:E{DEBT1}', '0', '100000000', 'Enter a positive amount.')
    ws.conditional_formatting.add(f'L{DEBT0}:L{DEBT1}', FormulaRule(formula=[f'L{DEBT0}<>""'], fill=fill(WARN)))
    ws.column_dimensions['M'].hidden = True

    # compare
    section(ws, 'B21', 'Compare the methods', span_to=5)
    header(ws, 22, [('B', ''), ('C', 'Minimums only'), ('D', 'Snowball'), ('E', 'Avalanche')], height=22)
    blocks = {'C': (MIN_TOT, MIN_INT), 'D': (SNOW['tot'], SNOW['int']), 'E': (AVAL['tot'], AVAL['int'])}
    rows = [(23, f'="Paid each month"&{CUR}'), (24, 'Months until debt-free'), (25, 'Debt-free date'),
            (26, f'="Total interest"&{CUR}'), (27, f'="Interest saved vs minimums"&{CUR}')]
    for r, label in rows:
        put(ws, f'B{r}', label, border=B_ROW)
    for col, (totc, intc) in blocks.items():
        put(ws, f'{col}23', f'=$E${t}' if col == 'C' else f'={DS}!${SNOW["pool"]}$7' if col == 'D' else f'={DS}!${AVAL["pool"]}$7',
            fmt=NUMZ, border=B_ROW)
        put(ws, f'{col}24', f'={DS}!{totc}$8', fmt='0;-0;"–"', border=B_ROW, al=RIGHT)
        put(ws, f'{col}25', f'=IF($C${t}=0,"–",IF(ISNUMBER({col}24),DATE(YEAR({ASOF}),MONTH({ASOF})+{col}24,1),{col}24))',
            fmt=MONTH_F, border=B_ROW, al=RIGHT)
        put(ws, f'{col}26', f'=IF(ISNUMBER({col}24),{DS}!{intc}$8,"–")', fmt=NUMZ, border=B_ROW, al=RIGHT)
        put(ws, f'{col}27', '="–"' if col == 'C' else f'=IF(AND(ISNUMBER($C$26),ISNUMBER({col}26)),$C$26-{col}26,"–")',
            fmt=NUMZ, border=B_ROW, al=RIGHT)
    for col, meth in (('D', 'Snowball'), ('E', 'Avalanche')):
        ws.conditional_formatting.add(f'{col}23:{col}27', FormulaRule(formula=[f'$C$4="{meth}"'], fill=fill(TEAL_S),
                                                                      font=Font(name=ARIAL, bold=True)))
        ws.conditional_formatting.add(f'{col}22', FormulaRule(formula=[f'$C$4="{meth}"'], fill=fill(TEAL_D),
                                                              font=Font(name=ARIAL, bold=True, color='FFFFFF', underline='single')))
    put(ws, 'B28', f'=IF($C${t}=0,"Your plan: enter your debts above.","Your plan: "&$C$4&IF(ISNUMBER(M28),", saving "&FIXED(M28,2)&" in interest compared with minimums only.",""))',
        f=font(10, True, TEAL_D))
    put(ws, 'M28', '=IF($C$4="Snowball",D27,E27)', f=font(9, color=MUTED), fmt=NUMZ)
    put(ws, 'B29', 'If the minimum payments do not cover the interest, a debt is never paid off with minimums only ("Over 30 years").',
        f=font(9, color=MUTED))

    # payoff order
    def order_list(r0, label, order_col, sched):
        section(ws, f'B{r0}', label, span_to=6)
        header(ws, r0 + 1, [('B', 'Debt'), ('C', f'="Balance"&{CUR}'), ('D', 'APR'), ('E', f'="Minimum"&{CUR}'), ('F', 'Paid off')],
               height=22)
        for k in range(1, N_DEBT + 1):
            r = r0 + 1 + k
            put(ws, f'A{r}', k, f=font(9, True, MUTED), al=RIGHT)
            m = f'MATCH({k},${order_col}${DEBT0}:${order_col}${DEBT1},0)'
            put(ws, f'B{r}', f'=IFERROR(INDEX($B${DEBT0}:$B${DEBT1},{m}),"")', border=B_ROW)
            put(ws, f'C{r}', f'=IFERROR(INDEX($C${DEBT0}:$C${DEBT1},{m}),"")', fmt=NUMZ, border=B_ROW)
            put(ws, f'D{r}', f'=IFERROR(INDEX($D${DEBT0}:$D${DEBT1},{m}),"")', fmt=PCT_IN, border=B_ROW)
            put(ws, f'E{r}', f'=IFERROR(INDEX($E${DEBT0}:$E${DEBT1},{m}),"")', fmt=NUMZ, border=B_ROW)
            mo = f'INDEX({DS}!${sched[0]}$8:${sched[-1]}$8,1,{k})'
            put(ws, f'F{r}', f'=IF(B{r}="","",IF(ISNUMBER({mo}),DATE(YEAR({ASOF}),MONTH({ASOF})+{mo},1),{mo}))', fmt=MONTH_F,
                border=B_ROW, al=Alignment(horizontal='center'))
        return r0 + 1 + N_DEBT

    e1 = order_list(31, 'Avalanche order: highest APR first', 'I', AVAL['bal'])
    e2 = order_list(e1 + 2, 'Snowball order: smallest balance first', 'J', SNOW['bal'])
    for r0, meth in ((31, 'Avalanche'), (e1 + 2, 'Snowball')):
        ws.conditional_formatting.add(f'B{r0 + 2}:F{r0 + 2}', FormulaRule(formula=[f'AND($C$4="{meth}",$B${r0 + 2}<>"")'],
                                                                      fill=fill(TEAL_S), font=Font(name=ARIAL, bold=True)))
    # chart data
    c0 = e2 + 3
    section(ws, f'B{c0}', 'Chart data: total balance over time', span_to=6)
    header(ws, c0 + 1, [('B', 'Month'), ('C', 'Minimums only'), ('D', 'Snowball'), ('E', 'Avalanche'), ('F', 'Month #')], height=22)
    put(ws, f'H{c0 + 1}', 'Months per point', f=font(9, color=MUTED), al=RIGHT)
    step = f'$I${c0 + 1}'
    put(ws, f'I{c0 + 1}', '=MAX(1,ROUNDUP(MAX(IF(ISNUMBER(D24),D24,360),IF(ISNUMBER(E24),E24,360))/24,0))', f=font(9, True))
    for i in range(25):
        r = c0 + 2 + i
        put(ws, f'F{r}', f'=MIN({HORIZON},{i}*{step})', f=font(9, color=MUTED), fmt='0', border=B_ROW)
        d = f'DATE(YEAR({ASOF}),MONTH({ASOF})+F{r},1)'
        put(ws, f'B{r}', f'={ym(d)}', border=B_ROW)
        for col, totc in (('C', MIN_TOT), ('D', SNOW['tot']), ('E', AVAL['tot'])):
            put(ws, f'{col}{r}', f'=INDEX({DS}!${totc}${S0}:${totc}${S1},F{r}+1)', fmt=NUMZ, border=B_ROW)
    last = c0 + 26
    ch = LineChart()
    ch.title = chart_title('Total debt balance over time')
    ch.add_data(Reference(ws, min_col=3, max_col=5, min_row=c0 + 1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=c0 + 2, max_row=last))
    for s, color, w, dash in zip(ch.series, (CH_GREY, CH_EXP, TEAL), (19050, 22225, 28575), ('dash', None, None)):
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = w
        if dash:
            s.graphicalProperties.line.dashStyle = dash
        s.smooth = False
        s.marker.symbol = 'none'
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.x_axis.tickLblSkip = 4
    light_grid(ch.y_axis)
    ch.legend.position = 'b'
    ch.width, ch.height = 15.5, 9.0
    ws.add_chart(ch, 'H31')
    ws.freeze_panes = f'C{DEBT0}'
    print_setup(ws, area=f'A1:L{e2}')
    ws.row_breaks.append(Break(id=29))
    protect(ws)
    return ws


# ---------------------------------------------------------------- assemblaggio
def build(example):
    wb = Workbook()
    base = Font(name=ARIAL, size=10)
    wb._fonts = IndexedList([base])
    wb._named_styles['Normal'].font = base
    wb.remove(wb.active)
    cats = CATEGORIES
    sheet_start(wb)
    sheet_settings(wb, EXAMPLE_SETTINGS if example else None)
    sheet_categories(wb, cats)
    sheet_transactions(wb, example_transactions() if example else [])
    sheet_budget(wb, cats, EXAMPLE_BUDGET if example else None, EXAMPLE_FROM_FUNDS if example else None)
    sheet_summary(wb)
    dash = sheet_dashboard(wb)
    sheet_goals(wb, EXAMPLE_GOALS if example else None)
    sheet_funds(wb, EXAMPLE_FUNDS if example else None)
    sheet_debts(wb, EXAMPLE_DEBTS if example else None, EXAMPLE_DEBT_OPTIONS if example else None)
    sheet_schedule(wb)
    sheet_lists(wb)
    for name in ('Start Here', 'Annual Dashboard', 'Monthly Summary'):
        wb[name].sheet_properties.tabColor = TEAL
    for name in ('Settings', 'Categories', 'Transactions', 'Monthly Budget', 'Savings Goals', 'Sinking Funds', 'Debt Payoff'):
        wb[name].sheet_properties.tabColor = 'E3C766'
    if example:
        wb.active = wb.sheetnames.index('Annual Dashboard')
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = ws is dash
    wb.calculation.fullCalcOnLoad = True
    p = wb.properties
    p.creator = p.lastModifiedBy = 'NettoChiaro'
    p.title = f'{YEAR} Budget Planner' + (' (example)' if example else '')
    p.subject = 'Budget planner: transactions, monthly budget, summary, dashboard, savings goals, sinking funds, debt payoff'
    p.created = p.modified = FIXED
    return wb


def save_deterministic(wb, path):
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)', rb'\g<1>2026-01-01T00:00:00Z\g<2>', data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    with open(path, 'wb') as fh:
        fh.write(out.getvalue())


def main():
    os.makedirs(OUT, exist_ok=True)
    for example, name in ((False, f'{BASE}.xlsx'), (True, f'{BASE}-example.xlsx')):
        path = os.path.join(OUT, name)
        save_deterministic(build(example), path)
        print(f'written {os.path.relpath(path, HERE)} ({os.path.getsize(path) // 1024} KB)')


if __name__ == '__main__':
    main()
