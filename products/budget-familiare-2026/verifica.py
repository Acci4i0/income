#!/usr/bin/env python3
"""Verifica delle formule di Budget familiare 2026.

Uso:  python3 build.py && python3 verifica.py

1. Ricalcola i file di files/ con LibreOffice headless (soffice --convert-to xlsx in una cartella temporanea).
2. Cerca errori (#REF!, #NAME?, #VALUE!, #DIV/0!, #N/A) e funzioni non compatibili.
3. Confronta i valori del file di esempio con totali calcolati qui in Python, a partire dagli stessi
   dati di esempio (somme per categoria e mese, budget, dashboard, obiettivi, debiti).
4. Prova i controlli del registro su un file con righe sbagliate apposta.
5. Scrive images/src/dati.js con i numeri dell'esempio, usati dalle immagini dell'inserzione.
Esce con codice 1 se qualcosa non torna.
"""
import datetime as dt
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

sys.dont_write_bytecode = True  # niente __pycache__ nella cartella del prodotto
import build as B  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ERR_RE = re.compile(r'^(#REF!|#NAME\?|#VALUE!|#DIV/0!|#N/A|#NUM!|#NULL!|Err:\d+)')
VIETATE = re.compile(r'\b(LET|LAMBDA|XLOOKUP|XMATCH|FILTER|SORT|SORTBY|UNIQUE|SEQUENCE|TEXTJOIN|IFS|SWITCH|MAXIFS|MINIFS|CONCAT|INDIRECT|OFFSET)\(')
problems = []


def check(cond, msg):
    if not cond:
        problems.append(msg)


def close(a, b, tol=0.006):
    return a is not None and b is not None and abs(float(a) - float(b)) <= tol


def recalc(paths, outdir):
    profile = tempfile.mkdtemp(prefix='lo-profile-')
    try:
        subprocess.run(['soffice', f'-env:UserInstallation=file://{profile}', '--headless', '--convert-to', 'xlsx',
                        '--outdir', outdir, *paths], check=True, capture_output=True, timeout=180)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    return [os.path.join(outdir, os.path.basename(p)) for p in paths]


def scan_errors(path, label):
    wb = load_workbook(path, data_only=True)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and ERR_RE.match(c.value):
                    n += 1
                    if n <= 20:
                        problems.append(f'{label}: errore {c.value} in {ws.title}!{c.coordinate}')
    return wb


def scan_formulas(path, label):
    wb = load_workbook(path)
    count = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    count += 1
                    m = VIETATE.search(c.value.upper())
                    if m:
                        problems.append(f'{label}: funzione non ammessa {m.group(1)} in {ws.title}!{c.coordinate}')
    return count


def nper(rate, pmt, pv):
    if rate == 0:
        return pv / pmt
    return -math.log(1 - rate * pv / pmt) / math.log(1 + rate)


def months_between(ref, d):
    return max(0, (d.year - ref.year) * 12 + d.month - ref.month)


def verifica_esempio(wb):
    cats = B.esempio_categorie()
    tipo = {c[0]: c[1] for c in cats}
    mov = B.esempio_movimenti()
    # consuntivo atteso per categoria e mese
    act = {c[0]: [0.0] * 12 for c in cats}
    for d, _, cat, imp, _, _ in mov:
        act[cat][d.month - 1] += imp
    ms = wb['Mensile']
    for i, (name, _, _) in enumerate(cats):
        r = B.G0 + i
        check(ms[f'A{r}'].value == name, f'Mensile A{r}: atteso {name}, trovato {ms[f"A{r}"].value}')
        for m in range(12):
            v = ms.cell(r, 3 + m).value
            check(close(v, act[name][m]), f'Mensile {name} mese {m + 1}: atteso {act[name][m]:.2f}, trovato {v}')
    by_type = {t: [sum(act[c][m] for c in act if tipo[c] == t) for m in range(12)] for t in B.TIPI}
    entr, fis, var, ris = (by_type[t] for t in B.TIPI)
    spese = [fis[m] + var[m] for m in range(12)]
    avanzo = [entr[m] - spese[m] - ris[m] for m in range(12)]
    saldo, s = [], B.ESEMPIO_IMPOSTAZIONI['saldo']
    for m in range(12):
        s += avanzo[m]
        saldo.append(s)
    for r, serie in [(6, entr), (7, fis), (8, var), (9, spese), (10, ris), (11, avanzo), (12, saldo)]:
        for m in range(12):
            check(close(ms.cell(r, 3 + m).value, serie[m]), f'Mensile riga {r} mese {m + 1}: atteso {serie[m]:.2f}, trovato {ms.cell(r, 3 + m).value}')
    tot = {k: sum(v) for k, v in [('E', entr), ('S', spese), ('R', ris), ('F', fis), ('V', var)]}
    check(close(ms['O6'].value, tot['E']), 'Mensile O6 entrate anno')
    check(close(ms['O12'].value, saldo[-1]), 'Mensile O12 saldo fine anno')

    # budget atteso
    bud = B.esempio_budget()
    exp_b = {}
    for name, _, _ in cats:
        typical, ov = bud.get(name, (None, {}))
        exp_b[name] = [ov.get(m + 1, typical or 0) for m in range(12)]
    bs = wb['Budget']
    for i, (name, _, _) in enumerate(cats):
        r = B.G0 + i
        check(close(bs[f'P{r}'].value, sum(exp_b[name])), f'Budget {name} totale: atteso {sum(exp_b[name])}, trovato {bs[f"P{r}"].value}')
        check(close(ms[f'P{r}'].value, sum(exp_b[name])), f'Mensile budget {name}')
        check(close(ms[f'Q{r}'].value, sum(act[name]) - sum(exp_b[name])), f'Mensile differenza {name}')
    b_type = {t: sum(sum(exp_b[c]) for c in exp_b if tipo[c] == t) for t in B.TIPI}
    check(close(bs['P6'].value, b_type['Entrata']), 'Budget P6 entrate')
    check(close(bs['P9'].value, b_type['Spesa fissa'] + b_type['Spesa variabile']), 'Budget P9 spese')
    b_sal = B.ESEMPIO_IMPOSTAZIONI['saldo'] + b_type['Entrata'] - b_type['Spesa fissa'] - b_type['Spesa variabile'] - b_type['Risparmio e investimenti']
    check(close(bs['P12'].value, b_sal), f'Budget saldo previsto: atteso {b_sal}, trovato {bs["P12"].value}')
    # mese di riferimento = dicembre (data di riferimento 31/12) -> da inizio anno = anno intero
    check(wb['Impostazioni']['C12'].value == 12, 'Impostazioni C12 mese di riferimento')
    check(close(ms['X9'].value, tot['S']) and close(ms['W9'].value, b_type['Spesa fissa'] + b_type['Spesa variabile']), 'Mensile da inizio anno')
    check(close(ms['U6'].value, entr[11]), 'Mensile consuntivo dicembre')

    # dashboard
    ds = wb['Dashboard']
    check(close(ds['B5'].value, tot['E']), 'Dashboard entrate')
    check(close(ds['D5'].value, tot['S']), 'Dashboard spese')
    check(close(ds['F5'].value, tot['E'] - tot['S']), 'Dashboard risparmio')
    check(close(ds['H5'].value, (tot['E'] - tot['S']) / tot['E'], 1e-9), 'Dashboard tasso di risparmio')
    check(close(ds['J5'].value, saldo[-1]), 'Dashboard saldo di cassa')
    debiti_tot = sum(d[1] for d in B.ESEMPIO_DEBITI)
    check(close(ds['L5'].value, debiti_tot), 'Dashboard debiti residui')
    for m in range(12):
        check(close(ds[f'C{10 + m}'].value, entr[m]) and close(ds[f'D{10 + m}'].value, spese[m]), f'Dashboard mese {m + 1}')
    spese_cat = sorted(((sum(v), c) for c, v in act.items() if tipo[c].startswith('Spesa') and sum(v) > 0),
                       key=lambda x: (-x[0], [n for n, _, _ in cats].index(x[1])))
    for k in range(10):
        exp_tot, exp_name = spese_cat[k]
        check(ds[f'C{36 + k}'].value == exp_name and close(ds[f'F{36 + k}'].value, exp_tot),
              f'Dashboard top {k + 1}: atteso {exp_name} {exp_tot:.2f}, trovato {ds[f"C{36 + k}"].value} {ds[f"F{36 + k}"].value}')
    per = {}
    for d, _, cat, imp, _, chi in mov:
        if tipo[cat].startswith('Spesa'):
            per[chi] = per.get(chi, 0) + imp
    for i, nome in enumerate(B.ESEMPIO_IMPOSTAZIONI['membri']):
        check(close(ds[f'D{50 + i}'].value, per.get(nome, 0)), f'Dashboard spese di {nome}')
    check(close(ds['D56'].value, per.get('', 0)), 'Dashboard spese comuni')
    check(ds['M25'].value == len(mov), f'Dashboard movimenti registrati: atteso {len(mov)}, trovato {ds["M25"].value}')
    check(ds['M26'].value == 0 and ds['M27'].value == 0 and ds['M28'].value == 0, 'Dashboard controlli non a zero')

    # obiettivi
    ref = B.ESEMPIO_IMPOSTAZIONI['data_rif']
    ob = wb['Obiettivi']
    tot_q = 0
    for i, (nome, imp, scad, vers) in enumerate(B.ESEMPIO_OBIETTIVI):
        r = B.OBI0 + i
        manc = max(0, imp - vers)
        mesi = months_between(ref, scad)
        quota = 0 if manc == 0 else manc / max(1, mesi)
        tot_q += quota
        stato = 'Raggiunto' if manc == 0 else ('Scaduto' if scad < ref else 'In corso')
        check(close(ob[f'F{r}'].value, manc), f'Obiettivi {nome} mancante')
        check(ob[f'I{r}'].value == mesi, f'Obiettivi {nome} mesi: atteso {mesi}, trovato {ob[f"I{r}"].value}')
        check(close(ob[f'J{r}'].value, quota), f'Obiettivi {nome} quota: attesa {quota:.2f}, trovata {ob[f"J{r}"].value}')
        check(ob[f'K{r}'].value == stato, f'Obiettivi {nome} stato')
    t = B.OBI0 + B.N_OBI
    check(close(ob[f'J{t}'].value, tot_q), 'Obiettivi quota totale')

    # debiti
    de = wb['Debiti']
    valid = [d for d in B.ESEMPIO_DEBITI]
    val_order = sorted(range(len(valid)), key=lambda i: (-valid[i][2], valid[i][1], i))
    snow_order = sorted(range(len(valid)), key=lambda i: (valid[i][1], -valid[i][2], i))
    for i, (nome, saldo_d, tan, rata) in enumerate(valid):
        r = B.DEB0 + i
        n = nper(tan / 12, rata, saldo_d)
        check(de[f'G{r}'].value == math.ceil(n - 1e-9), f'Debiti {nome} mesi: atteso {math.ceil(n)}, trovato {de[f"G{r}"].value}')
        interessi = 0 if tan == 0 else max(0, rata * n - saldo_d)
        check(close(de[f'I{r}'].value, interessi, 0.02), f'Debiti {nome} interessi: attesi {interessi:.2f}, trovati {de[f"I{r}"].value}')
        check(de[f'J{r}'].value == val_order.index(i) + 1, f'Debiti {nome} ordine valanga')
        check(de[f'K{r}'].value == snow_order.index(i) + 1, f'Debiti {nome} ordine palla di neve')
    for k, i in enumerate(val_order):
        check(de[f'B{23 + k}'].value == valid[i][0], f'Debiti piano valanga posizione {k + 1}')
    for k, i in enumerate(snow_order):
        check(de[f'F{23 + k}'].value == valid[i][0], f'Debiti piano palla di neve posizione {k + 1}')
    first = valid[val_order[0]]
    extra = B.ESEMPIO_DEBITI_OPZIONI['extra']
    n0, n1 = nper(first[2] / 12, first[3], first[1]), nper(first[2] / 12, first[3] + extra, first[1])
    check(de['C35'].value == first[0], 'Debiti: debito su cui concentrare l\'extra')
    check(de['C40'].value == math.ceil(n1), f'Debiti mesi con extra: attesi {math.ceil(n1)}, trovati {de["C40"].value}')
    check(de['C41'].value == math.ceil(n0) - math.ceil(n1), 'Debiti mesi in meno')
    saved = (first[3] * n0 - first[1]) - ((first[3] + extra) * n1 - first[1])
    check(close(de['C42'].value, saved, 0.02), f'Debiti interessi risparmiati: attesi {saved:.2f}, trovati {de["C42"].value}')
    return dict(cats=cats, act=act, entr=entr, spese=spese, ris=ris, avanzo=avanzo, saldo=saldo, fis=fis, var=var,
                tot=tot, exp_b=exp_b, spese_cat=spese_cat, mov=mov)


def verifica_vuoto(wb):
    ds = wb['Dashboard']
    for ref in ('B5', 'D5', 'F5', 'H5', 'J5', 'L5', 'M25', 'M26', 'M27', 'M28'):
        check(ds[ref].value in (0, None), f'Vuoto: Dashboard {ref} = {ds[ref].value}')
    check(ds['C36'].value in ('', None), 'Vuoto: top categorie non vuota')
    check(wb['Mensile']['O9'].value == 0, 'Vuoto: Mensile spese non a zero')
    check(wb['Impostazioni']['C11'].value is not None, 'Vuoto: data di riferimento (oggi) mancante')


def verifica_controlli(tmp):
    """File di prova con righe sbagliate apposta: il registro deve segnalarle e non conteggiarle."""
    wb = B.build(False)
    ws = wb['Movimenti']
    righe = [
        (dt.date(2026, 3, 10), 'ok spesa', 'Luce', 50, 'Valida'),
        (dt.date(2026, 3, 12), 'rimborso', 'Luce', -10, 'Valida'),
        (dt.date(2025, 12, 31), 'anno prima', 'Luce', 99, 'Fuori dall\'anno'),
        (None, 'senza data', 'Luce', 20, 'Data mancante o non valida'),
        (dt.date(2026, 3, 1), 'senza categoria', None, 30, 'Manca la categoria'),
        (dt.date(2026, 3, 1), 'senza importo', 'Luce', None, 'Manca l\'importo'),
        (dt.date(2026, 4, 1), 'stipendio', 'Stipendio 1', 1000, 'Valida'),
        (dt.date(2026, 4, 2), 'versamento', 'Fondo emergenze', 200, 'Valida'),
    ]
    for i, (d, desc, cat, imp, _) in enumerate(righe):
        r = B.MOV0 + i
        ws[f'A{r}'], ws[f'B{r}'], ws[f'C{r}'], ws[f'D{r}'] = d, desc, cat, imp
    # categoria non in elenco (scritta a mano, aggirando il menu)
    ws[f'A{B.MOV0 + 8}'], ws[f'C{B.MOV0 + 8}'], ws[f'D{B.MOV0 + 8}'] = dt.date(2026, 5, 1), 'Categoria inventata', 15
    righe.append((None, None, None, None, 'Categoria non trovata'))
    # debito con rata che non copre gli interessi, debito senza nome, debito senza rata, debito estinto (saldo 0)
    de = wb['Debiti']
    de['B8'], de['C8'], de['D8'], de['E8'] = 'Prova', 10000, 0.12, 50
    de['C9'], de['D9'], de['E9'] = 500, 0.05, 20
    de['B10'], de['C10'], de['D10'] = 'Senza rata', 1000, 0.05
    de['B11'], de['C11'], de['D11'], de['E11'] = 'Estinto', 0, 0.05, 20
    # categoria doppia e categoria con un carattere jolly (* sommerebbe anche "Svago")
    wb['Categorie'][f'A{B.CAT0 + 40}'], wb['Categorie'][f'B{B.CAT0 + 40}'] = 'Luce', 'Spesa fissa'
    wb['Categorie'][f'A{B.CAT0 + 41}'], wb['Categorie'][f'B{B.CAT0 + 41}'] = 'Svago*', 'Spesa variabile'
    # mese di riferimento vuoto = automatico; nome ripetuto tra i membri della famiglia
    imp = wb['Impostazioni']
    imp['C7'] = None
    imp[f'C{B.MEM0}'], imp[f'C{B.MEM0 + 1}'] = 'Anna', 'Anna'
    path = os.path.join(tmp, 'prova-controlli.xlsx')
    wb.save(path)
    out = recalc([path], os.path.join(tmp, 'out'))[0]
    v = load_workbook(out, data_only=True)
    mv = v['Movimenti']
    for i, (_, desc, _, _, atteso) in enumerate(righe):
        got = mv[f'H{B.MOV0 + i}'].value or 'Valida'
        check(got == atteso, f'Controlli riga {i + 1}: atteso «{atteso}», trovato «{got}»')
    ms = v['Mensile']
    # Luce e' la 9a categoria: riga G0+8. Doppione: la riga 47 (G0+40) viene sommata anche lei, ma e' segnalata.
    check(close(ms[f'E{B.G0 + 8}'].value, 40), f'Controlli: Luce marzo attesa 40 (50 - 10), trovata {ms[f"E{B.G0 + 8}"].value}')
    check(close(ms['F6'].value, 1000) and close(ms['F10'].value, 200), 'Controlli: entrate e risparmio di aprile')
    check(v['Dashboard']['M26'].value == 5, f'Controlli: righe da controllare attese 5, trovate {v["Dashboard"]["M26"].value}')
    check(v['Dashboard']['M27'].value == 3, f'Controlli: categorie da controllare attese 3, trovate {v["Dashboard"]["M27"].value}')
    check(v['Categorie'][f'D{B.CAT0 + 41}'].value == 'Simbolo non ammesso', 'Controlli: categoria con carattere jolly non segnalata')
    de = v['Debiti']
    for ref, atteso in [('G8', 'Rata troppo bassa'), ('G9', 'Manca il nome'), ('G10', 'Manca la rata'), ('G11', None)]:
        check(de[ref].value == atteso, f'Controlli: Debiti {ref} atteso «{atteso}», trovato «{de[ref].value}»')
    check(v['Dashboard']['M28'].value == 3, f'Controlli: debiti da controllare attesi 3, trovati {v["Dashboard"]["M28"].value}')
    oggi = dt.date.today()
    mese = oggi.month if oggi.year == B.ANNO else (12 if oggi.year > B.ANNO else 1)
    check(v['Impostazioni']['C12'].value == mese, f'Controlli: mese di riferimento vuoto, atteso {mese}, trovato {v["Impostazioni"]["C12"].value}')
    ds = v['Dashboard']
    check(ds['B50'].value == 'Anna' and ds['B51'].value in ('', None), 'Controlli: nome ripetuto tra i membri contato due volte')
    scan_errors(out, 'prova-controlli')


def scrivi_dati_immagini(wb, exp):
    """Numeri dell'esempio per le immagini dell'inserzione (images/src/dati.js)."""
    ms, ds, ob, de, mv = wb['Mensile'], wb['Dashboard'], wb['Obiettivi'], wb['Debiti'], wb['Movimenti']
    cats = []
    for i in range(B.N_CAT):
        r = B.G0 + i
        if not ms[f'A{r}'].value:
            continue
        cats.append(dict(nome=ms[f'A{r}'].value, tipo=ms[f'B{r}'].value,
                         mesi=[round(ms.cell(r, 3 + m).value or 0, 2) for m in range(12)],
                         totale=round(ms[f'O{r}'].value or 0, 2), budget=round(ms[f'P{r}'].value or 0, 2),
                         diff=round(ms[f'Q{r}'].value or 0, 2)))
    summary = {}
    for r, key in [(6, 'entrate'), (7, 'fisse'), (8, 'variabili'), (9, 'spese'), (10, 'risparmio'), (11, 'avanzo'), (12, 'saldo')]:
        summary[key] = dict(mesi=[round(ms.cell(r, 3 + m).value or 0, 2) for m in range(12)],
                            totale=round(ms[f'O{r}'].value or 0, 2), budget=round(ms[f'P{r}'].value or 0, 2),
                            diff=round(ms[f'Q{r}'].value or 0, 2))
    movimenti = []
    for r in range(B.MOV0, B.MOV1 + 1):
        if mv[f'A{r}'].value is None:
            break
        d = mv[f'A{r}'].value
        movimenti.append(dict(data=d.strftime('%d/%m/%Y'), desc=mv[f'B{r}'].value, cat=mv[f'C{r}'].value,
                              imp=mv[f'D{r}'].value, conto=mv[f'E{r}'].value, chi=mv[f'F{r}'].value or '',
                              tipo=mv[f'G{r}'].value))
    obiettivi = []
    for r in range(B.OBI0, B.OBI0 + B.N_OBI):
        if not ob[f'B{r}'].value:
            continue
        obiettivi.append(dict(nome=ob[f'B{r}'].value, importo=ob[f'C{r}'].value, scadenza=ob[f'D{r}'].value.strftime('%d/%m/%Y'),
                              versato=ob[f'E{r}'].value, mancante=ob[f'F{r}'].value, perc=ob[f'G{r}'].value,
                              mesi=ob[f'I{r}'].value, quota=round(ob[f'J{r}'].value, 2), stato=ob[f'K{r}'].value))
    debiti = []
    for r in range(B.DEB0, B.DEB0 + B.N_DEB):
        if not de[f'B{r}'].value:
            continue
        debiti.append(dict(nome=de[f'B{r}'].value, saldo=de[f'C{r}'].value, tan=de[f'D{r}'].value, rata=de[f'E{r}'].value,
                           interessi_mese=round(de[f'F{r}'].value, 2), mesi=de[f'G{r}'].value,
                           fine=de[f'H{r}'].value.strftime('%m/%Y') if de[f'H{r}'].value else '',
                           interessi=round(de[f'I{r}'].value, 2), valanga=de[f'J{r}'].value, neve=de[f'K{r}'].value))
    top = [dict(nome=ds[f'C{36 + k}'].value, totale=ds[f'F{36 + k}'].value, perc=ds[f'G{36 + k}'].value) for k in range(10)]
    dati = dict(
        anno=B.ANNO, famiglia=B.ESEMPIO_IMPOSTAZIONI['famiglia'], mesi=B.MESI_BREVI,
        kpi=dict(entrate=ds['B5'].value, spese=ds['D5'].value, risparmio=ds['F5'].value, tasso=ds['H5'].value,
                 saldo=ds['J5'].value, debiti=ds['L5'].value, rate=ds['L6'].value, versati=ds['F6'].value,
                 entrate_budget=ds['B6'].value, spese_budget=ds['D6'].value, saldo_iniziale=ds['J6'].value),
        summary=summary, categorie=cats, top=top, movimenti=movimenti, obiettivi=obiettivi, debiti=debiti,
        quota_totale=round(ob[f'J{B.OBI0 + B.N_OBI}'].value, 2),
        extra=dict(debito=de['C35'].value, mesi_prima=de['C39'].value, mesi_dopo=de['C40'].value,
                   mesi_meno=de['C41'].value, interessi=round(de['C42'].value, 2), somma=B.ESEMPIO_DEBITI_OPZIONI['extra']),
        n_movimenti=len(exp['mov']),
    )
    path = os.path.join(HERE, 'images', 'src', 'dati.js')
    with open(path, 'w') as fh:
        fh.write('// Generato da verifica.py dai valori ricalcolati del file di esempio. Non modificare a mano.\n')
        fh.write('window.DATI = ' + json.dumps(dati, ensure_ascii=False, indent=1) + ';\n')
    return path


def main():
    files = [os.path.join(HERE, 'files', f'{B.SLUG}.xlsx'), os.path.join(HERE, 'files', f'{B.SLUG}-esempio.xlsx')]
    for f in files:
        check(os.path.exists(f), f'manca {f}: esegui prima python3 build.py')
    if problems:
        print('\n'.join(problems))
        sys.exit(1)
    tmp = tempfile.mkdtemp(prefix='verifica-budget-')
    try:
        n_formule = [scan_formulas(f, os.path.basename(f)) for f in files]
        vuoto, esempio = recalc(files, tmp)
        wv = scan_errors(vuoto, 'vuoto')
        we = scan_errors(esempio, 'esempio')
        verifica_vuoto(wv)
        exp = verifica_esempio(we)
        verifica_controlli(tmp)
        dati = scrivi_dati_immagini(we, exp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    t = exp['tot']
    print(f'formule: {n_formule[0]} (vuoto), {n_formule[1]} (esempio)')
    print(f'esempio: {len(exp["mov"])} movimenti · entrate {t["E"]:.2f} · spese {t["S"]:.2f} · '
          f'risparmio versato {t["R"]:.2f} · saldo fine anno {exp["saldo"][-1]:.2f}')
    print(f'dati per le immagini: {os.path.relpath(dati, HERE)}')
    if problems:
        print(f'\n{len(problems)} PROBLEMI:')
        print('\n'.join(problems[:80]))
        sys.exit(1)
    print('OK: nessun errore, tutti i valori attesi coincidono.')


if __name__ == '__main__':
    main()
