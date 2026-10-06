#!/usr/bin/env python3
"""Budget familiare 2026 (NettoChiaro): rigenera i file consegnati all'acquirente.

Uso:  python3 build.py

Scrive in files/:
  budget-familiare-2026.xlsx          modello da compilare
  budget-familiare-2026-esempio.xlsx  stesso modello con i dati di una famiglia di fantasia

Il risultato e' deterministico: stessi dati, stessi byte (le date dei metadati sono fissate).
Solo formule comuni a Excel 2016+, Google Sheets e LibreOffice (niente LET, XLOOKUP, FILTER, macro).
Verifica delle formule: python3 verifica.py (ricalcola con LibreOffice e confronta i totali).
"""
import calendar
import datetime as dt
import io
import os
import random
import re
import zipfile

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation

ANNO = 2026
VERSIONE = '2026.1'
SLUG = 'budget-familiare-2026'
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'files')
FIXED = dt.datetime(2026, 1, 1, 0, 0, 0)

# ---------------------------------------------------------------- dimensioni (riferimenti fissi)
CAT0, N_CAT = 6, 60                 # Categorie!A6:A65
CAT1 = CAT0 + N_CAT - 1
MOV0, N_MOV = 6, 2000               # Movimenti!A6:H2005
MOV1 = MOV0 + N_MOV - 1
G0 = 14                             # prima riga categorie in Budget e Mensile
G1 = G0 + N_CAT - 1                 # 73
MEM0, N_MEM = 16, 6                 # Impostazioni!C16:C21
CON0, N_CON = 24, 8                 # Impostazioni!C24:C31
OBI0, N_OBI = 5, 12                 # Obiettivi!B5:B16
DEB0, N_DEB = 8, 10                 # Debiti!B8:B17
DEB1 = DEB0 + N_DEB - 1

ANNO_REF = 'Impostazioni!$C$4'
SALDO_REF = 'Impostazioni!$C$5'
DATA_REF = 'Impostazioni!$C$11'
MESE_REF = 'Impostazioni!$C$12'
MESE_NOME_REF = 'Impostazioni!$C$13'

M_DATA = f'Movimenti!$A${MOV0}:$A${MOV1}'
M_CAT = f'Movimenti!$C${MOV0}:$C${MOV1}'
M_IMP = f'Movimenti!$D${MOV0}:$D${MOV1}'
M_CHI = f'Movimenti!$F${MOV0}:$F${MOV1}'
M_TIPO = f'Movimenti!$G${MOV0}:$G${MOV1}'
M_CTRL = f'Movimenti!$H${MOV0}:$H${MOV1}'
C_NOMI = f'Categorie!$A${CAT0}:$A${CAT1}'
C_TIPI = f'Categorie!$B${CAT0}:$B${CAT1}'

MESI = ['Gennaio', 'Febbraio', 'Marzo', 'Aprile', 'Maggio', 'Giugno', 'Luglio', 'Agosto',
        'Settembre', 'Ottobre', 'Novembre', 'Dicembre']
MESI_BREVI = ['Gen', 'Feb', 'Mar', 'Apr', 'Mag', 'Giu', 'Lug', 'Ago', 'Set', 'Ott', 'Nov', 'Dic']
TIPI = ['Entrata', 'Spesa fissa', 'Spesa variabile', 'Risparmio e investimenti']

CATEGORIE = [
    ('Stipendio 1', 'Entrata', 'Stipendio o pensione netta. Rinominala con il nome della persona.'),
    ('Stipendio 2', 'Entrata', 'Secondo stipendio o pensione.'),
    ('Tredicesima e quattordicesima', 'Entrata', 'Mensilità aggiuntive.'),
    ('Assegno unico figli', 'Entrata', 'Assegno unico e universale INPS.'),
    ('Entrate extra', 'Entrata', 'Lavori occasionali, vendita di oggetti usati.'),
    ('Rimborsi e altre entrate', 'Entrata', 'Rimborso del 730, rimborsi spese, regali in denaro.'),
    ('Affitto o rata mutuo', 'Spesa fissa', 'Per stimare la rata del mutuo: nettochiaro.com/calcolo-rata-mutuo'),
    ('Condominio', 'Spesa fissa', ''),
    ('Luce', 'Spesa fissa', ''),
    ('Gas', 'Spesa fissa', ''),
    ('Acqua', 'Spesa fissa', ''),
    ('TARI', 'Spesa fissa', 'Tassa sui rifiuti del Comune.'),
    ('Internet e telefono', 'Spesa fissa', ''),
    ('Assicurazioni', 'Spesa fissa', 'Auto, casa, salute, vita.'),
    ('Bollo auto', 'Spesa fissa', 'Tassa automobilistica regionale.'),
    ('Abbonamenti', 'Spesa fissa', 'Streaming, palestra, quotidiani, app.'),
    ('Rate di prestiti', 'Spesa fissa', 'Prestiti personali, auto, acquisti a rate. Dettaglio nel foglio Debiti.'),
    ('Spesa alimentare', 'Spesa variabile', 'Supermercato, mercato, prodotti per la casa.'),
    ('Carburante', 'Spesa variabile', ''),
    ('Trasporti', 'Spesa variabile', 'Mezzi pubblici, treni, parcheggi, pedaggi.'),
    ('Farmacia e salute', 'Spesa variabile', 'Farmaci, visite, ticket, dentista.'),
    ('Scuola e figli', 'Spesa variabile', 'Mensa, libri, gite, sport, centri estivi.'),
    ('Casa e manutenzione', 'Spesa variabile', 'Piccoli lavori, arredi, elettrodomestici.'),
    ('Abbigliamento', 'Spesa variabile', ''),
    ('Svago', 'Spesa variabile', 'Ristoranti, bar, cinema, hobby.'),
    ('Regali', 'Spesa variabile', ''),
    ('Vacanze', 'Spesa variabile', ''),
    ('Imprevisti', 'Spesa variabile', 'Riparazioni e spese non previste.'),
    ('Fondo emergenze', 'Risparmio e investimenti', 'Versamenti sul conto di riserva.'),
    ('Obiettivi di risparmio', 'Risparmio e investimenti', 'Versamenti per gli obiettivi del foglio Obiettivi.'),
    ('Investimenti', 'Risparmio e investimenti', 'Piani di accumulo, titoli, buoni.'),
    ('Previdenza complementare', 'Risparmio e investimenti', 'Versamenti volontari al fondo pensione.'),
]
CONTI = ['Conto corrente', 'Carta di debito', 'Carta di credito', 'Contanti', 'Carta prepagata']

# ---------------------------------------------------------------- stile
ARIAL = 'Arial'
TEAL, TEAL_D, TEAL_S = '0F766E', '0B5C56', 'E0F2EF'
TEXT, MUTED, LINE, SURF2 = '17201C', '5B6862', 'D9E0DC', 'EEF2EF'
INPUT, INPUT_LINE = 'FFF6D5', 'E3D49B'
WARN, WARN_T = 'FFE4C7', '9A3412'
GOOD, BAD, BAD_S = '15803D', 'B91C1C', 'FDECEC'
CH_ENTRATE, CH_SPESE = '0D9488', 'C2410C'

EUR = '#,##0.00 "€";-#,##0.00 "€";"–"'
EUR_IN = '#,##0.00 "€";-#,##0.00 "€";0.00 "€"'
EUR0 = '#,##0 "€";-#,##0 "€";"0 €"'
EUR0D = '#,##0 "€";-#,##0 "€";"–"'
PCT = '0.0%'
PCT_IN = '0.00%'
DATE = 'DD/MM/YYYY'


def font(size=10, bold=False, color=TEXT, italic=False, underline=None):
    return Font(name=ARIAL, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(color):
    return PatternFill('solid', fgColor=color)


THIN = Side(style='thin', color=LINE)
THIN_IN = Side(style='thin', color=INPUT_LINE)
B_ROW = Border(bottom=THIN)
B_IN = Border(left=THIN_IN, right=THIN_IN, top=THIN_IN, bottom=THIN_IN)
B_TOT = Border(top=Side(style='thin', color=TEAL), bottom=Side(style='thin', color=TEAL))
UNLOCK = Protection(locked=False)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_C = Alignment(horizontal='left', vertical='center')
RIGHT_C = Alignment(horizontal='right', vertical='center')


def chart_title(text):
    cp = CharacterProperties(sz=1100, b=True, solidFill=TEXT)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def put(ws, ref, value=None, f=None, bg=None, fmt=None, al=None, border=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = f or font()
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if border:
        c.border = border
    return c


def inp(ws, ref, value=None, fmt=None, al=None):
    """Cella da compilare: giallo chiaro, sbloccata."""
    c = put(ws, ref, value, bg=INPUT, fmt=fmt, al=al or LEFT_C, border=B_IN)
    c.protection = UNLOCK
    return c


def title(ws, ref, text, sub=None, sub_ref=None):
    put(ws, ref, text, f=font(18, True, TEAL))
    ws.row_dimensions[ws[ref].row].height = 30
    if sub:
        put(ws, sub_ref, sub, f=font(10, color=MUTED))


def header(ws, row, cols_texts, height=30):
    for col, text in cols_texts:
        put(ws, f'{col}{row}', text, f=font(10, True, 'FFFFFF'), bg=TEAL, al=CENTER)
    ws.row_dimensions[row].height = height


def section(ws, ref, text, span_to=None):
    put(ws, ref, text, f=font(11, True, TEAL_D))
    if span_to:
        r = ws[ref].row
        c0 = ws[ref].column
        for c in range(c0, span_to + 1):
            ws.cell(r, c).border = Border(bottom=Side(style='medium', color=TEAL))


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def print_setup(ws, title_rows=None, area=None):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if title_rows:
        ws.print_title_rows = title_rows
    if area:
        ws.print_area = area


def protect(ws, sort=False):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False   # False = consentito
    p.formatRows = False
    p.formatCells = False
    if sort:
        p.sort = False
        p.autoFilter = False


def dv_list(ws, src, ranges, title_, msg):
    dv = DataValidation(type='list', formula1=src, allow_blank=True, showErrorMessage=True,
                        errorTitle=title_, error=msg)
    ws.add_data_validation(dv)
    dv.add(ranges)


def dv_date(ws, ranges):
    dv = DataValidation(type='date', operator='between', formula1='36526', formula2='73050',
                        allow_blank=True, showErrorMessage=True, errorTitle='Data non valida',
                        error='Inserisci una data nel formato gg/mm/aaaa.')
    ws.add_data_validation(dv)
    dv.add(ranges)


def dv_num(ws, ranges, lo='-100000000', hi='100000000', msg='Inserisci un numero (es. 125,50).'):
    dv = DataValidation(type='decimal', operator='between', formula1=lo, formula2=hi,
                        allow_blank=True, showErrorMessage=True, errorTitle='Valore non valido', error=msg)
    ws.add_data_validation(dv)
    dv.add(ranges)


# ---------------------------------------------------------------- dati di esempio (famiglia di fantasia)
def last_day(m):
    return calendar.monthrange(ANNO, m)[1]


def esempio_categorie():
    cats = [list(c) for c in CATEGORIE]
    cats[0][0], cats[0][2] = 'Stipendio Marco', 'Stipendio netto.'
    cats[1][0], cats[1][2] = 'Stipendio Giulia', 'Stipendio netto.'
    return [tuple(c) for c in cats]


def esempio_movimenti():
    """Un anno di movimenti realistici per una coppia con una figlia. Deterministico (seed fisso)."""
    rng = random.Random(2026)
    mov = []

    def add(d, desc, cat, imp, conto='Conto corrente', chi=''):
        mov.append((d, desc, cat, round(imp, 2), conto, chi))

    luce = [82, 76, 70, 61, 57, 64, 84, 88, 71, 63, 69, 80]
    gas = [168, 152, 121, 74, 41, 29, 24, 23, 31, 62, 115, 160]
    acqua = {2: 91.40, 5: 87.20, 8: 104.60, 11: 95.10}
    for m in range(1, 13):
        def D(day):
            return dt.date(ANNO, m, min(day, last_day(m)))
        add(D(1), 'Rata mutuo', 'Affitto o rata mutuo', 680)
        add(D(5), 'Fibra casa', 'Internet e telefono', 27.90, 'Carta di credito')
        add(D(5), 'Telefono Marco', 'Internet e telefono', 9.99, 'Carta di credito', 'Marco')
        add(D(5), 'Telefono Giulia', 'Internet e telefono', 7.99, 'Carta di credito', 'Giulia')
        add(D(8), 'Streaming video', 'Abbonamenti', 13.99, 'Carta di credito')
        add(D(8), 'Musica in streaming', 'Abbonamenti', 10.99, 'Carta di credito', 'Giulia')
        if m in (1, 4, 7, 10):
            add(D(10), 'Rata condominio trimestrale', 'Condominio', 270)
        add(D(15), 'Rata prestito auto', 'Rate di prestiti', 210)
        add(D(15), 'Rata carta revolving', 'Rate di prestiti', 60)
        add(D(15), 'Rata lavatrice a tasso zero', 'Rate di prestiti', 40)
        add(D(18), 'Bolletta luce', 'Luce', luce[m - 1] + rng.uniform(-4, 4))
        add(D(18), 'Bolletta gas', 'Gas', gas[m - 1] + rng.uniform(-5, 5))
        if m in acqua:
            add(D(22), 'Bolletta acqua', 'Acqua', acqua[m])
        if m in (4, 7, 11):
            add(D(16), 'TARI rata', 'TARI', 108.67)
        add(D(20), 'Assegno unico INPS', 'Assegno unico figli', 145.30)
        add(D(27), 'Stipendio', 'Stipendio Marco', 1850, chi='Marco')
        add(D(27), 'Stipendio', 'Stipendio Giulia', 1480, chi='Giulia')
        add(D(28), 'Versamento fondo emergenze', 'Fondo emergenze', 150)
        add(D(28), 'Accantonamento obiettivi', 'Obiettivi di risparmio', 200)
        add(D(28), 'Piano di accumulo mensile', 'Investimenti', 100)
        if m not in (7, 8):
            add(D(3), 'Abbonamento bus Sofia', 'Trasporti', 25, 'Carta di debito', 'Sofia')
        for _ in range(rng.randint(1, 2)):
            add(D(rng.randint(2, 27)), 'Parcheggio', 'Trasporti', rng.uniform(2, 9), 'Contanti')
        if m in (1, 2, 3, 4, 5, 10, 11, 12):
            add(D(4), 'Mensa scolastica', 'Scuola e figli', 78, chi='Sofia')
        if m == 9:
            add(D(22), 'Mensa scolastica (mezzo mese)', 'Scuola e figli', 39, chi='Sofia')
        for _ in range(rng.randint(1, 3)):
            add(D(rng.randint(2, 27)), 'Farmacia', 'Farmacia e salute', rng.uniform(6, 38), 'Carta di debito')
        for _ in range(rng.randint(3, 4)):
            desc, lo, hi = rng.choice([('Pizzeria', 38, 62), ('Cinema', 22, 30), ('Gelateria', 9, 16),
                                       ('Bar', 6, 14), ('Cena con amici', 45, 80), ('Libri', 12, 28)])
            add(D(rng.randint(2, 28)), desc, 'Svago', rng.uniform(lo, hi), 'Carta di debito')
        if m == 12:
            add(D(19), 'Cena di Natale', 'Svago', 96, 'Carta di credito')

    # spesa settimanale (sabato) e piccola spesa (mercoledì)
    d = dt.date(ANNO, 1, 3)
    while d.year == ANNO:
        k = 1.25 if d.month == 12 else 1.0
        add(d, 'Spesa settimanale', 'Spesa alimentare', rng.uniform(88, 132) * k, 'Carta di debito')
        mid = d + dt.timedelta(days=4)
        if mid.year == ANNO and mid.isocalendar()[1] % 2 == 0:
            add(mid, 'Spesa al mercato', 'Spesa alimentare', rng.uniform(16, 38), 'Contanti')
        d += dt.timedelta(days=7)
    # carburante ogni 9-12 giorni
    d, chi = dt.date(ANNO, 1, 5), 'Marco'
    while d.year == ANNO:
        add(d, 'Rifornimento', 'Carburante', rng.uniform(52, 68), 'Carta di credito', chi)
        chi = 'Giulia' if chi == 'Marco' else 'Marco'
        d += dt.timedelta(days=rng.randint(9, 12) - (3 if d.month == 8 else 0))

    una_tantum = [
        ((3, 14), 'Vendita bici usata', 'Entrate extra', 120, 'Contanti', 'Marco'),
        ((6, 20), 'Lavoro extra nel fine settimana', 'Entrate extra', 300, 'Conto corrente', 'Giulia'),
        ((7, 27), 'Rimborso 730 in busta paga', 'Rimborsi e altre entrate', 412, 'Conto corrente', 'Marco'),
        ((11, 12), 'Rimborso spese mediche', 'Rimborsi e altre entrate', 64, 'Conto corrente', ''),
        ((12, 18), 'Tredicesima', 'Tredicesima e quattordicesima', 1790, 'Conto corrente', 'Marco'),
        ((12, 18), 'Tredicesima', 'Tredicesima e quattordicesima', 1430, 'Conto corrente', 'Giulia'),
        ((3, 2), 'RC auto', 'Assicurazioni', 486, 'Conto corrente', ''),
        ((6, 9), 'Assicurazione casa', 'Assicurazioni', 168, 'Conto corrente', ''),
        ((4, 28), 'Bollo auto', 'Bollo auto', 232.40, 'Carta di credito', 'Marco'),
        ((9, 7), 'Palestra Giulia (annuale)', 'Abbonamenti', 290, 'Carta di debito', 'Giulia'),
        ((1, 12), 'Corso di nuoto Sofia', 'Scuola e figli', 165, 'Carta di debito', 'Sofia'),
        ((9, 14), 'Corso di nuoto Sofia', 'Scuola e figli', 165, 'Carta di debito', 'Sofia'),
        ((9, 5), 'Libri e materiale scolastico', 'Scuola e figli', 238.40, 'Carta di debito', 'Sofia'),
        ((4, 20), 'Gita scolastica', 'Scuola e figli', 42, 'Contanti', 'Sofia'),
        ((7, 1), 'Centro estivo', 'Scuola e figli', 310, 'Conto corrente', 'Sofia'),
        ((5, 13), 'Visita dal dentista', 'Farmacia e salute', 180, 'Carta di debito', 'Giulia'),
        ((9, 24), 'Occhiali Sofia', 'Farmacia e salute', 160, 'Carta di debito', 'Sofia'),
        ((2, 7), 'Ferramenta', 'Casa e manutenzione', 23.50, 'Contanti', ''),
        ((3, 21), 'Piante per il balcone', 'Casa e manutenzione', 34, 'Carta di debito', ''),
        ((6, 3), 'Idraulico', 'Casa e manutenzione', 120, 'Conto corrente', ''),
        ((10, 17), 'Tende soggiorno', 'Casa e manutenzione', 89, 'Carta di credito', ''),
        ((12, 6), 'Addobbi di Natale', 'Casa e manutenzione', 28, 'Contanti', ''),
        ((1, 10), 'Saldi invernali', 'Abbigliamento', 142.80, 'Carta di credito', 'Giulia'),
        ((4, 11), 'Maglieria', 'Abbigliamento', 39.90, 'Carta di debito', 'Marco'),
        ((5, 23), 'Scarpe Sofia', 'Abbigliamento', 54.90, 'Carta di debito', 'Sofia'),
        ((9, 12), 'Abbigliamento per la scuola', 'Abbigliamento', 128.50, 'Carta di debito', 'Sofia'),
        ((11, 14), 'Giaccone', 'Abbigliamento', 109, 'Carta di credito', 'Marco'),
        ((12, 27), 'Saldi', 'Abbigliamento', 72, 'Carta di debito', 'Giulia'),
        ((2, 15), 'Compleanno nonna', 'Regali', 40, 'Contanti', ''),
        ((5, 9), 'Festa della mamma', 'Regali', 30, 'Carta di debito', 'Marco'),
        ((6, 6), 'Compleanno Sofia', 'Regali', 85, 'Carta di debito', ''),
        ((10, 3), 'Regalo di nozze amici', 'Regali', 150, 'Conto corrente', ''),
        ((12, 12), 'Regali di Natale', 'Regali', 95.50, 'Carta di credito', ''),
        ((12, 16), 'Regali di Natale', 'Regali', 120, 'Carta di credito', ''),
        ((12, 21), 'Regali di Natale', 'Regali', 64.90, 'Carta di debito', ''),
        ((4, 4), 'Weekend di Pasqua', 'Vacanze', 284, 'Carta di credito', ''),
        ((7, 10), 'Acconto casa vacanze', 'Vacanze', 400, 'Conto corrente', ''),
        ((8, 8), 'Saldo casa vacanze', 'Vacanze', 850, 'Conto corrente', ''),
        ((8, 16), 'Ristoranti in vacanza', 'Vacanze', 186.50, 'Carta di credito', ''),
        ((2, 19), 'Riparazione auto', 'Imprevisti', 384.20, 'Carta di debito', 'Marco'),
        ((10, 22), 'Riparazione caldaia', 'Imprevisti', 165, 'Conto corrente', ''),
    ]
    for (m, g), desc, cat, imp, conto, chi in una_tantum:
        add(dt.date(ANNO, m, g), desc, cat, imp, conto, chi)
    mov.sort(key=lambda r: (r[0], r[2], r[1]))
    return mov


def esempio_budget():
    """{categoria: (importo tipico mensile o None, {mese: importo diverso})}"""
    return {
        'Stipendio Marco': (1850, {}), 'Stipendio Giulia': (1480, {}),
        'Tredicesima e quattordicesima': (None, {12: 3200}),
        'Assegno unico figli': (145, {}), 'Rimborsi e altre entrate': (None, {7: 350}),
        'Affitto o rata mutuo': (680, {}),
        'Condominio': (None, {1: 270, 4: 270, 7: 270, 10: 270}),
        'Luce': (75, {}),
        'Gas': (None, {1: 160, 2: 150, 3: 120, 4: 80, 5: 45, 6: 30, 7: 25, 8: 25, 9: 30,
                       10: 60, 11: 110, 12: 160}),
        'Acqua': (None, {2: 95, 5: 95, 8: 95, 11: 95}),
        'TARI': (None, {4: 110, 7: 110, 11: 110}),
        'Internet e telefono': (46, {}), 'Assicurazioni': (None, {3: 480, 6: 170}),
        'Bollo auto': (None, {4: 235}), 'Abbonamenti': (25, {9: 315}), 'Rate di prestiti': (310, {}),
        'Spesa alimentare': (520, {12: 620}), 'Carburante': (170, {8: 220}), 'Trasporti': (35, {}),
        'Farmacia e salute': (50, {}), 'Scuola e figli': (80, {7: 320, 8: 0, 9: 420}),
        'Casa e manutenzione': (30, {}), 'Abbigliamento': (60, {}), 'Svago': (150, {12: 220}),
        'Regali': (30, {12: 300}), 'Vacanze': (None, {4: 300, 7: 400, 8: 1100}), 'Imprevisti': (60, {}),
        'Fondo emergenze': (150, {}), 'Obiettivi di risparmio': (200, {}), 'Investimenti': (100, {}),
    }


ESEMPIO_IMPOSTAZIONI = {'saldo': 3200, 'famiglia': 'Famiglia Rossi', 'data_rif': dt.date(ANNO, 12, 31),
                        'membri': ['Marco', 'Giulia', 'Sofia']}
ESEMPIO_OBIETTIVI = [
    ('Fondo per le emergenze', 12000, dt.date(2028, 12, 31), 6300),
    ('Vacanza estate 2027', 2400, dt.date(2027, 6, 30), 1200),
    ('Auto nuova', 9000, dt.date(2029, 12, 31), 2000),
    ('Regali di Natale 2026', 400, dt.date(2026, 12, 1), 400),
    ('Corso di inglese Sofia', 600, dt.date(2027, 8, 31), 150),
]
ESEMPIO_DEBITI = [  # nome, saldo residuo, TAN annuo, rata mensile
    ('Mutuo casa', 118400, 0.0295, 680),
    ('Prestito auto', 6150, 0.069, 210),
    ('Carta revolving', 1120, 0.159, 60),
    ('Lavatrice a tasso zero', 360, 0.0, 40),
]
ESEMPIO_DEBITI_OPZIONI = {'metodo': 'Valanga', 'extra': 100}


# ---------------------------------------------------------------- fogli
def sheet_elenchi(wb):
    ws = wb.create_sheet('Elenchi')
    ws.sheet_state = 'hidden'
    for col, head, values in [('A', 'Mesi', MESI), ('B', 'Mesi brevi', MESI_BREVI), ('C', 'Tipi', TIPI),
                              ('D', 'Mese di riferimento', ['Automatico'] + MESI),
                              ('E', 'Metodo', ['Valanga', 'Palla di neve'])]:
        put(ws, f'{col}1', head, f=font(10, True))
        for i, v in enumerate(values):
            put(ws, f'{col}{i + 2}', v)
    widths(ws, {'A': 14, 'B': 10, 'C': 26, 'D': 20, 'E': 16})
    put(ws, 'G1', 'Elenchi usati dai menu a tendina. Non modificare.', f=font(10, color=MUTED))
    protect(ws)


def sheet_istruzioni(wb):
    ws = wb.create_sheet('Istruzioni')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, 'B': 5, 'C': 112})
    put(ws, 'B1', f'Budget familiare {ANNO}', f=font(20, True, TEAL))
    ws.row_dimensions[1].height = 34
    put(ws, 'B2', f'Istruzioni d\'uso · NettoChiaro · versione {VERSIONE}', f=font(10, color=MUTED))
    r = 4

    def sec(text):
        nonlocal r
        r += 1
        section(ws, f'B{r}', text, span_to=3)
        ws.row_dimensions[r].height = 20
        r += 1

    def line(text, bullet='•', link=None):
        nonlocal r
        put(ws, f'B{r}', bullet, f=font(10, True, TEAL), al=Alignment(horizontal='right', vertical='top'))
        c = put(ws, f'C{r}', text, f=font(10, underline='single', color=TEAL) if link else font(), al=WRAP)
        if link:
            c.hyperlink = link
        lines = max(1, -(-len(text) // 118))
        ws.row_dimensions[r].height = 15 * lines + 3
        r += 1

    sec('Legenda')
    inp(ws, f'B{r}')
    put(ws, f'C{r}', 'Celle gialle: da compilare. Tutte le altre contengono formule e si aggiornano da sole.', al=LEFT_C)
    ws[f'B{r}'].protection = Protection(locked=True)
    r += 1
    put(ws, f'B{r}', '', bg=WARN)
    put(ws, f'C{r}', 'Celle arancioni: segnalano un dato da controllare (data mancante, categoria non trovata, nome doppio).', al=LEFT_C)
    r += 1
    put(ws, f'B{r}', '', bg=TEAL)
    put(ws, f'C{r}', 'Intestazioni verdi: titoli delle colonne. I fogli sono protetti senza password per evitare di cancellare le formule.', al=LEFT_C)
    r += 1

    sec('Come iniziare')
    steps = [
        'Impostazioni: anno, saldo di cassa al 1° gennaio (conto corrente più contanti), nomi dei membri della famiglia (facoltativi) e conti o metodi di pagamento.',
        'Categorie: controlla l\'elenco. Puoi rinominare le voci, aggiungerne nelle righe vuote e scegliere il tipo: Entrata, Spesa fissa, Spesa variabile, Risparmio e investimenti.',
        'Budget: scrivi l\'importo tipico mensile di ogni categoria; i 12 mesi si compilano da soli. Per le spese che non sono mensili (TARI, bollo auto, assicurazioni) scrivi l\'importo nel mese in cui le paghi.',
        'Movimenti: registra ogni entrata e uscita con data, descrizione, categoria scelta dal menu e importo. Conto e persona sono facoltativi.',
        'Mensile e Dashboard: si aggiornano da soli. Mostrano il consuntivo per categoria e per mese, la differenza dal budget, il risparmio e i grafici.',
        'Obiettivi e Debiti: elenca cosa vuoi mettere da parte e i debiti in corso. Il foglio calcola la quota mensile necessaria e l\'ordine di estinzione consigliato.',
    ]
    for i, s in enumerate(steps, 1):
        line(s, bullet=f'{i}.')

    sec('Regole utili')
    for s in [
        'Gli importi si scrivono sempre positivi: è la categoria a dire se sono un\'entrata o un\'uscita. Un rimborso su una spesa si registra con importo negativo nella stessa categoria.',
        'I versamenti su conti di risparmio, investimenti o fondo pensione vanno nelle categorie di tipo «Risparmio e investimenti». Non sono spese, ma escono dal saldo di cassa.',
        'Risparmio = entrate − spese. Tasso di risparmio = risparmio / entrate. Avanzo = entrate − spese − versamenti a risparmio.',
        'Saldo di cassa = saldo al 1° gennaio + entrate − spese − versamenti a risparmio. Se registri tutti i movimenti, si avvicina al totale di conto corrente e contanti (le spese con carta di credito escono dal conto più tardi).',
        'Contano solo i movimenti con una data dell\'anno impostato. La colonna «Controllo» dei Movimenti segnala date mancanti, categorie non trovate e movimenti di altri anni.',
        'Il registro ha 2.000 righe. In Categorie non eliminare le righe: svuota le celle. Se rinomini una categoria già usata, correggi anche i movimenti con il vecchio nome.',
        'Mese di riferimento (Impostazioni): «Automatico» usa il mese della data di riferimento (se la data è vuota, quello di oggi). Serve per i confronti «da inizio anno» di Mensile e Dashboard.',
        'Per un mese di budget diverso dal solito scrivi il valore nella cella del mese. Per tornare all\'importo tipico copia la cella di un mese non modificato.',
    ]:
        line(s)

    sec('Protezione, stampa e nuovo anno')
    for s in [
        'Per togliere la protezione (non c\'è password): Excel, Revisione > Rimuovi protezione foglio; LibreOffice, Strumenti > Proteggi foglio; Google Sheets, Dati > Proteggi fogli e intervalli.',
        'Per ordinare i movimenti togli prima la protezione del foglio Movimenti. Per stampare solo una parte del registro seleziona le righe e scegli «Stampa selezione».',
        'Nuovo anno: salva una copia del file, cambia l\'anno in Impostazioni, scrivi come saldo iniziale il saldo di cassa di fine anno, svuota i movimenti e aggiorna budget, obiettivi e debiti.',
    ]:
        line(s)

    sec('Compatibilità')
    line('Excel 2016 o successivo (Windows e Mac), Google Sheets (carica il file su Google Drive e aprilo con Google Sheets), LibreOffice Calc. Nessuna macro, nessun collegamento esterno.')
    line('Excel apre i file scaricati da internet in «Visualizzazione protetta», dove i calcoli non si aggiornano: fai clic su «Abilita modifica».')
    line('In Google Sheets la protezione dei fogli può non restare attiva: scrivi solo nelle celle gialle. I grafici possono avere uno stile leggermente diverso.')

    sec('Avvertenze')
    line('Strumento di organizzazione personale: non è consulenza finanziaria né fiscale. Le stime su obiettivi e debiti ipotizzano tassi e rate costanti; controlla sempre i dati del tuo contratto.')
    line('Calcolo dello stipendio netto: nettochiaro.com/calcolo-stipendio-netto', link='https://nettochiaro.com/calcolo-stipendio-netto/')
    line('Calcolo della rata del mutuo: nettochiaro.com/calcolo-rata-mutuo', link='https://nettochiaro.com/calcolo-rata-mutuo/')
    print_setup(ws)
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToHeight = 1   # una pagina sola: le avvertenze non finiscono su un foglio a parte
    protect(ws)
    return ws


def sheet_impostazioni(wb, ex):
    ws = wb.create_sheet('Impostazioni')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, 'B': 40, 'C': 24, 'D': 78})
    title(ws, 'B1', 'Impostazioni', 'Compila le celle gialle. Le altre si aggiornano da sole.', 'B2')
    rows = [
        (4, 'Anno del budget', ANNO, '0', 'Contano solo i movimenti con una data di quest\'anno.'),
        (5, 'Saldo di cassa al 1° gennaio', ex['saldo'] if ex else 0, EUR_IN,
         'Conto corrente più contanti a inizio anno. Esclude i conti di risparmio e gli investimenti.'),
        (6, 'Nome della famiglia (facoltativo)', ex['famiglia'] if ex else None, None, 'Compare nel titolo della Dashboard.'),
        (7, 'Mese di riferimento', 'Automatico', None,
         '«Automatico» (o vuoto) usa il mese della data di riferimento. Serve per i confronti «da inizio anno».'),
        (8, 'Data di riferimento (facoltativa)', ex['data_rif'] if ex else None, DATE,
         'Lascia vuoto per usare la data di oggi. Serve per obiettivi, debiti e mese automatico.'),
    ]
    for r, label, val, fmt, note in rows:
        put(ws, f'B{r}', label, f=font(10, True), al=LEFT_C, border=B_ROW)
        inp(ws, f'C{r}', val, fmt=fmt, al=RIGHT_C if fmt else LEFT_C)
        put(ws, f'D{r}', note, f=font(9, color=MUTED), al=LEFT_C, border=B_ROW)
        ws.row_dimensions[r].height = 20
    dv_num(ws, 'C4', '2000', '2099', 'Inserisci un anno, ad esempio 2026.')
    dv_num(ws, 'C5')
    dv_list(ws, 'Elenchi!$D$2:$D$14', 'C7', 'Mese non valido', 'Scegli «Automatico» o un mese dal menu.')
    dv_date(ws, 'C8')

    section(ws, 'B10', 'Valori usati dal foglio (calcolati, non modificare)', span_to=4)
    calc = [
        (11, 'Data di riferimento usata', '=IF(C8="",TODAY(),C8)', DATE),
        (12, 'Mese di riferimento usato (numero)',
         '=IF(OR(C7="",C7="Automatico"),IF(YEAR(C11)=C4,MONTH(C11),IF(YEAR(C11)>C4,12,1)),IFERROR(MATCH(C7,Elenchi!$A$2:$A$13,0),12))', '0'),
        (13, 'Mese di riferimento usato (nome)', '=INDEX(Elenchi!$A$2:$A$13,C12)', None),
    ]
    for r, label, formula, fmt in calc:
        put(ws, f'B{r}', label, al=LEFT_C, border=B_ROW)
        put(ws, f'C{r}', formula, f=font(10, True), fmt=fmt, al=RIGHT_C, border=B_ROW)
        ws.row_dimensions[r].height = 18

    section(ws, 'B15', 'Membri della famiglia (facoltativo)', span_to=4)
    put(ws, 'D15', 'Per la colonna «Chi» dei Movimenti e il riepilogo per persona in Dashboard.', f=font(9, color=MUTED))
    for i in range(N_MEM):
        r = MEM0 + i
        put(ws, f'B{r}', f'Membro {i + 1}', al=LEFT_C, border=B_ROW)
        val = ex['membri'][i] if ex and i < len(ex['membri']) else None
        inp(ws, f'C{r}', val)
    section(ws, f'B{CON0 - 1}', 'Conti e metodi di pagamento', span_to=4)
    put(ws, f'D{CON0 - 1}', 'Per la colonna «Conto o metodo» dei Movimenti. Puoi rinominarli.', f=font(9, color=MUTED))
    for i in range(N_CON):
        r = CON0 + i
        put(ws, f'B{r}', f'Conto {i + 1}', al=LEFT_C, border=B_ROW)
        inp(ws, f'C{r}', CONTI[i] if i < len(CONTI) else None)
    print_setup(ws)
    protect(ws)
    return ws


def sheet_categorie(wb, cats):
    ws = wb.create_sheet('Categorie')
    widths(ws, {'A': 32, 'B': 26, 'C': 70, 'D': 24})
    title(ws, 'A1', 'Categorie', 'Rinomina, aggiungi o svuota le categorie. Non eliminare le righe: per togliere una categoria cancella nome e tipo.', 'A2')
    put(ws, 'A3', 'Se rinomini una categoria già usata, correggi anche i movimenti: il registro segnala «Categoria non trovata».', f=font(10, color=MUTED))
    put(ws, 'A4', 'Nei nomi non usare i simboli * ? ~ e non iniziare con < > =.', f=font(10, color=MUTED))
    header(ws, 5, [('A', 'Categoria'), ('B', 'Tipo'), ('C', 'Note'), ('D', 'Controllo')], height=22)
    for i in range(N_CAT):
        r = CAT0 + i
        name, tipo, note = cats[i] if i < len(cats) else (None, None, None)
        inp(ws, f'A{r}', name)
        inp(ws, f'B{r}', tipo)
        inp(ws, f'C{r}', note or None)
        # * ? ~ sono caratteri jolly per SUMIFS/COUNTIF/MATCH e < > = all'inizio sono operatori: il nome
        # verrebbe interpretato come criterio e sommerebbe anche movimenti di altre categorie.
        simboli = (f'OR(ISNUMBER(FIND("*",A{r})),ISNUMBER(FIND("?",A{r})),ISNUMBER(FIND("~",A{r})),'
                   f'ISNUMBER(FIND(LEFT(A{r},1),"<>=")))')
        put(ws, f'D{r}', f'=IF(AND(A{r}="",B{r}=""),"",IF(A{r}="","Manca il nome",IF({simboli},"Simbolo non ammesso",'
                         f'IF(COUNTIF($A${CAT0}:$A${CAT1},A{r})>1,"Nome doppio",IF(B{r}="","Manca il tipo","")))))',
            f=font(10, True, WARN_T), border=B_ROW)
    dv_list(ws, 'Elenchi!$C$2:$C$5', f'B{CAT0}:B{CAT1}', 'Tipo non valido', 'Scegli il tipo dal menu.')
    ws.conditional_formatting.add(f'D{CAT0}:D{CAT1}', FormulaRule(formula=[f'LEN($D{CAT0})>0'], fill=fill(WARN)))
    ws.freeze_panes = f'A{CAT0}'
    print_setup(ws, title_rows='5:5', area=f'A1:D{CAT1}')
    protect(ws)
    return ws


def sheet_movimenti(wb, mov):
    ws = wb.create_sheet('Movimenti')
    widths(ws, {'A': 12, 'B': 36, 'C': 28, 'D': 14, 'E': 18, 'F': 14, 'G': 24, 'H': 26})
    title(ws, 'A1', 'Movimenti', 'Un movimento per riga. Importi sempre positivi: il tipo lo decide la categoria. Un rimborso su una spesa si scrive in negativo.', 'A2')
    put(ws, 'C3', 'Movimenti registrati', f=font(10, color=MUTED), al=RIGHT_C)
    put(ws, 'D3', f'=COUNT(D{MOV0}:D{MOV1})', f=font(10, True), fmt='0', al=RIGHT_C)
    put(ws, 'G3', 'Righe da controllare', f=font(10, color=MUTED), al=RIGHT_C)
    put(ws, 'H3', f'=COUNTIF(H{MOV0}:H{MOV1},"?*")', f=font(10, True), fmt='0', al=LEFT_C)
    ws.conditional_formatting.add('H3', FormulaRule(formula=['$H$3>0'], fill=fill(WARN)))
    header(ws, 5, [('A', 'Data'), ('B', 'Descrizione'), ('C', 'Categoria'), ('D', 'Importo'),
                   ('E', 'Conto o metodo'), ('F', 'Chi (facoltativo)'), ('G', 'Tipo'), ('H', 'Controllo')], height=22)
    for i in range(N_MOV):
        r = MOV0 + i
        row = mov[i] if i < len(mov) else (None,) * 6
        d, desc, cat, imp, conto, chi = row
        inp(ws, f'A{r}', d, fmt=DATE, al=Alignment(horizontal='center', vertical='center'))
        inp(ws, f'B{r}', desc)
        inp(ws, f'C{r}', cat)
        inp(ws, f'D{r}', imp, fmt=EUR_IN, al=RIGHT_C)
        inp(ws, f'E{r}', conto, al=Alignment(horizontal='left', vertical='center', indent=1))
        inp(ws, f'F{r}', chi or None)
        put(ws, f'G{r}', f'=IF(C{r}="","",IFERROR(INDEX({C_TIPI},MATCH(C{r},{C_NOMI},0))&"",""))',
            f=font(10, color=MUTED), al=LEFT_C, border=B_ROW)
        put(ws, f'H{r}',
            f'=IF(AND(A{r}="",C{r}="",D{r}=""),"",IF(D{r}="","Manca l\'importo",IF(NOT(ISNUMBER(D{r})),"Importo non valido",'
            f'IF(NOT(ISNUMBER(A{r})),"Data mancante o non valida",IF(YEAR(A{r})<>{ANNO_REF},"Fuori dall\'anno",'
            f'IF(C{r}="","Manca la categoria",IF(ISNA(MATCH(C{r},{C_NOMI},0)),"Categoria non trovata",'
            f'IF(G{r}="","Categoria senza tipo",""))))))))',
            f=font(10, True, WARN_T), al=LEFT_C, border=B_ROW)
    dv_date(ws, f'A{MOV0}:A{MOV1}')
    dv_list(ws, C_NOMI, f'C{MOV0}:C{MOV1}', 'Categoria non valida',
            'Scegli una categoria dal menu. Per aggiungerne una nuova usa il foglio Categorie.')
    dv_num(ws, f'D{MOV0}:D{MOV1}')
    dv_list(ws, f'Impostazioni!$C${CON0}:$C${CON0 + N_CON - 1}', f'E{MOV0}:E{MOV1}', 'Conto non valido',
            'Scegli dal menu. I conti si modificano nel foglio Impostazioni.')
    dv_list(ws, f'Impostazioni!$C${MEM0}:$C${MEM0 + N_MEM - 1}', f'F{MOV0}:F{MOV1}', 'Nome non valido',
            'Scegli dal menu. I nomi si inseriscono nel foglio Impostazioni.')
    ws.conditional_formatting.add(f'H{MOV0}:H{MOV1}', FormulaRule(formula=[f'LEN($H{MOV0})>0'], fill=fill(WARN)))
    ws.auto_filter.ref = f'A5:H{MOV1}'
    ws.freeze_panes = f'A{MOV0}'
    print_setup(ws, title_rows='5:5')
    protect(ws, sort=True)
    return ws


SUMMARY = [  # riga, etichetta, tipo (per i colori), formula per colonna X
    (6, 'Entrate', 'Entrata', lambda X: f'SUMIFS({X}${G0}:{X}${G1},$B${G0}:$B${G1},"Entrata")'),
    (7, 'Spese fisse', 'Spesa fissa', lambda X: f'SUMIFS({X}${G0}:{X}${G1},$B${G0}:$B${G1},"Spesa fissa")'),
    (8, 'Spese variabili', 'Spesa variabile', lambda X: f'SUMIFS({X}${G0}:{X}${G1},$B${G0}:$B${G1},"Spesa variabile")'),
    (9, 'Totale spese', 'Spesa totale', lambda X: f'{X}7+{X}8'),
    (10, 'Risparmio e investimenti', 'Risparmio e investimenti',
     lambda X: f'SUMIFS({X}${G0}:{X}${G1},$B${G0}:$B${G1},"Risparmio e investimenti")'),
    (11, 'Avanzo (entrate − spese − risparmio)', 'Risultato', lambda X: f'{X}6-{X}9-{X}10'),
]


def grid_common(ws, first_month_col, last_total_col):
    """Righe di riepilogo e intestazioni categoria comuni a Budget e Mensile."""
    put(ws, 'A13', 'Dettaglio per categoria', f=font(10, True, TEAL_D), bg=TEAL_S)
    for c in range(1, last_total_col + 1):
        ws.cell(13, c).fill = fill(TEAL_S)
    for i in range(N_CAT):
        r = G0 + i
        cr = CAT0 + i
        put(ws, f'A{r}', f'=IF(Categorie!$A{cr}="","",Categorie!$A{cr})', border=B_ROW)
        put(ws, f'B{r}', f'=IF(A{r}="","",Categorie!$B{cr}&"")', f=font(9, color=MUTED), border=B_ROW)
    # mesi (riga 4 nascosta) per i confronti da inizio anno
    for m in range(12):
        col = L(first_month_col + m)
        put(ws, f'{col}4', m + 1, f=font(8, color=MUTED))
        put(ws, f'{col}5', MESI_BREVI[m], f=font(10, True, 'FFFFFF'), bg=TEAL, al=CENTER)
    ws.row_dimensions[4].hidden = True


def summary_style(ws, r, label, tipo, ncols):
    bold = r in (9, 11, 12)
    put(ws, f'A{r}', label, f=font(10, True), bg=SURF2 if bold else None, border=B_TOT if bold else B_ROW)
    put(ws, f'B{r}', tipo, f=font(9, color=MUTED), bg=SURF2 if bold else None, border=B_TOT if bold else B_ROW)
    for c in range(3, ncols + 1):
        cell = ws.cell(r, c)
        cell.font = font(10, bold)
        cell.border = B_TOT if bold else B_ROW
        if bold:
            cell.fill = fill(SURF2)


def sheet_budget(wb, cats, budget):
    ws = wb.create_sheet('Budget')
    widths(ws, {'A': 34, 'B': 22, 'C': 14, **{L(c): 11.5 for c in range(4, 16)}, 'P': 13})
    title(ws, 'A1', 'Budget', 'Scrivi in colonna C l\'importo tipico mensile: i mesi si compilano da soli. Per un mese diverso (TARI, bollo, assicurazioni) scrivi l\'importo nella cella del mese.', 'A2')
    put(ws, 'A3', 'Per tornare all\'importo tipico copia la cella di un mese non modificato. Le righe di riepilogo in alto si calcolano da sole.', f=font(10, color=MUTED))
    header(ws, 5, [('A', 'Voce'), ('B', 'Tipo'), ('C', 'Importo tipico mensile'), ('P', 'Totale anno')])
    grid_common(ws, 4, 16)
    for r, label, tipo, fx in SUMMARY:
        summary_style(ws, r, label, tipo, 16)
        for c in range(3, 17):
            X = L(c)
            ws[f'{X}{r}'] = '=' + fx(X)
            ws[f'{X}{r}'].number_format = EUR
    r = 12
    summary_style(ws, r, 'Saldo di cassa previsto a fine mese', 'Saldo', 16)
    for m in range(12):
        X = L(4 + m)
        prev = SALDO_REF if m == 0 else f'{L(3 + m)}12'
        ws[f'{X}12'] = f'={prev}+{X}11'
        ws[f'{X}12'].number_format = EUR
    ws['P12'] = '=O12'
    ws['P12'].number_format = EUR
    names = [c[0] for c in cats]
    for i in range(N_CAT):
        r = G0 + i
        name = names[i] if i < len(names) else None
        typical, override = budget.get(name, (None, {})) if budget else (None, {})
        inp(ws, f'C{r}', typical, fmt=EUR_IN, al=Alignment(horizontal='right'))
        for m in range(12):
            X = L(4 + m)
            val = override.get(m + 1, f'=$C{r}')
            inp(ws, f'{X}{r}', val, fmt=EUR, al=Alignment(horizontal='right'))
        put(ws, f'P{r}', f'=SUM(D{r}:O{r})', f=font(10, True), fmt=EUR, border=B_ROW)
    dv_num(ws, f'C{G0}:O{G1}')
    ws.freeze_panes = 'C6'
    print_setup(ws, title_rows='5:5', area=f'A1:P{G1}')
    protect(ws)
    return ws


def diff_rules(ws, col, r0, r1):
    ref = f'{col}{r0}'
    red = (f'AND($B{r0}<>"",OR(AND(LEFT($B{r0},5)="Spesa",{ref}>0.004),'
           f'AND(LEFT($B{r0},5)<>"Spesa",{ref}<-0.004)))')
    green = (f'AND($B{r0}<>"",OR(AND(LEFT($B{r0},5)="Spesa",{ref}<-0.004),'
             f'AND(LEFT($B{r0},5)<>"Spesa",{ref}>0.004)))')
    ws.conditional_formatting.add(f'{col}{r0}:{col}{r1}', FormulaRule(formula=[red], font=Font(name=ARIAL, color=BAD, bold=True), fill=fill(BAD_S)))
    ws.conditional_formatting.add(f'{col}{r0}:{col}{r1}', FormulaRule(formula=[green], font=Font(name=ARIAL, color=GOOD, bold=True)))


def sheet_mensile(wb):
    ws = wb.create_sheet('Mensile')
    widths(ws, {'A': 34, 'B': 22, **{L(c): 11.5 for c in range(3, 15)}, 'O': 13, 'P': 13, 'Q': 13, 'R': 10,
                'S': 2, 'T': 13, 'U': 13, 'V': 13, 'W': 13, 'X': 13, 'Y': 13, 'Z': 8})
    title(ws, 'A1', 'Mensile: consuntivo e differenze dal budget',
          'Somme dei movimenti per categoria e mese. Differenza = consuntivo − budget: verde se va bene, rosso se va male.', 'A2')
    put(ws, 'A3', f'="Mese di riferimento: "&{MESE_NOME_REF}&" (si cambia in Impostazioni)"', f=font(10, True, TEAL_D))
    header(ws, 5, [('A', 'Voce'), ('B', 'Tipo'), ('O', 'Totale anno'), ('P', 'Budget anno'),
                   ('Q', 'Differenza anno'), ('R', '% del budget'),
                   ('W', 'Budget da inizio anno'), ('X', 'Consuntivo da inizio anno'), ('Y', 'Differenza da inizio anno'),
                   ('Z', 'Ordine (tecnico)')], height=42)
    for col, txt in [('T', 'Budget '), ('U', 'Consuntivo '), ('V', 'Differenza ')]:
        put(ws, f'{col}5', f'="{txt}"&LOWER({MESE_NOME_REF})', f=font(10, True, 'FFFFFF'), bg=TEAL, al=CENTER)
    grid_common(ws, 3, 26)
    ws.cell(13, 19).fill = PatternFill()
    put(ws, 'T13', '=' + f'"Mese di riferimento: "&LOWER({MESE_NOME_REF})', f=font(10, True, TEAL_D), bg=TEAL_S)
    put(ws, 'W13', 'Da gennaio al mese di riferimento', f=font(10, True, TEAL_D), bg=TEAL_S)

    cols_val = [L(c) for c in range(3, 18)] + ['T', 'U', 'V', 'W', 'X', 'Y']   # C..Q + T..Y
    for r, label, tipo, fx in SUMMARY:
        summary_style(ws, r, label, tipo, 25)
        ws.cell(r, 19).fill = PatternFill()
        ws.cell(r, 19).border = Border()
        for X in cols_val:
            if X in ('Q', 'V', 'Y'):
                prev = {'Q': ('O', 'P'), 'V': ('U', 'T'), 'Y': ('X', 'W')}[X]
                f = f'={prev[0]}{r}-{prev[1]}{r}'
            else:
                f = '=' + fx(X)
            ws[f'{X}{r}'] = f
            ws[f'{X}{r}'].number_format = EUR
        ws[f'R{r}'] = f'=IF(P{r}=0,"",O{r}/P{r})'
        ws[f'R{r}'].number_format = '0%'
    # saldo di cassa
    summary_style(ws, 12, 'Saldo di cassa a fine mese', 'Saldo', 25)
    ws.cell(12, 19).fill = PatternFill()
    ws.cell(12, 19).border = Border()
    for m in range(12):
        X = L(3 + m)
        prev = SALDO_REF if m == 0 else f'{L(2 + m)}12'
        ws[f'{X}12'] = f'={prev}+{X}11'
    ws['O12'] = '=N12'
    ws['P12'] = '=Budget!P12'
    ws['Q12'] = '=O12-P12'
    ws['T12'] = f'=INDEX(Budget!$D$12:$O$12,{MESE_REF})'
    ws['U12'] = f'=INDEX($C$12:$N$12,{MESE_REF})'
    ws['V12'] = '=U12-T12'
    for X in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'T', 'U', 'V']:
        ws[f'{X}12'].number_format = EUR

    for i in range(N_CAT):
        r = G0 + i
        for m in range(12):
            X = L(3 + m)
            put(ws, f'{X}{r}',
                f'=IF($A{r}="",0,SUMIFS({M_IMP},{M_CAT},$A{r},{M_DATA},">="&DATE({ANNO_REF},{X}$4,1),{M_DATA},"<"&DATE({ANNO_REF},{X}$4+1,1)))',
                fmt=EUR, border=B_ROW)
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})', f=font(10, True), fmt=EUR, border=B_ROW)
        put(ws, f'P{r}', f'=Budget!P{r}', fmt=EUR, border=B_ROW)
        put(ws, f'Q{r}', f'=O{r}-P{r}', fmt=EUR, border=B_ROW)
        put(ws, f'R{r}', f'=IF(P{r}=0,"",O{r}/P{r})', fmt='0%', border=B_ROW)
        put(ws, f'T{r}', f'=INDEX(Budget!$D{r}:$O{r},{MESE_REF})', fmt=EUR, border=B_ROW)
        put(ws, f'U{r}', f'=INDEX($C{r}:$N{r},{MESE_REF})', fmt=EUR, border=B_ROW)
        put(ws, f'V{r}', f'=U{r}-T{r}', fmt=EUR, border=B_ROW)
        put(ws, f'W{r}', f'=SUMIFS(Budget!$D{r}:$O{r},Budget!$D$4:$O$4,"<="&{MESE_REF})', fmt=EUR, border=B_ROW)
        put(ws, f'X{r}', f'=SUMIFS($C{r}:$N{r},$C$4:$N$4,"<="&{MESE_REF})', fmt=EUR, border=B_ROW)
        put(ws, f'Y{r}', f'=X{r}-W{r}', fmt=EUR, border=B_ROW)
        # chiave per la classifica delle spese: totale meno un piccolo scarto, così a pari importo viene prima
        # la categoria più in alto nell'elenco (lo scarto massimo, 73/10^7, è molto sotto il centesimo)
        put(ws, f'Z{r}', f'=IF(AND(O{r}>0,LEFT(B{r},5)="Spesa"),O{r}-ROW()/10000000,0)', f=font(8, color=MUTED), fmt='0.00')
    for col in ('Q', 'V', 'Y'):
        diff_rules(ws, col, 6, G1)
    ws.conditional_formatting.add(f'R6:R{G1}', FormulaRule(formula=['AND(LEFT($B6,5)="Spesa",ISNUMBER(R6),R6>1.00001)'],
                                                           font=Font(name=ARIAL, color=BAD, bold=True)))
    ws.column_dimensions['Z'].hidden = True
    ws.freeze_panes = 'C6'
    print_setup(ws, title_rows='5:5', area=f'A1:Y{G1}')
    ws.page_setup.fitToWidth = 2    # 25 colonne su una pagina sarebbero illeggibili: due pagine di larghezza
    ws.print_title_cols = 'A:B'     # con voce e tipo ripetuti
    protect(ws)
    return ws


def sheet_dashboard(wb):
    ws = wb.create_sheet('Dashboard')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, **{L(c): 12.5 for c in range(2, 14)}, 'N': 2})
    put(ws, 'B1', f'="Budget familiare "&{ANNO_REF}&IF(Impostazioni!$C$6="",""," · "&Impostazioni!$C$6)', f=font(20, True, TEAL))
    ws.row_dimensions[1].height = 34
    put(ws, 'B2', f'="Totali dei movimenti registrati nell\'anno. Confronti con il budget fino a "&LOWER({MESE_NOME_REF})&"."',
        f=font(10, color=MUTED))

    tiles = [
        ('B', 'Entrate', '=Mensile!$O$6', EUR0, '=IF(Mensile!$P$6=0,"",Mensile!$O$6/Mensile!$P$6)', '0%" del budget annuo"'),
        ('D', 'Spese', '=Mensile!$O$9', EUR0, '=IF(Mensile!$P$9=0,"",Mensile!$O$9/Mensile!$P$9)', '0%" del budget annuo"'),
        ('F', 'Risparmio', '=Mensile!$O$6-Mensile!$O$9', EUR0, '=Mensile!$O$10',
         '"di cui versati: "#,##0" €";"di cui versati: -"#,##0" €";"nessun versamento"'),
        ('H', 'Tasso di risparmio', '=IF(Mensile!$O$6=0,0,(Mensile!$O$6-Mensile!$O$9)/Mensile!$O$6)', PCT,
         'risparmio / entrate', None),
        ('J', 'Saldo di cassa', '=Mensile!$O$12', EUR0, f'={SALDO_REF}',
         '"al 1° gennaio: "#,##0" €";"al 1° gennaio: -"#,##0" €";"al 1° gennaio: 0 €"'),
        ('L', 'Debiti residui', f'=Debiti!$C${DEB1 + 1}', EUR0, f'=Debiti!$E${DEB1 + 1}',
         '"rate: "#,##0" € al mese";"rate: -"#,##0" €";"nessuna rata"'),
    ]
    for col, label, value, fmt, note, note_fmt in tiles:
        c2 = L(ws[f'{col}1'].column + 1)
        for r in (4, 5, 6):
            ws.merge_cells(f'{col}{r}:{c2}{r}')
            for cc in (col, c2):
                ws[f'{cc}{r}'].fill = fill(TEAL_S)
        for cc in (col, c2):
            ws[f'{cc}4'].border = Border(top=Side(style='thick', color=TEAL))
        put(ws, f'{col}4', label, f=font(10, True, TEAL_D), bg=TEAL_S, al=Alignment(horizontal='left', vertical='bottom', indent=1))
        put(ws, f'{col}5', value, f=font(18, True), bg=TEAL_S, fmt=fmt, al=Alignment(horizontal='left', vertical='center', indent=1))
        put(ws, f'{col}6', note, f=font(9, color=MUTED), bg=TEAL_S, fmt=note_fmt, al=Alignment(horizontal='left', vertical='top', indent=1))
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 34
    ws.row_dimensions[6].height = 18

    # andamento mensile
    section(ws, 'B8', 'Andamento mensile', span_to=7)
    header(ws, 9, [('B', 'Mese'), ('C', 'Entrate'), ('D', 'Spese'), ('E', 'Risparmio e investimenti'),
                   ('F', 'Avanzo'), ('G', 'Saldo di cassa')], height=30)
    for m in range(12):
        r = 10 + m
        X = L(3 + m)
        put(ws, f'B{r}', MESI_BREVI[m], border=B_ROW)
        for col, src in [('C', 6), ('D', 9), ('E', 10), ('F', 11), ('G', 12)]:
            put(ws, f'{col}{r}', f'=Mensile!{X}${src}', fmt=EUR0D, border=B_ROW)
    put(ws, 'B22', 'Totale', f=font(10, True), bg=SURF2, border=B_TOT)
    for col in 'CDEF':
        put(ws, f'{col}22', f'=SUM({col}10:{col}21)', f=font(10, True), bg=SURF2, fmt=EUR0D, border=B_TOT)
    put(ws, 'G22', '=G21', f=font(10, True), bg=SURF2, fmt=EUR0D, border=B_TOT)

    ch = BarChart()
    ch.type, ch.grouping = 'col', 'clustered'
    ch.title = chart_title('Entrate e spese per mese')
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=9, max_row=21), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=10, max_row=21))
    for s, color in zip(ch.series, (CH_ENTRATE, CH_SPESE)):
        s.graphicalProperties.solidFill = color
        s.graphicalProperties.line.solidFill = color
    ch.gapWidth = 60
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.legend.position = 'b'
    ch.width, ch.height = 12.0, 8.2
    ws.add_chart(ch, 'I8')

    # budget e consuntivo da inizio anno
    section(ws, 'B24', 'Budget e consuntivo da inizio anno', span_to=7)
    ws.merge_cells('B25:C25')
    header(ws, 25, [('B', 'Voce'), ('D', 'Budget'), ('E', 'Consuntivo'), ('F', 'Differenza'), ('G', '% del budget')], height=22)
    rows = [(26, 'Entrate', 6, 'Entrata'), (27, 'Spese fisse', 7, 'Spesa fissa'), (28, 'Spese variabili', 8, 'Spesa variabile'),
            (29, 'Totale spese', 9, 'Spesa totale'), (30, 'Risparmio e investimenti', 10, 'Risparmio'), (31, 'Avanzo', 11, 'Risultato')]
    for r, label, src, tipo in rows:
        ws.merge_cells(f'B{r}:C{r}')
        bold = src in (9, 11)
        bg = SURF2 if bold else None
        put(ws, f'B{r}', label, f=font(10, bold), bg=bg, border=B_ROW)
        ws[f'C{r}'].border = B_ROW
        put(ws, f'D{r}', f'=Mensile!$W${src}', f=font(10, bold), bg=bg, fmt=EUR0D, border=B_ROW)
        put(ws, f'E{r}', f'=Mensile!$X${src}', f=font(10, bold), bg=bg, fmt=EUR0D, border=B_ROW)
        put(ws, f'F{r}', f'=E{r}-D{r}', f=font(10, bold), bg=bg, fmt=EUR0D, border=B_ROW)
        put(ws, f'G{r}', f'=IF(D{r}=0,"",E{r}/D{r})', f=font(10, bold), bg=bg, fmt='0%', border=B_ROW)
        put(ws, f'N{r}', tipo, f=font(8, color='FFFFFF'))  # tipo per i colori (testo bianco)
    red = 'OR(AND(LEFT($N26,5)="Spesa",F26>0.5),AND(LEFT($N26,5)<>"Spesa",F26<-0.5))'
    green = 'OR(AND(LEFT($N26,5)="Spesa",F26<-0.5),AND(LEFT($N26,5)<>"Spesa",F26>0.5))'
    ws.conditional_formatting.add('F26:F31', FormulaRule(formula=[red], font=Font(name=ARIAL, color=BAD, bold=True)))
    ws.conditional_formatting.add('F26:F31', FormulaRule(formula=[green], font=Font(name=ARIAL, color=GOOD, bold=True)))
    put(ws, 'B32', f'="Da gennaio a "&LOWER({MESE_NOME_REF})&". Differenza = consuntivo − budget."', f=font(9, color=MUTED))

    # controlli
    section(ws, 'I24', 'Controlli', span_to=13)
    checks = [
        (25, 'Movimenti registrati', f'=COUNT({M_IMP})'),
        (26, 'Righe del registro da controllare', f'=COUNTIF({M_CTRL},"?*")'),
        (27, 'Categorie da controllare', f'=COUNTIF(Categorie!$D${CAT0}:$D${CAT1},"?*")'),
        (28, 'Debiti da controllare',   # non "?*": LibreOffice lo applica anche ai numeri
         f'=COUNTIF(Debiti!$G${DEB0}:$G${DEB1},"Manca*")+COUNTIF(Debiti!$G${DEB0}:$G${DEB1},"Rata*")'),
    ]
    for r, label, f in checks:
        ws.merge_cells(f'I{r}:L{r}')
        put(ws, f'I{r}', label, border=B_ROW)
        for cc in 'JKL':
            ws[f'{cc}{r}'].border = B_ROW
        put(ws, f'M{r}', f, f=font(10, True), fmt='0', al=RIGHT_C, border=B_ROW)
    ws.conditional_formatting.add('M26:M28', FormulaRule(formula=['M26>0'], fill=fill(WARN)))
    ws.merge_cells('I29:M30')
    put(ws, 'I29', '=IF(M26+M27+M28=0,"Nessun problema rilevato.","Ci sono dati da controllare: vedi le celle arancioni nei fogli Movimenti, Categorie e Debiti.")',
        f=font(9, color=MUTED), al=WRAP)

    # top 10 categorie di spesa
    section(ws, 'B34', 'Le 10 categorie di spesa più alte (anno)', span_to=7)
    ws.merge_cells('C35:E35')
    header(ws, 35, [('B', '#'), ('C', 'Categoria'), ('F', 'Totale anno'), ('G', '% sulle spese')], height=22)
    key = f'Mensile!$Z${G0}:$Z${G1}'
    for k in range(1, 11):
        r = 35 + k
        ws.merge_cells(f'C{r}:E{r}')
        put(ws, f'B{r}', k, f=font(10, True, TEAL), al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'C{r}', f'=IF(LARGE({key},B{r})<=0,"",INDEX(Mensile!$A${G0}:$A${G1},MATCH(LARGE({key},B{r}),{key},0)))', border=B_ROW)
        for cc in 'DE':
            ws[f'{cc}{r}'].border = B_ROW
        put(ws, f'F{r}', f'=IF(C{r}="",0,INDEX(Mensile!$O${G0}:$O${G1},MATCH(LARGE({key},B{r}),{key},0)))', fmt=EUR0D, border=B_ROW)
        put(ws, f'G{r}', f'=IF(OR(F{r}=0,Mensile!$O$9=0),"",F{r}/Mensile!$O$9)', fmt=PCT, border=B_ROW)
    ch2 = BarChart()
    ch2.type = 'bar'
    ch2.title = chart_title('Spese per categoria (anno)')
    ch2.add_data(Reference(ws, min_col=6, min_row=35, max_row=45), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=3, min_row=36, max_row=45))
    ch2.series[0].graphicalProperties.solidFill = TEAL
    ch2.series[0].graphicalProperties.line.solidFill = TEAL
    ch2.legend = None
    ch2.gapWidth = 50
    ch2.x_axis.scaling.orientation = 'maxMin'
    ch2.x_axis.delete = False
    ch2.y_axis.delete = False
    ch2.y_axis.numFmt = '#,##0'
    ch2.y_axis.crosses = 'max'
    ch2.width, ch2.height = 12.0, 7.4
    ws.add_chart(ch2, 'I34')

    # per persona
    section(ws, 'B48', 'Spese ed entrate per persona (anno)', span_to=7)
    ws.merge_cells('B49:C49')
    header(ws, 49, [('B', 'Persona'), ('D', 'Spese'), ('E', 'Entrate'), ('F', '% sulle spese')], height=22)
    per_anno = f'{M_DATA},">="&DATE({ANNO_REF},1,1),{M_DATA},"<"&DATE({ANNO_REF}+1,1,1)'
    for i in range(N_MEM):
        r = 50 + i
        mr = MEM0 + i
        ws.merge_cells(f'B{r}:C{r}')
        # un nome ripetuto in Impostazioni compare una volta sola (altrimenti verrebbe contato due volte)
        put(ws, f'B{r}', f'=IF(OR(Impostazioni!$C${mr}="",COUNTIF(Impostazioni!$C${MEM0 - 1}:$C${mr - 1},Impostazioni!$C${mr})>0),"",Impostazioni!$C${mr})', border=B_ROW)
        ws[f'C{r}'].border = B_ROW
        put(ws, f'D{r}', f'=IF(B{r}="",0,SUMIFS({M_IMP},{M_CHI},B{r},{M_TIPO},"Spesa*",{per_anno}))', fmt=EUR0D, border=B_ROW)
        put(ws, f'E{r}', f'=IF(B{r}="",0,SUMIFS({M_IMP},{M_CHI},B{r},{M_TIPO},"Entrata",{per_anno}))', fmt=EUR0D, border=B_ROW)
        put(ws, f'F{r}', f'=IF(OR(D{r}=0,Mensile!$O$9=0),"",D{r}/Mensile!$O$9)', fmt=PCT, border=B_ROW)
    r = 50 + N_MEM
    ws.merge_cells(f'B{r}:C{r}')
    put(ws, f'B{r}', 'Comuni o senza nome', f=font(10, italic=True), border=B_ROW)
    ws[f'C{r}'].border = B_ROW
    put(ws, f'D{r}', f'=Mensile!$O$9-SUM(D50:D{r - 1})', fmt=EUR0D, border=B_ROW)
    put(ws, f'E{r}', f'=Mensile!$O$6-SUM(E50:E{r - 1})', fmt=EUR0D, border=B_ROW)
    put(ws, f'F{r}', f'=IF(OR(D{r}=0,Mensile!$O$9=0),"",D{r}/Mensile!$O$9)', fmt=PCT, border=B_ROW)
    put(ws, f'B{r + 1}', 'La persona è la colonna «Chi» dei Movimenti; i nomi si inseriscono in Impostazioni.', f=font(9, color=MUTED))
    print_setup(ws, area=f'A1:N{r + 1}')
    protect(ws)
    return ws


def sheet_obiettivi(wb, obiettivi):
    ws = wb.create_sheet('Obiettivi')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 2, 'B': 40, 'C': 15, 'D': 13, 'E': 15, 'F': 15, 'G': 11, 'H': 24, 'I': 11, 'J': 17, 'K': 18})
    title(ws, 'B1', 'Obiettivi di risparmio',
          'Per ogni obiettivo indica importo, scadenza e quanto hai già messo da parte. La quota mensile parte dalla data di riferimento (Impostazioni).', 'B2')
    header(ws, 4, [('B', 'Obiettivo'), ('C', 'Importo obiettivo'), ('D', 'Scadenza'), ('E', 'Già versato'),
                   ('F', 'Mancante'), ('G', '% raggiunta'), ('H', 'Avanzamento'), ('I', 'Mesi rimanenti'),
                   ('J', 'Quota mensile necessaria'), ('K', 'Stato')], height=32)
    for i in range(N_OBI):
        r = OBI0 + i
        o = obiettivi[i] if obiettivi and i < len(obiettivi) else (None,) * 4
        inp(ws, f'B{r}', o[0])
        inp(ws, f'C{r}', o[1], fmt=EUR_IN, al=Alignment(horizontal='right'))
        inp(ws, f'D{r}', o[2], fmt=DATE, al=Alignment(horizontal='center'))
        inp(ws, f'E{r}', o[3], fmt=EUR_IN, al=Alignment(horizontal='right'))
        put(ws, f'F{r}', f'=IF(B{r}="","",MAX(0,N(C{r})-N(E{r})))', fmt=EUR, border=B_ROW)
        put(ws, f'G{r}', f'=IF(OR(B{r}="",N(C{r})<=0),"",MIN(1,N(E{r})/C{r}))', fmt='0%', border=B_ROW)
        put(ws, f'H{r}', f'=IF(G{r}="","",REPT("█",ROUND(G{r}*20,0))&REPT("░",20-ROUND(G{r}*20,0)))',
            f=font(9, color=TEAL), border=B_ROW)
        put(ws, f'I{r}', f'=IF(OR(B{r}="",D{r}=""),"",MAX(0,(YEAR(D{r})-YEAR({DATA_REF}))*12+MONTH(D{r})-MONTH({DATA_REF})))',
            fmt='0', al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'J{r}', f'=IF(OR(B{r}="",N(C{r})<=0),"",IF(F{r}=0,0,IF(I{r}="","",F{r}/MAX(1,I{r}))))',
            f=font(10, True), fmt=EUR, border=B_ROW)
        put(ws, f'K{r}', f'=IF(B{r}="","",IF(N(C{r})<=0,"Manca l\'importo",IF(F{r}=0,"Raggiunto",IF(D{r}="","Manca la scadenza",IF(D{r}<{DATA_REF},"Scaduto","In corso")))))',
            al=Alignment(horizontal='center'), border=B_ROW)
    t = OBI0 + N_OBI
    put(ws, f'B{t}', 'Totale', f=font(10, True), bg=SURF2, border=B_TOT)
    for col, f, fmt in [('C', f'=SUM(C{OBI0}:C{t - 1})', EUR), ('D', None, None), ('E', f'=SUM(E{OBI0}:E{t - 1})', EUR),
                        ('F', f'=SUM(F{OBI0}:F{t - 1})', EUR), ('G', f'=IF(C{t}=0,"",E{t}/C{t})', '0%'), ('H', None, None),
                        ('I', None, None), ('J', f'=SUM(J{OBI0}:J{t - 1})', EUR), ('K', None, None)]:
        put(ws, f'{col}{t}', f, f=font(10, True), bg=SURF2, fmt=fmt, border=B_TOT)
    dv_num(ws, f'C{OBI0}:C{t - 1}', '0', '100000000', 'Inserisci un importo positivo.')
    dv_num(ws, f'E{OBI0}:E{t - 1}', '0', '100000000', 'Inserisci un importo positivo.')
    dv_date(ws, f'D{OBI0}:D{t - 1}')
    ws.conditional_formatting.add(f'K{OBI0}:K{t - 1}', FormulaRule(formula=[f'$K{OBI0}="Raggiunto"'], font=Font(name=ARIAL, color=GOOD, bold=True)))
    ws.conditional_formatting.add(f'K{OBI0}:K{t - 1}', FormulaRule(formula=[f'OR($K{OBI0}="Scaduto",LEFT($K{OBI0},5)="Manca")'],
                                                                 font=Font(name=ARIAL, color=BAD, bold=True), fill=fill(BAD_S)))

    s = t + 2
    section(ws, f'B{s}', 'Confronto con il budget', span_to=6)
    rows = [
        (s + 1, 'Quota mensile per tutti gli obiettivi', f'=J{t}'),
        (s + 2, 'Risparmio previsto nel budget (al mese)', '=Budget!P10/12'),
        (s + 3, 'Versato a risparmio (media mensile)', f'=IF({MESE_REF}=0,0,Mensile!X10/{MESE_REF})'),
        (s + 4, 'Margine del budget sugli obiettivi', f'=C{s + 2}-C{s + 1}'),
    ]
    for r, label, f in rows:
        put(ws, f'B{r}', label, border=B_ROW)
        put(ws, f'C{r}', f, f=font(10, True), fmt=EUR, border=B_ROW)
    ws.conditional_formatting.add(f'C{s + 4}', FormulaRule(formula=[f'C{s + 4}<0'], font=Font(name=ARIAL, color=BAD, bold=True)))
    ws.conditional_formatting.add(f'C{s + 4}', FormulaRule(formula=[f'C{s + 4}>0'], font=Font(name=ARIAL, color=GOOD, bold=True)))
    put(ws, f'B{s + 5}', 'I mesi rimanenti vanno dal mese della data di riferimento a quello della scadenza. Se la scadenza è nel mese in corso, la quota è l\'intero importo mancante.',
        f=font(9, color=MUTED))
    ws.freeze_panes = f'A{OBI0}'
    print_setup(ws, area=f'A1:K{s + 5}')
    protect(ws)
    return ws


def sheet_debiti(wb, debiti, opz):
    ws = wb.create_sheet('Debiti')
    ws.sheet_view.showGridLines = False
    widths(ws, {'A': 4, 'B': 34, 'C': 15, 'D': 14, 'E': 14, 'F': 14, 'G': 18, 'H': 13, 'I': 15, 'J': 12, 'K': 12})
    title(ws, 'B1', 'Debiti e piano di estinzione',
          'Elenca mutuo, prestiti e carte rateali. TAN = tasso annuo nominale del contratto. Le stime ipotizzano rata e tasso costanti.', 'B2')
    put(ws, 'B4', 'Metodo di estinzione scelto', f=font(10, True), al=LEFT_C, border=B_ROW)
    inp(ws, 'C4', (opz or {}).get('metodo', 'Valanga'))
    ws.merge_cells('D4:K4')
    put(ws, 'D4', 'Valanga: prima il debito con il tasso più alto (paghi meno interessi). Palla di neve: prima il debito più piccolo (lo chiudi prima).',
        f=font(9, color=MUTED), al=LEFT_C)
    put(ws, 'B5', 'Somma extra al mese per i debiti', f=font(10, True), al=LEFT_C, border=B_ROW)
    inp(ws, 'C5', (opz or {}).get('extra', 0), fmt=EUR_IN, al=RIGHT_C)
    ws.merge_cells('D5:K5')
    put(ws, 'D5', 'Si aggiunge alla rata del primo debito dell\'ordine scelto. L\'effetto è calcolato in fondo al foglio.',
        f=font(9, color=MUTED), al=LEFT_C)
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20
    dv_list(ws, 'Elenchi!$E$2:$E$3', 'C4', 'Metodo non valido', 'Scegli Valanga o Palla di neve.')
    dv_num(ws, 'C5', '0', '100000000', 'Inserisci un importo positivo.')

    header(ws, 7, [('B', 'Debito'), ('C', 'Saldo residuo'), ('D', 'TAN annuo'), ('E', 'Rata mensile'),
                   ('F', 'Interessi del mese (stima)'), ('G', 'Mesi alla fine'), ('H', 'Fine prevista'),
                   ('I', 'Interessi residui (stima)'), ('J', 'Ordine valanga'), ('K', 'Ordine palla di neve')], height=32)
    valid = f'$B${DEB0}:$B${DEB1},"<>",$C${DEB0}:$C${DEB1},">0",$D${DEB0}:$D${DEB1},">=0"'
    for i in range(N_DEB):
        r = DEB0 + i
        d = debiti[i] if debiti and i < len(debiti) else (None,) * 4
        inp(ws, f'B{r}', d[0])
        inp(ws, f'C{r}', d[1], fmt=EUR_IN, al=Alignment(horizontal='right'))
        inp(ws, f'D{r}', d[2], fmt=PCT_IN, al=Alignment(horizontal='right'))
        inp(ws, f'E{r}', d[3], fmt=EUR_IN, al=Alignment(horizontal='right'))
        put(ws, f'F{r}', f'=IF(OR(B{r}="",N(C{r})<=0),"",C{r}*N(D{r})/12)', fmt=EUR, border=B_ROW)
        # mesi alla fine, oppure un avviso (testo) se mancano dati: saldo 0 = debito estinto, nessun avviso
        put(ws, f'G{r}', f'=IF(B{r}="",IF(N(C{r})>0,"Manca il nome",""),IF(C{r}="","Manca il saldo",IF(N(C{r})<=0,"",'
                         f'IF(NOT(ISNUMBER(D{r})),"Manca il TAN",IF(N(E{r})<=0,"Manca la rata",IF(D{r}=0,ROUNDUP(C{r}/E{r},0),'
                         f'IF(E{r}<=C{r}*D{r}/12,"Rata troppo bassa",ROUNDUP(NPER(D{r}/12,-E{r},C{r}),0))))))))',
            fmt='0', al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'H{r}', f'=IF(ISNUMBER(G{r}),DATE(YEAR({DATA_REF}),MONTH({DATA_REF})+G{r},1),"")', fmt='MM/YYYY',
            al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'I{r}', f'=IF(ISNUMBER(G{r}),IF(N(D{r})=0,0,MAX(0,E{r}*NPER(D{r}/12,-E{r},C{r})-C{r})),"")', fmt=EUR, border=B_ROW)
        guard = f'OR(B{r}="",N(C{r})<=0,NOT(ISNUMBER(D{r})))'
        prev = f'$B${DEB0 - 1}:B{r - 1},"<>",$C${DEB0 - 1}:C{r - 1},C{r},$D${DEB0 - 1}:D{r - 1},D{r}'
        put(ws, f'J{r}', f'=IF({guard},"",1+COUNTIFS({valid},$D${DEB0}:$D${DEB1},">"&D{r})'
                         f'+COUNTIFS({valid},$D${DEB0}:$D${DEB1},D{r},$C${DEB0}:$C${DEB1},"<"&C{r})+COUNTIFS({prev}))',
            f=font(10, True, TEAL), al=Alignment(horizontal='center'), border=B_ROW)
        put(ws, f'K{r}', f'=IF({guard},"",1+COUNTIFS({valid},$C${DEB0}:$C${DEB1},"<"&C{r})'
                         f'+COUNTIFS({valid},$C${DEB0}:$C${DEB1},C{r},$D${DEB0}:$D${DEB1},">"&D{r})+COUNTIFS({prev}))',
            f=font(10, True, TEAL), al=Alignment(horizontal='center'), border=B_ROW)
    t = DEB1 + 1
    put(ws, f'B{t}', 'Totale', f=font(10, True), bg=SURF2, border=B_TOT)
    for col, f, fmt in [('C', f'=SUM(C{DEB0}:C{DEB1})', EUR), ('D', f'=IF(C{t}=0,"",F{t}*12/C{t})', PCT_IN),
                        ('E', f'=SUM(E{DEB0}:E{DEB1})', EUR), ('F', f'=SUM(F{DEB0}:F{DEB1})', EUR),
                        ('G', None, None), ('H', None, None), ('I', f'=SUM(I{DEB0}:I{DEB1})', EUR), ('J', None, None), ('K', None, None)]:
        put(ws, f'{col}{t}', f, f=font(10, True), bg=SURF2, fmt=fmt, border=B_TOT)
    put(ws, f'B{t + 1}', 'Il TAN nella riga del totale è la media pesata sui saldi. Per un finanziamento a tasso zero scrivi 0%.', f=font(9, color=MUTED))
    dv_num(ws, f'C{DEB0}:C{DEB1}', '0', '100000000', 'Inserisci un importo positivo.')
    dv_num(ws, f'D{DEB0}:D{DEB1}', '0', '1', 'Inserisci il tasso in percentuale, ad esempio 6,9%.')
    dv_num(ws, f'E{DEB0}:E{DEB1}', '0', '100000000', 'Inserisci un importo positivo.')
    ws.conditional_formatting.add(f'G{DEB0}:G{DEB1}', FormulaRule(formula=[f'AND(ISTEXT($G{DEB0}),LEN($G{DEB0})>0)'], fill=fill(WARN),
                                                                font=Font(name=ARIAL, color=WARN_T, bold=True)))

    p = t + 3   # 21
    section(ws, f'B{p}', 'Ordine di estinzione consigliato', span_to=11)
    ws.merge_cells(f'F{p + 1}:G{p + 1}')
    header(ws, p + 1, [('B', 'Valanga: prima il tasso più alto'), ('C', 'TAN'), ('D', 'Saldo'),
                       ('F', 'Palla di neve: prima il saldo più basso'), ('H', 'Saldo'), ('I', 'TAN')], height=32)
    for k in range(1, N_DEB + 1):
        r = p + 1 + k
        put(ws, f'A{r}', k, f=font(9, True, MUTED), al=Alignment(horizontal='right'))
        mv = f'MATCH({k},$J${DEB0}:$J${DEB1},0)'
        ms = f'MATCH({k},$K${DEB0}:$K${DEB1},0)'
        put(ws, f'B{r}', f'=IFERROR(INDEX($B${DEB0}:$B${DEB1},{mv}),"")', border=B_ROW)
        put(ws, f'C{r}', f'=IFERROR(INDEX($D${DEB0}:$D${DEB1},{mv}),"")', fmt=PCT_IN, border=B_ROW)
        put(ws, f'D{r}', f'=IFERROR(INDEX($C${DEB0}:$C${DEB1},{mv}),"")', fmt=EUR, border=B_ROW)
        ws.merge_cells(f'F{r}:G{r}')
        put(ws, f'F{r}', f'=IFERROR(INDEX($B${DEB0}:$B${DEB1},{ms}),"")', border=B_ROW)
        ws[f'G{r}'].border = B_ROW
        put(ws, f'H{r}', f'=IFERROR(INDEX($C${DEB0}:$C${DEB1},{ms}),"")', fmt=EUR, border=B_ROW)
        put(ws, f'I{r}', f'=IFERROR(INDEX($D${DEB0}:$D${DEB1},{ms}),"")', fmt=PCT_IN, border=B_ROW)
    first_v, first_s = f'B{p + 2}', f'F{p + 2}'
    ws.conditional_formatting.add(f'B{p + 2}:D{p + 2}', FormulaRule(formula=[f'AND($C$4<>"Palla di neve",$B${p + 2}<>"")'],
                                                                  fill=fill(TEAL_S), font=Font(name=ARIAL, bold=True)))
    ws.conditional_formatting.add(f'F{p + 2}:I{p + 2}', FormulaRule(formula=[f'AND($C$4="Palla di neve",$F${p + 2}<>"")'],
                                                                  fill=fill(TEAL_S), font=Font(name=ARIAL, bold=True)))
    e = p + 2 + N_DEB + 1   # 34
    section(ws, f'B{e}', 'Effetto della somma extra sul primo debito', span_to=6)
    tgt = f'C{e + 1}'
    m = f'MATCH({tgt},$B${DEB0}:$B${DEB1},0)'
    rows = [
        (e + 1, 'Debito su cui concentrare l\'extra', f'=IF($C$4="Palla di neve",{first_s},{first_v})', None),
        (e + 2, 'Saldo residuo', f'=IF({tgt}="","",INDEX($C${DEB0}:$C${DEB1},{m}))', EUR),
        (e + 3, 'TAN annuo', f'=IF({tgt}="","",INDEX($D${DEB0}:$D${DEB1},{m}))', PCT_IN),
        (e + 4, 'Rata attuale più somma extra', f'=IF({tgt}="","",INDEX($E${DEB0}:$E${DEB1},{m})+N($C$5))', EUR),
        (e + 5, 'Mesi alla fine con la sola rata', f'=IF({tgt}="","",INDEX($G${DEB0}:$G${DEB1},{m}))', '0'),
        (e + 6, 'Mesi alla fine con rata più extra',
         f'=IF(OR({tgt}="",N(C{e + 4})<=0),"",IF(C{e + 3}=0,ROUNDUP(C{e + 2}/C{e + 4},0),IF(C{e + 4}<=C{e + 2}*C{e + 3}/12,"Rata troppo bassa",ROUNDUP(NPER(C{e + 3}/12,-C{e + 4},C{e + 2}),0))))', '0'),
        (e + 7, 'Mesi in meno', f'=IF(AND(ISNUMBER(C{e + 5}),ISNUMBER(C{e + 6})),C{e + 5}-C{e + 6},"")', '0'),
        (e + 8, 'Interessi risparmiati (stima)',
         f'=IF(AND(ISNUMBER(C{e + 5}),ISNUMBER(C{e + 6})),IF(C{e + 3}=0,0,MAX(0,INDEX($I${DEB0}:$I${DEB1},{m})-(C{e + 4}*NPER(C{e + 3}/12,-C{e + 4},C{e + 2})-C{e + 2}))),"")', EUR),
    ]
    for r, label, f, fmt in rows:
        put(ws, f'B{r}', label, border=B_ROW)
        put(ws, f'C{r}', f, f=font(10, True), fmt=fmt, al=Alignment(horizontal='right') if fmt else LEFT_C, border=B_ROW)
    ws.merge_cells(f'C{e + 1}:E{e + 1}')
    put(ws, f'B{e + 9}', 'Quando estingui il primo debito, aggiungi la sua rata a quella del debito successivo: è l\'effetto «valanga» o «palla di neve».',
        f=font(9, color=MUTED))
    ws.freeze_panes = f'A{DEB0}'
    print_setup(ws, area=f'A1:K{e + 9}')
    ws.page_setup.fitToHeight = 1   # una pagina: l'effetto della somma extra non si spezza
    protect(ws)
    return ws


# ---------------------------------------------------------------- assemblaggio
def build(example):
    wb = Workbook()
    base = Font(name=ARIAL, size=10)
    wb._fonts = IndexedList([base])
    wb._named_styles['Normal'].font = base
    wb.remove(wb.active)
    cats = esempio_categorie() if example else CATEGORIE
    sheet_istruzioni(wb)
    sheet_impostazioni(wb, ESEMPIO_IMPOSTAZIONI if example else None)
    sheet_categorie(wb, cats)
    sheet_movimenti(wb, esempio_movimenti() if example else [])
    sheet_budget(wb, cats, esempio_budget() if example else None)
    sheet_mensile(wb)
    dash = sheet_dashboard(wb)
    sheet_obiettivi(wb, ESEMPIO_OBIETTIVI if example else None)
    sheet_debiti(wb, ESEMPIO_DEBITI if example else None, ESEMPIO_DEBITI_OPZIONI if example else None)
    sheet_elenchi(wb)
    for name in ('Istruzioni', 'Dashboard'):
        wb[name].sheet_properties.tabColor = TEAL
    for name in ('Impostazioni', 'Categorie', 'Movimenti', 'Budget', 'Obiettivi', 'Debiti'):
        wb[name].sheet_properties.tabColor = 'E3C766'
    if example:
        wb.active = wb.sheetnames.index('Dashboard')
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = ws is dash
    wb.calculation.fullCalcOnLoad = True
    p = wb.properties
    p.creator = p.lastModifiedBy = 'NettoChiaro'
    p.title = f'Budget familiare {ANNO}' + (' (esempio)' if example else '')
    p.subject = 'Budget familiare annuale: movimenti, budget, consuntivo, obiettivi e debiti'
    p.created = p.modified = FIXED
    return wb


def save_deterministic(wb, path):
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)', rb'\g<1>2026-01-01T00:00:00Z\g<2>', data)
            if info.filename == 'xl/workbook.xml':
                # titoli di stampa con righe e colonne: openpyxl scrive prima le righe, Excel prima le colonne
                data = re.sub(rb">('[^']+'!\$\d+:\$\d+),('[^']+'!\$[A-Z]+:\$[A-Z]+)<", rb'>\g<2>,\g<1><', data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    with open(path, 'wb') as fh:
        fh.write(out.getvalue())


def main():
    os.makedirs(OUT, exist_ok=True)
    for example, name in ((False, f'{SLUG}.xlsx'), (True, f'{SLUG}-esempio.xlsx')):
        path = os.path.join(OUT, name)
        save_deterministic(build(example), path)
        print(f'scritto {os.path.relpath(path, HERE)} ({os.path.getsize(path) // 1024} KB)')


if __name__ == '__main__':
    main()
