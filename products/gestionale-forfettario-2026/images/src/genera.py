#!/usr/bin/env python3
"""Rigenera le immagini dell'inserzione dai dati del file di esempio.

Uso: python3 images/src/genera.py   (dopo python3 build.py)

1. ricalcola files/...-esempio.xlsx con LibreOffice headless e ne estrae i valori in images/src/dati.js;
2. fotografa con Playwright (node images/src/render.mjs) ogni pagina images/src/NN-*.html a 2000×1500;
3. ricava images/card.jpg (1200×900, ≤ 200 KB) dalla copertina.
"""
import datetime as dt
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

SRC = Path(__file__).resolve().parent
IMAGES = SRC.parent
PROD = IMAGES.parent
sys.dont_write_bytecode = True
sys.path.insert(0, str(PROD))
import build  # noqa: E402

ESEMPIO = PROD / 'files' / f'{build.SLUG}-esempio.xlsx'


def ricalcola(path, tmp):
    profilo = (tmp / 'lo-profile').as_uri()
    subprocess.run(['soffice', f'-env:UserInstallation={profilo}', '--headless', '--calc', '--convert-to', 'xlsx',
                    '--outdir', str(tmp / 'out'), str(path)], capture_output=True, check=True, timeout=300)
    return tmp / 'out' / path.name


def data_it(v):
    return v.strftime('%d/%m/%Y') if isinstance(v, (dt.date, dt.datetime)) else ''


def estrai():
    P = build.carica_parametri()
    formule = load_workbook(ESEMPIO)
    with tempfile.TemporaryDirectory(prefix='img-forf-') as t:
        wb = load_workbook(ricalcola(ESEMPIO, Path(t)), data_only=True)
    imp, das, rie, fat, sca = (wb[n] for n in ('Impostazioni', 'Dashboard', 'Riepilogo mensile', 'Fatture',
                                               'Scadenze'))
    d = {
        'anno': P['ANNO'],
        'impostazioni': {k: imp[c].value for k, c in (('nome', 'B5'), ('attivita', 'B6'), ('previdenza', 'B7'),
                                                       ('imposta', 'B9'), ('riduzione', 'B10'), ('bollo', 'B11'),
                                                       ('coeff', 'B16'), ('aliqC', 'B19'))},
        'dashboard': {k: das[c].value for k, c in (
            ('incassato', 'F11'), ('coeff', 'F12'), ('reddito', 'F13'), ('fissi', 'F14'), ('variabili', 'F15'),
            ('contributi', 'F16'), ('dedotti', 'F17'), ('imponibile', 'F18'), ('aliquota', 'F19'),
            ('imposta', 'F20'), ('totale', 'F21'), ('incidenza', 'F22'), ('netto', 'F23'), ('spese', 'F24'),
            ('nettoSpese', 'F25'), ('nettoMese', 'F26'), ('pct85', 'L11'), ('manca85', 'L13'), ('pct100', 'L14'),
            ('emesse', 'L19'), ('totEmesso', 'L20'), ('incassate', 'L21'), ('daIncN', 'L22'), ('daIncTot', 'L23'),
            ('bolloAnno', 'L24'), ('messaggio', 'B8'))},
        'controlli': [x for x in (das['B31'].value or '').split('\n') if x],
        'mesi': [{
            'mese': rie[f'A{r}'].value, 'incassato': rie[f'B{r}'].value, 'n': rie[f'C{r}'].value,
            'cumulato': rie[f'D{r}'].value, 'pct85': rie[f'E{r}'].value, 'tasseCum': rie[f'J{r}'].value,
            'accantonare': rie[f'K{r}'].value, 'spese': rie[f'L{r}'].value, 'netto': rie[f'M{r}'].value,
        } for r in range(5, 17)],
        'totaleMesi': {k: rie[f'{c}17'].value for k, c in (('incassato', 'B'), ('accantonare', 'K'),
                                                           ('spese', 'L'), ('netto', 'M'))},
        'fatture': [],
        'scadenze': [],
        'coefficienti': [{'label': c['label'], 'coeff': c['coeff'], 'ateco': c['ateco']}
                         for c in P['FORFETTARIO']['coefficienti']],
        'gestioni': [g['label'] for g in P['GESTIONI']],
        'soglie': {'s85': P['FORFETTARIO']['sogliaRicavi'], 's100': P['FORFETTARIO']['sogliaUscitaImmediata']},
        'bollo': P['BOLLO'],
        'gs': P['INPS']['gestioneSeparata'],
    }
    for r in range(5, 5 + len(build.ESEMPIO['fatture'])):
        d['fatture'].append({k: (data_it(fat[f'{c}{r}'].value) if c in 'BC' else fat[f'{c}{r}'].value)
                             for k, c in (('n', 'A'), ('emissione', 'B'), ('incasso', 'C'), ('cliente', 'D'),
                                          ('descrizione', 'E'), ('compenso', 'F'), ('rivalsaSi', 'G'),
                                          ('rivalsa', 'H'), ('imponibile', 'I'), ('bollo', 'J'), ('bolloCliente', 'K'),
                                          ('totale', 'L'), ('ricavo', 'M'), ('stato', 'N'))})
    for r in range(5, 5 + len(build.SCADENZE)):
        d['scadenze'].append({'data': data_it(sca[f'A{r}'].value), 'cosa': sca[f'B{r}'].value,
                              'chi': sca[f'C{r}'].value, 'codice': sca[f'D{r}'].value,
                              'importo': sca[f'E{r}'].value})
    # Testi dei tre avvisi sulle soglie, presi dalla formula della Dashboard.
    f = formule['Dashboard']['B8'].value
    d['avvisiSoglie'] = [s.replace('""', '"') for s in re.findall(r'"((?:[^"]|"")+)"', f) if len(s) > 30]
    (SRC / 'dati.js').write_text('// Generato da genera.py: valori ricalcolati del file di esempio.\n'
                                 f'window.DATI = {json.dumps(d, ensure_ascii=False, indent=1)};\n', encoding='utf-8')
    return d


def card():
    img = Image.open(IMAGES / '01-cover.png').convert('RGB').resize((1200, 900), Image.LANCZOS)
    for q in (88, 84, 80, 75, 70):
        img.save(IMAGES / 'card.jpg', 'JPEG', quality=q, optimize=True, progressive=True)
        if (IMAGES / 'card.jpg').stat().st_size <= 200 * 1024:
            break
    return (IMAGES / 'card.jpg').stat().st_size


def main():
    estrai()
    subprocess.run(['node', str(SRC / 'render.mjs')], check=True)
    size = card()
    print(f'card.jpg: {size // 1024} KB')


if __name__ == '__main__':
    main()
