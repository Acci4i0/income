#!/usr/bin/env python3
"""Verifica le formule del gestionale con LibreOffice e le confronta con src/lib/forfettario.js.

Uso: python3 verifica.py   (dopo python3 build.py)

1. Per ogni scenario compila una copia del file vuoto (Impostazioni + una fattura), la ricalcola con
   LibreOffice headless e confronta i valori della Dashboard con calcolaForfettario() di forfettario.js:
   devono coincidere al centesimo.
2. Ricalcola il file di esempio e controlla i totali con un calcolo indipendente in Python + forfettario.js.
3. Controlla che nessuna cella dei due file contenga errori (#REF!, #NAME?, #VALUE!, #DIV/0!...).
4. Scadenze: niente acconti 2027 d'imposta sostitutiva oltre 85.000 €, niente saldo 2026 oltre 100.000 €.
5. Un file con dati sbagliati (testo al posto di numeri e date) non deve avere celle in errore: le righe
   non valide restano fuori dai totali e i Controlli della Dashboard le segnalano.
Esce con codice 1 se un controllo fallisce. Stampa la tabella di confronto usata nel README.
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.dont_write_bytecode = True
sys.path.insert(0, str(HERE))
import build  # noqa: E402

VUOTO = HERE / 'files' / f'{build.SLUG}.xlsx'
ESEMPIO = HERE / 'files' / f'{build.SLUG}-esempio.xlsx'
ERRORI = ('#REF!', '#NAME?', '#VALUE!', '#DIV/0!', '#N/A', '#NUM!', '#NULL!', 'Err:')

SCENARI = [
    dict(nome='Gestione Separata, 40.000 € al 78%', ricavi=40000, coefficiente=0.78, gestione='separata'),
    dict(nome='Commercianti, 50.000 € al 40%', ricavi=50000, coefficiente=0.40, gestione='commercianti'),
    dict(nome='Artigiano, 30.000 € al 67%, riduzione 35%', ricavi=30000, coefficiente=0.67, gestione='artigiani',
         riduzione35=True),
    dict(nome='Startup 5%, Gestione Separata, 30.000 € al 78%', ricavi=30000, coefficiente=0.78,
         gestione='separata', startup=True),
    dict(nome='GS pensionato, 25.000 € al 67%', ricavi=25000, coefficiente=0.67, gestione='separata-pensionati'),
    dict(nome='Cassa 14,5%, 60.000 € al 78%', ricavi=60000, coefficiente=0.78, gestione='cassa', aliquotaCassa=0.145),
    dict(nome='Artigiano, 95.000 € all\'86% (oltre 56.224 € di reddito)', ricavi=95000, coefficiente=0.86,
         gestione='artigiani'),
    dict(nome='Commercianti, 20.000 € al 40% (sotto il minimale), riduzione 35%', ricavi=20000, coefficiente=0.40,
         gestione='commercianti', riduzione35=True),
    dict(nome='Gestione Separata, 160.000 € al 78% (fuori regime, oltre il massimale)', ricavi=160000,
         coefficiente=0.78, gestione='separata'),
    dict(nome='Artigiano, 15.000 € al 67% (sotto il minimale, senza riduzione)', ricavi=15000, coefficiente=0.67,
         gestione='artigiani'),
    dict(nome='Gestione Separata, 90.000 € al 78% (oltre 85.000 €)', ricavi=90000, coefficiente=0.78,
         gestione='separata'),
    dict(nome='Artigiano, 120.000 € all\'86% (oltre 100.000 €)', ricavi=120000, coefficiente=0.86,
         gestione='artigiani'),
]

# Celle della Dashboard (colonna F, righe del "Dettaglio del calcolo").
DASH = {'incassato': 'F11', 'coeff': 'F12', 'reddito': 'F13', 'fissi': 'F14', 'variabili': 'F15',
        'contributi': 'F16', 'dedotti': 'F17', 'imponibile': 'F18', 'aliquota': 'F19', 'imposta': 'F20',
        'totale': 'F21', 'incidenza': 'F22', 'netto': 'F23', 'spese': 'F24', 'nettoSpese': 'F25'}


def node_calcola(scenari):
    lib = (ROOT / 'src/lib/forfettario.js').as_uri()
    js = (f"import('{lib}').then(({{calcolaForfettario}}) => console.log(JSON.stringify("
          f"{json.dumps(scenari)}.map((s) => calcolaForfettario(s)))))")
    return json.loads(subprocess.run(['node', '-e', js], capture_output=True, text=True, check=True).stdout)


def ricalcola(paths, tmp):
    """Converte con LibreOffice (profilo dedicato, per non scontrarsi con altre istanze) e ritorna i nuovi path."""
    out = tmp / 'out'
    out.mkdir(exist_ok=True)
    profilo = (tmp / 'lo-profile').as_uri()
    subprocess.run(['soffice', f'-env:UserInstallation={profilo}', '--headless', '--calc', '--convert-to', 'xlsx',
                    '--outdir', str(out)] + [str(p) for p in paths], capture_output=True, check=True, timeout=300)
    return [out / p.name for p in paths]


def errori(wb):
    trovati = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERRORI):
                    trovati.append(f'{ws.title}!{c.coordinate}={c.value}')
    return trovati


def etichette(P):
    coeff = {round(c['coeff'], 4): c['label'] for c in reversed(P['FORFETTARIO']['coefficienti'])}
    gest = {g['id']: g['label'] for g in P['GESTIONI']}
    return coeff, gest


def main():
    P = build.carica_parametri()
    coeff_label, gest_label = etichette(P)
    imposte = [f'{build.it_pct(P["FORFETTARIO"]["impostaOrdinaria"])} (aliquota ordinaria)',
               f'{build.it_pct(P["FORFETTARIO"]["impostaStartup"])} (nuova attività, primi 5 anni)']
    falliti = []
    with tempfile.TemporaryDirectory(prefix='verifica-forf-') as t:
        tmp = Path(t)
        # ------------------------------------------------ scenari
        files = []
        for i, s in enumerate(SCENARI):
            wb = load_workbook(VUOTO)
            imp = wb['Impostazioni']
            imp['B6'] = coeff_label[round(s['coefficiente'], 4)]
            imp['B7'] = gest_label[s['gestione']]
            imp['B8'] = s.get('aliquotaCassa')
            imp['B9'] = imposte[1 if s.get('startup') else 0]
            imp['B10'] = 'Sì' if s.get('riduzione35') else 'No'
            fat = wb['Fatture']
            fat['A5'], fat['B5'], fat['C5'] = '1/2026', build.D(2026, 3, 1), build.D(2026, 3, 31)
            fat['F5'], fat['G5'], fat['K5'] = s['ricavi'], 'No', 'No'  # bollo a carico: non entra nei ricavi
            p = tmp / f'scenario{i}.xlsx'
            wb.save(p)
            files.append(p)
        # scenario extra: contributi versati inseriti a mano (deduzione reale)
        wb = load_workbook(VUOTO)
        imp = wb['Impostazioni']
        imp['B6'], imp['B7'], imp['B12'], imp['B13'] = coeff_label[0.78], gest_label['separata'], 5000, 2000
        fat = wb['Fatture']
        fat['B5'], fat['C5'], fat['F5'], fat['K5'] = build.D(2026, 5, 1), build.D(2026, 6, 1), 40000, 'No'
        p_vers = tmp / 'versati.xlsx'
        wb.save(p_vers)
        # scenario extra: dati sbagliati (testo al posto di numeri e date), non devono propagare errori
        wb = load_workbook(VUOTO)
        imp = wb['Impostazioni']
        imp['B6'], imp['B7'], imp['B12'] = coeff_label[0.78], gest_label['separata'], 'mille'
        fat = wb['Fatture']
        for r_, (em, inc, comp) in enumerate([(build.D(2026, 1, 5), build.D(2026, 2, 5), 10000),
                                              (build.D(2026, 1, 5), build.D(2026, 2, 5), '1.000,00'),
                                              (build.D(2026, 1, 5), '31/03/2026', 500),
                                              (build.D(2026, 4, 1), None, 700)], start=5):
            fat[f'B{r_}'], fat[f'C{r_}'], fat[f'F{r_}'], fat[f'K{r_}'] = em, inc, comp, 'No'
        spe = wb['Spese']
        spe['A5'], spe['D5'], spe['A6'], spe['D6'] = build.D(2026, 3, 1), '100,00', build.D(2026, 3, 2), 40
        p_err = tmp / 'errori.xlsx'
        wb.save(p_err)
        ricalcolati = ricalcola(files + [p_vers, p_err, VUOTO, ESEMPIO], tmp)
        attesi = node_calcola([{k: v for k, v in s.items() if k != 'nome'} for s in SCENARI])

        print('| Scenario | Valore | Foglio (LibreOffice) | forfettario.js | Diff. |')
        print('|---|---|---:|---:|---:|')
        for s, r, path in zip(SCENARI, attesi, ricalcolati):
            wb = load_workbook(path, data_only=True)
            d = wb['Dashboard']
            if errori(wb):
                falliti.append(f'{s["nome"]}: errori {errori(wb)[:5]}')
            v = {k: d[c].value for k, c in DASH.items()}
            confronti = [('Reddito', v['reddito'], r['redditoLordo']),
                         ('Contributi fissi', v['fissi'], r['contributi']['fissi']),
                         ('Contributi a percentuale', v['variabili'], r['contributi']['variabili']),
                         ('Contributi totali', v['contributi'], r['contributi']['totale']),
                         ('Imposta sostitutiva', v['imposta'], r['imposta']),
                         ('Totale tasse e contributi', v['totale'], r['totaleTasse']),
                         ('Netto', v['netto'], r['netto'])]
            for lab, foglio, js in confronti:
                diff = abs(foglio - js)
                ok = diff < 0.005
                if not ok:
                    falliti.append(f'{s["nome"]} {lab}: foglio {foglio} js {js}')
                print(f'| {s["nome"]} | {lab} | {build.it_num(foglio, 2)} | {build.it_num(js, 2)} | '
                      f'{build.it_num(diff, 4)} |')
            # somma mensile da accantonare = totale annuo
            rie = wb['Riepilogo mensile']
            if abs(rie['K17'].value - rie['J16'].value) > 0.005:
                falliti.append(f'{s["nome"]}: somma accantonamenti {rie["K17"].value} != totale {rie["J16"].value}')
            msg = d['B8'].value
            atteso = ('esci dal forfettario' if r['fuoriRegime'] else
                      'dal prossimo passi' if s['ricavi'] > P['FORFETTARIO']['sogliaRicavi'] else 'nessun limite')
            if atteso not in msg:
                falliti.append(f'{s["nome"]}: messaggio soglie inatteso: {msg}')
            avviso_minimale = 'minimale' in (d['B31'].value or '')
            if avviso_minimale != any('minimale' in a for a in r['avvisi']):
                falliti.append(f'{s["nome"]}: avviso minimale non coerente con forfettario.js')
            # oltre 85.000 € niente acconti d'imposta sostitutiva 2027; oltre 100.000 € niente saldo 2026
            sca = wb['Scadenze']
            righe = {sca[f'B{r_}'].value: sca[f'E{r_}'].value for r_ in range(5, 5 + len(build.SCADENZE))}
            acc = righe['Imposta sostitutiva: 1° acconto 2027']
            if (s['ricavi'] > P['FORFETTARIO']['sogliaRicavi']) != (isinstance(acc, str) and acc.startswith('Non dovuto')):
                falliti.append(f'{s["nome"]}: acconto 2027 in Scadenze = {acc}')
            saldo = righe['Imposta sostitutiva: saldo 2026']
            if r['fuoriRegime'] != (isinstance(saldo, str) and saldo.startswith('Non dovuto')):
                falliti.append(f'{s["nome"]}: saldo 2026 in Scadenze = {saldo}')

        # deduzione reale: imposta = (reddito - versati) * 15%
        wb = load_workbook(ricalcolati[len(SCENARI)], data_only=True)
        d = wb['Dashboard']
        atteso = (40000 * 0.78 - 5000) * 0.15
        print(f'\nDeduzione dei contributi versati (5.000 €): imposta {d["F20"].value:.2f}, attesa {atteso:.2f}')
        if abs(d['F20'].value - atteso) > 0.005:
            falliti.append('deduzione contributi versati')
        sca = wb['Scadenze']
        saldo = next(sca[f'E{r_}'].value for r_ in range(5, 30) if sca[f'B{r_}'].value == 'Imposta sostitutiva: saldo 2026')
        if abs(saldo - (atteso - 2000)) > 0.005:
            falliti.append(f'saldo 2026 in Scadenze: {saldo}')

        # dati sbagliati: nessun errore, le righe non valide escluse, i controlli le segnalano
        wb = load_workbook(ricalcolati[len(SCENARI) + 1], data_only=True)
        d = wb['Dashboard']
        ctrl = d['B31'].value or ''
        print(f'\nDati sbagliati: incassato {d["F11"].value}, spese {d["F24"].value}, controlli: {ctrl!r}')
        if errori(wb):
            falliti.append(f'dati sbagliati: errori {errori(wb)[:5]}')
        if d['F11'].value != 10000 or d['F24'].value != 40:
            falliti.append(f'dati sbagliati: incassato {d["F11"].value}, spese {d["F24"].value}')
        for atteso in ('compenso scritto come testo: 1', 'Date delle fatture scritte come testo: 1',
                       'Spese senza una data valida', 'In Impostazioni un importo'):
            if atteso not in ctrl:
                falliti.append(f'dati sbagliati: manca il controllo "{atteso}"')
        stati = [wb['Fatture'][f'N{r_}'].value for r_ in range(5, 9)]
        if stati != ['Incassata', 'Compenso non valido', 'Data non valida', 'Da incassare']:
            falliti.append(f'dati sbagliati: stati {stati}')

        # ------------------------------------------------ file vuoto ed esempio
        for path in ricalcolati[-2:]:
            wb = load_workbook(path, data_only=True)
            e = errori(wb)
            print(f'\n{path.name}: {len(e)} celle con errori')
            if e:
                falliti.append(f'{path.name}: {e[:10]}')
        wb = load_workbook(ricalcolati[-1], data_only=True)
        E = build.ESEMPIO
        anno = P['ANNO']
        rivalsa, bollo, soglia_b = P['INPS']['gestioneSeparata']['rivalsa'], P['BOLLO']['importo'], P['BOLLO']['soglia']
        mesi = [0.0] * 12
        for n, em, inc, cli, desc, comp, riv, bol in E['fatture']:
            if inc is None or inc.year != anno:
                continue
            imponibile = comp + (round(comp * rivalsa, 2) if riv == 'Sì' else 0)
            ricavo = imponibile + (bollo if imponibile > soglia_b and bol != 'No' else 0)
            mesi[inc.month - 1] += ricavo
        spese = sum(x[3] for x in E['spese'] if x[0].year == anno)
        ricavi = round(sum(mesi), 2)
        r = node_calcola([dict(ricavi=ricavi, coefficiente=0.78, gestione='separata', costiReali=spese)])[0]
        d = wb['Dashboard']
        rie = wb['Riepilogo mensile']
        print(f'Esempio: incassato {d["F11"].value:.2f} (atteso {ricavi:.2f}), spese {d["F24"].value:.2f} '
              f'(attese {spese:.2f})')
        print(f'Esempio: contributi {d["F16"].value:.2f} / {r["contributi"]["totale"]:.2f}, imposta '
              f'{d["F20"].value:.2f} / {r["imposta"]:.2f}, netto dopo spese {d["F25"].value:.2f} / {r["netto"]:.2f}')
        for k, (a, b) in {'incassato': (d['F11'].value, ricavi), 'spese': (d['F24'].value, spese),
                          'contributi': (d['F16'].value, r['contributi']['totale']),
                          'imposta': (d['F20'].value, r['imposta']),
                          'netto dopo spese': (d['F25'].value, r['netto'])}.items():
            if abs(a - b) > 0.005:
                falliti.append(f'esempio {k}: {a} != {b}')
        for m in range(12):
            if abs(rie[f'B{5 + m}'].value - mesi[m]) > 0.005:
                falliti.append(f'esempio mese {m + 1}: {rie[f"B{5 + m}"].value} != {mesi[m]}')
        if abs(rie['K17'].value - rie['J16'].value) > 0.005:
            falliti.append('esempio: somma accantonamenti diversa dal totale')
        ctrl = [x for x in (d['B31'].value or '').split('\n') if x]
        print('Esempio, controlli:', ctrl)
        print('Esempio, scadenze:', [(wb['Scadenze'][f'A{r_}'].value.strftime('%d/%m/%Y'),
                                      wb['Scadenze'][f'E{r_}'].value) for r_ in range(5, 5 + len(build.SCADENZE))])

    if falliti:
        print('\nCONTROLLI FALLITI:')
        for x in falliti:
            print(' -', x)
        sys.exit(1)
    print('\nTutti i controlli superati.')


if __name__ == '__main__':
    main()
