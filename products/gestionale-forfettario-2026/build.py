#!/usr/bin/env python3
"""Gestionale Partita IVA forfettaria 2026: rigenera i file consegnati all'acquirente.

Uso (dalla cartella del prodotto o da qualunque altra):
    python3 build.py

Scrive:
    files/gestionale-forfettario-2026.xlsx           foglio vuoto, pronto da compilare
    files/gestionale-forfettario-2026-esempio.xlsx   stesso foglio con un anno di dati di esempio

I parametri fiscali arrivano da src/lib/params.js (letto con node), così il foglio usa
esattamente i valori del sito. Le fonti sono riportate nel foglio "Parametri <anno>".
Le formule replicano src/lib/forfettario.js (vedi README.md per il confronto).
L'output è deterministico: cambiano solo i metadati di data del file xlsx.
"""
import datetime as dt
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
SLUG = 'gestionale-forfettario-2026'

# --------------------------------------------------------------------------------------
# Parametri: da params.js (fonte di verità) + valori sulle scadenze verificati a parte
# --------------------------------------------------------------------------------------


def carica_parametri():
    params = (ROOT / 'src/lib/params.js').as_uri()
    forf = (ROOT / 'src/lib/forfettario.js').as_uri()
    js = (
        f"Promise.all([import('{params}'), import('{forf}')]).then(([p, f]) => console.log(JSON.stringify({{"
        "ANNO: p.ANNO, AGGIORNATO: p.AGGIORNATO, FORFETTARIO: p.FORFETTARIO, INPS: p.INPS, BOLLO: p.BOLLO,"
        " GESTIONI: f.GESTIONI })))"
    )
    try:
        out = subprocess.run(['node', '-e', js], capture_output=True, text=True, check=True).stdout
    except (OSError, subprocess.CalledProcessError) as e:
        sys.exit(f'Impossibile leggere src/lib/params.js con node: {e}')
    return json.loads(out)


# Valori che params.js non contiene (servono solo per il foglio Scadenze), con la fonte.
ACCONTI = {
    'rata': 0.5,  # forfettari: acconto in due rate uguali, art. 58 DL 124/2019
    'minimo': 51.65,  # nessun acconto se l'imposta dell'anno non supera 51,65 €
    'primaRataMinima': 103,  # sotto 103 € di prima rata l'acconto si versa tutto a novembre
}
BOLLO_SOGLIA_RINVIO = 5000  # DM 17/06/2014 art. 6 c. 2, soglia portata a 5.000 € dal DL 73/2022

# Scadenze legate all'anno 2026 (verificate a ottobre 2026, fonti nel README).
# tipo: chi le riguarda. importo: chiave della formula usata nella colonna "Importo stimato".
SCADENZE = [
    (dt.date(2026, 5, 18), 'Contributi INPS fissi 2026: 1ª rata', 'Artigiani e commercianti',
     'Sezione INPS', 'fissi',
     'Il 16 maggio 2026 era sabato. Gli importi esatti sono nel Cassetto previdenziale INPS.'),
    (dt.date(2026, 6, 1), 'Bollo fatture elettroniche, 1° trimestre 2026', 'Tutti', '2521', 'bollo1',
     'Il 31 maggio 2026 era domenica. Fino a 5.000 € il versamento si può rinviare alla scadenza '
     'del trimestre successivo (importo già spostato qui).'),
    (dt.date(2026, 7, 20), 'Imposta sostitutiva: saldo 2025 e 1° acconto 2026', 'Tutti', '1792 saldo, 1790 acconto',
     'dich2025',
     'Termine ordinario 30 giugno, prorogato al 20 luglio dal DL 89/2026 (art. 6) per forfettari e '
     'soggetti ISA. Fino al 20 agosto con la maggiorazione dello 0,80%. Acconto: 50% + 50%.'),
    (dt.date(2026, 7, 20), 'Contributi INPS a percentuale: saldo 2025 e 1° acconto 2026',
     'Gestione Separata; artigiani e commercianti sul reddito oltre il minimale', 'Sezione INPS', 'dichInps',
     'Stesse scadenze delle imposte sui redditi, proroga compresa. Le casse professionali hanno '
     'scadenze proprie.'),
    (dt.date(2026, 8, 20), 'Contributi INPS fissi 2026: 2ª rata', 'Artigiani e commercianti', 'Sezione INPS',
     'fissi', ''),
    (dt.date(2026, 9, 30), 'Bollo fatture elettroniche, 2° trimestre 2026', 'Tutti', '2522', 'bollo2',
     'Se il 1° e il 2° trimestre insieme non superano 5.000 € si può rinviare al 30 novembre.'),
    (dt.date(2026, 11, 16), 'Contributi INPS fissi 2026: 3ª rata', 'Artigiani e commercianti', 'Sezione INPS',
     'fissi', ''),
    (dt.date(2026, 11, 30), 'Imposta sostitutiva: 2° acconto 2026', 'Tutti', '1791', 'dich2025',
     'Importo calcolato nella dichiarazione 2026 (metodo storico).'),
    (dt.date(2026, 11, 30), 'Contributi INPS a percentuale: 2° acconto 2026',
     'Gestione Separata; artigiani e commercianti sul reddito oltre il minimale', 'Sezione INPS', 'dichInps', ''),
    (dt.date(2026, 11, 30), 'Bollo fatture elettroniche, 3° trimestre 2026', 'Tutti', '2523', 'bollo3',
     'Comprende i trimestri precedenti rinviati (codici 2521 e 2522).'),
    (dt.date(2027, 2, 16), 'Contributi INPS fissi 2026: 4ª rata', 'Artigiani e commercianti', 'Sezione INPS',
     'fissi', 'Si deduce dal reddito del 2027, l\'anno in cui la versi.'),
    (dt.date(2027, 3, 1), 'Bollo fatture elettroniche, 4° trimestre 2026', 'Tutti', '2524', 'bollo4',
     'Il 28 febbraio 2027 è domenica.'),
    (dt.date(2027, 6, 30), 'Imposta sostitutiva: saldo 2026', 'Tutti', '1792', 'saldo2026',
     'Imposta 2026 stimata meno gli acconti 2026 versati (Impostazioni). Se è negativo hai un credito. '
     'Oltre 100.000 € di ricavi l\'imposta sostitutiva 2026 non è dovuta. Termine ordinario: controlla le proroghe.'),
    (dt.date(2027, 6, 30), 'Imposta sostitutiva: 1° acconto 2027', 'Tutti', '1790', 'acconto1',
     'Metodo storico: 50% dell\'imposta 2026. Nessun acconto se l\'imposta non supera 51,65 €; '
     'se la prima rata è sotto 103 € si versa tutto a novembre. Oltre 85.000 € di ricavi dal 2027 passi al '
     'regime ordinario: niente acconti d\'imposta sostitutiva, il foglio non stima quelli IRPEF.'),
    (dt.date(2027, 6, 30), 'Contributi INPS a percentuale: saldo 2026 e 1° acconto 2027',
     'Gestione Separata; artigiani e commercianti sul reddito oltre il minimale', 'Sezione INPS', 'dichInps2',
     'Base di partenza: i contributi a percentuale stimati nella Dashboard.'),
    (dt.date(2027, 11, 30), 'Imposta sostitutiva: 2° acconto 2027', 'Tutti', '1791', 'acconto2', ''),
]

# --------------------------------------------------------------------------------------
# Stile
# --------------------------------------------------------------------------------------
FONT = 'Arial'
TEAL = '0F766E'
TEAL_DARK = '134E4A'
TEAL_LIGHT = 'CCFBF1'
TEAL_BG = 'F0FDFA'
INPUT = 'FEF9C3'
INPUT_EDGE = 'E5C04A'
GRID = 'D1D5DB'
TEXT = '1F2937'
MUTED = '6B7280'
RED = 'B91C1C'
RED_BG = 'FEE2E2'
AMBER = '92400E'
AMBER_BG = 'FEF3C7'
GREEN = '166534'
GREEN_BG = 'DCFCE7'

EUR = '#,##0.00 "€"'
EUR0 = '#,##0 "€"'
PCT0 = '0%'
PCT1 = '0.0%'
PCT2 = '0.00%'
DATA = 'dd/mm/yyyy'

N_FATTURE = 300
N_SPESE = 200
R0 = 5  # prima riga dati in Fatture e Spese
MESI = ['Gennaio', 'Febbraio', 'Marzo', 'Aprile', 'Maggio', 'Giugno', 'Luglio', 'Agosto',
        'Settembre', 'Ottobre', 'Novembre', 'Dicembre']
CATEGORIE_SPESE = ['Software e abbonamenti', 'Attrezzatura', 'Commercialista', 'Formazione',
                   'Affitto e coworking', 'Telefono e internet', 'Trasporti e trasferte', 'Materiali',
                   'Marketing', 'Altro']

thin = Side(style='thin', color=GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
edge = Side(style='thin', color=INPUT_EDGE)
BORDER_INPUT = Border(left=edge, right=edge, top=edge, bottom=edge)


def f(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill('solid', start_color=color, end_color=color)


def it_num(x, dec=0):
    """Numero in formato italiano per i testi (1.234,56)."""
    s = f'{x:,.{dec}f}'
    return s.replace(',', 'X').replace('.', ',').replace('X', '.')


def it_pct(x):
    v = x * 100
    return (it_num(v, 2).rstrip('0').rstrip(',') if v != int(v) else str(int(v))) + '%'


def it_data(d):
    return d.strftime('%d/%m/%Y')


def serial(d):
    return (d - dt.date(1899, 12, 30)).days


def put(ws, coord, value, *, font=None, fmt=None, bg=None, align=None, border=None, locked=None):
    c = ws[coord]
    c.value = value
    if font is not None:
        c.font = font
    if fmt is not None:
        c.number_format = fmt
    if bg is not None:
        c.fill = fill(bg)
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    if locked is not None:
        c.protection = Protection(locked=locked)
    return c


def input_cell(c, fmt=None):
    c.fill = fill(INPUT)
    c.border = BORDER_INPUT
    c.protection = Protection(locked=False)
    if fmt:
        c.number_format = fmt


WRAP = Alignment(wrap_text=True, vertical='top')
WRAP_C = Alignment(wrap_text=True, vertical='center')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center')
LEFT_C = Alignment(horizontal='left', vertical='center', wrap_text=True)


def titolo(ws, testo, sotto, last_col):
    ws.merge_cells(f'A1:{last_col}1')
    put(ws, 'A1', testo, font=f(16, True, TEAL), align=Alignment(vertical='center'))
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f'A2:{last_col}2')
    put(ws, 'A2', sotto, font=f(10, color=MUTED), align=Alignment(vertical='center', wrap_text=True))
    ws.row_dimensions[2].height = 30


def intestazioni(ws, row, labels, height=32):
    for i, lab in enumerate(labels, start=1):
        put(ws, f'{get_column_letter(i)}{row}', lab, font=f(10, True, 'FFFFFF'), bg=TEAL, align=CENTER,
            border=BORDER)
    ws.row_dimensions[row].height = height


def larghezze(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def stampa(ws, landscape=True, titoli=None, area=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    if titoli:
        ws.print_title_rows = titoli
    if area:
        ws.print_area = area
    ws.oddFooter.center.text = '&A · pagina &P di &N'
    ws.oddFooter.center.size = 8


def proteggi(ws, filtri=False):
    p = ws.protection
    p.sheet = True
    p.formatColumns = False  # consentito: allargare le colonne
    p.formatRows = False
    if filtri:
        p.autoFilter = False
        p.sort = False


def section(ws, coord_range, testo):
    first = coord_range.split(':')[0]
    ws.merge_cells(coord_range)
    put(ws, first, testo, font=f(11, True, TEAL_DARK), bg=TEAL_LIGHT, align=LEFT_C)


# --------------------------------------------------------------------------------------
# Costruzione del workbook
# --------------------------------------------------------------------------------------


class Refs(dict):
    """Riferimenti assoluti alle celle chiave, per scrivere le formule in modo leggibile."""


def crea_workbook(P):
    anno = P['ANNO']
    F = P['FORFETTARIO']
    G = P['INPS']
    B = P['BOLLO']
    wb = Workbook()
    wb._named_styles['Normal'].font = Font(name=FONT, size=10)
    wb.properties.creator = 'NettoChiaro'
    wb.properties.title = f'Gestionale Partita IVA forfettaria {anno}'
    wb.properties.subject = 'Registro fatture, stima di imposta sostitutiva e contributi, scadenze'
    fixed = dt.datetime(2026, 10, 6, 12, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed

    ws_istr = wb.active
    ws_istr.title = 'Istruzioni'
    ws_imp = wb.create_sheet('Impostazioni')
    par_name = f'Parametri {anno}'
    ws_par = wb.create_sheet(par_name)
    ws_fat = wb.create_sheet('Fatture')
    ws_spe = wb.create_sheet('Spese')
    ws_rie = wb.create_sheet('Riepilogo mensile')
    ws_das = wb.create_sheet('Dashboard')
    ws_sca = wb.create_sheet('Scadenze')
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 100

    PAR = f"'{par_name}'"
    IMP = 'Impostazioni'
    FAT = 'Fatture'
    SPE = 'Spese'
    RIE = "'Riepilogo mensile'"
    DAS = 'Dashboard'
    R = Refs()

    # ---------------------------------------------------------------- Parametri
    ws = ws_par
    larghezze(ws, [58, 16, 78])
    titolo(ws, f'Parametri {anno}',
           f'Valori usati dalle formule, con la fonte. Aggiornati al {it_data(dt.date.fromisoformat(P["AGGIORNATO"]))}. '
           'Ogni anno basta aggiornare questa tabella (vedi Istruzioni).', 'C')
    intestazioni(ws, 4, ['Parametro', 'Valore', 'Fonte'], height=22)
    ws.freeze_panes = 'A5'
    row = 5

    def par(key, label, value, fmt, fonte):
        nonlocal row
        put(ws, f'A{row}', label, font=f(), align=LEFT_C, border=BORDER)
        put(ws, f'B{row}', value, font=f(bold=True), fmt=fmt, align=RIGHT, border=BORDER)
        put(ws, f'C{row}', fonte, font=f(9, color=MUTED), align=LEFT_C, border=BORDER)
        R[key] = f'{PAR}!$B${row}'
        row += 1

    def par_section(testo):
        nonlocal row
        section(ws, f'A{row}:C{row}', testo)
        ws.row_dimensions[row].height = 20
        row += 1

    gs, art, com = G['gestioneSeparata'], G['artigiani'], G['commercianti']
    par_section('Regime forfettario')
    par('anno', 'Anno di riferimento', anno, '0', 'Anno a cui si riferiscono i valori e gli incassi')
    par('soglia', 'Soglia di ricavi per restare nel regime l\'anno dopo', F['sogliaRicavi'], EUR0,
        'L. 190/2014, art. 1 c. 54')
    par('uscita', 'Soglia di uscita immediata (già nell\'anno)', F['sogliaUscitaImmediata'], EUR0,
        'L. 190/2014, art. 1 c. 71; circolare Agenzia delle Entrate 32/E/2023')
    par('imp15', 'Imposta sostitutiva ordinaria', F['impostaOrdinaria'], PCT0, 'L. 190/2014, art. 1 c. 64')
    par('imp5', 'Imposta sostitutiva per le nuove attività (primi 5 anni)', F['impostaStartup'], PCT0,
        'L. 190/2014, art. 1 c. 65 (con i requisiti di legge)')
    par('limDip', 'Limite di reddito da lavoro dipendente o pensione (solo informativo)',
        F['limiteRedditoDipendente'], EUR0,
        'L. 190/2014, art. 1 c. 57 lett. d-ter; 35.000 € per il 2026 (L. 199/2025, art. 1 c. 27). '
        'Il foglio non lo usa nei calcoli.')
    par_section('Imposta di bollo e rivalsa')
    par('bollo', 'Imposta di bollo su fatture senza IVA', B['importo'], EUR, 'DPR 642/1972, Tariffa parte I, art. 13')
    par('sogliaBollo', 'Il bollo è dovuto se l\'importo della fattura supera', B['soglia'], EUR,
        'DPR 642/1972, Tariffa parte I, art. 13 e note')
    par('rinvioBollo', 'Bollo fatture elettroniche: si può rinviare al trimestre dopo fino a', BOLLO_SOGLIA_RINVIO,
        EUR0, 'DM 17/06/2014, art. 6 c. 2 (soglia di 5.000 € dal 2023, DL 73/2022)')
    par('rivalsa', 'Rivalsa INPS facoltativa in fattura (solo Gestione Separata)', gs['rivalsa'], PCT0,
        'L. 662/1996, art. 1 c. 212')
    par_section('INPS Gestione Separata')
    par('gs', 'Aliquota professionisti senza altra previdenza obbligatoria', gs['aliquota'], PCT2,
        'Circolare INPS 8/2026 (25% IVS + 0,72% + 0,35% ISCRO)')
    par('gsPens', 'Aliquota pensionati o iscritti ad altra previdenza obbligatoria', gs['aliquotaPensionati'], PCT2,
        'Circolare INPS 8/2026')
    par('gsMax', 'Massimale di reddito', gs['massimale'], EUR0, 'Circolare INPS 8/2026')
    for key, nome, p in (('art', 'Artigiani', art), ('com', 'Commercianti', com)):
        par_section(f'INPS {nome}')
        par(f'{key}Min', 'Minimale di reddito', p['minimale'], EUR0, 'Circolare INPS 14/2026')
        par(f'{key}Fisso', 'Contributo fisso annuo sul minimale (quota maternità compresa)', p['contributoFisso'],
            EUR, 'Circolare INPS 14/2026')
        par(f'{key}Mat', 'Quota maternità compresa nel contributo fisso (non si riduce)', p['maternita'], EUR,
            'Circolare INPS 14/2026 (0,62 € al mese)')
        par(f'{key}Aliq', 'Aliquota fino alla soglia', p['aliquota'], PCT2, 'Circolare INPS 14/2026')
        par(f'{key}Soglia', 'Soglia oltre la quale l\'aliquota sale di un punto', p['sogliaAggiuntiva'], EUR0,
            'Circolare INPS 14/2026 (prima fascia di retribuzione pensionabile)')
        par(f'{key}Oltre', 'Aliquota oltre la soglia', p['aliquotaOltre'], PCT2, 'Circolare INPS 14/2026')
        par(f'{key}Max', 'Massimale di reddito', p['massimale'], EUR0, 'Circolare INPS 14/2026')
    par_section('Riduzione contributiva')
    par('rid', 'Riduzione per forfettari iscritti ad artigiani o commercianti (su domanda)', G['riduzioneForfettari'],
        PCT0, 'L. 190/2014, art. 1 c. 77')
    par_section('Acconti dell\'imposta sostitutiva (foglio Scadenze)')
    par('accRata', 'Ogni rata di acconto (due rate uguali)', ACCONTI['rata'], PCT0, 'DL 124/2019, art. 58')
    par('accMin', 'Nessun acconto se l\'imposta dell\'anno non supera', ACCONTI['minimo'], EUR,
        'Istruzioni del modello Redditi PF')
    par('accPrima', 'Acconto tutto a novembre se la prima rata è inferiore a', ACCONTI['primaRataMinima'], EUR,
        'Istruzioni del modello Redditi PF')

    row += 1
    section(ws, f'A{row}:C{row}', 'Coefficienti di redditività (Allegato 4, L. 190/2014)')
    row += 1
    for col, lab in zip('ABC', ['Gruppo di attività (menu in Impostazioni)', 'Coefficiente', 'Codici ATECO 2007']):
        put(ws, f'{col}{row}', lab, font=f(10, True, 'FFFFFF'), bg=TEAL, align=CENTER, border=BORDER)
    row += 1
    c_first = row
    for c in F['coefficienti']:
        put(ws, f'A{row}', c['label'], font=f(), align=LEFT_C, border=BORDER)
        put(ws, f'B{row}', c['coeff'], font=f(bold=True), fmt=PCT0, align=RIGHT, border=BORDER)
        put(ws, f'C{row}', c['ateco'], font=f(9, color=MUTED), align=LEFT_C, border=BORDER)
        ws.row_dimensions[row].height = 28 if len(c['label']) > 60 else 16
        row += 1
    c_last = row - 1
    R['coeffLabels'] = f'{PAR}!$A${c_first}:$A${c_last}'
    R['coeffValues'] = f'{PAR}!$B${c_first}:$B${c_last}'
    put(ws, f'A{row}', 'Il coefficiente si determina con il codice ATECO 2007 corrispondente anche dopo l\'adozione '
        'di ATECO 2025 (1/4/2025), finché non esce la nuova tabella.', font=f(9, color=MUTED), align=WRAP)
    ws.merge_cells(f'A{row}:C{row}')
    ws.row_dimensions[row].height = 26
    row += 2

    section(ws, f'A{row}:C{row}', 'Elenchi dei menu a tendina')
    row += 1
    for col, lab in zip('ABC', ['Previdenza', 'Codice', 'Come si calcola']):
        put(ws, f'{col}{row}', lab, font=f(10, True, 'FFFFFF'), bg=TEAL, align=CENTER, border=BORDER)
    row += 1
    g_first = row
    spieg = {
        'separata': 'Aliquota Gestione Separata sul reddito, fino al massimale',
        'separata-pensionati': 'Aliquota ridotta sul reddito, fino al massimale',
        'artigiani': 'Contributo fisso sul minimale + aliquota sul reddito oltre il minimale',
        'commercianti': 'Contributo fisso sul minimale + aliquota sul reddito oltre il minimale',
        'cassa': 'Aliquota inserita in Impostazioni sul reddito (minimi e contributo integrativo esclusi)',
    }
    for i, g in enumerate(P['GESTIONI'], start=1):
        put(ws, f'A{row}', g['label'], font=f(), align=LEFT_C, border=BORDER)
        put(ws, f'B{row}', i, font=f(), align=RIGHT, border=BORDER)
        put(ws, f'C{row}', spieg[g['id']], font=f(9, color=MUTED), align=LEFT_C, border=BORDER)
        row += 1
    R['gestioni'] = f'{PAR}!$A${g_first}:$A${row - 1}'
    ids = [g['id'] for g in P['GESTIONI']]
    assert ids == ['separata', 'separata-pensionati', 'artigiani', 'commercianti', 'cassa'], ids
    row += 1
    put(ws, f'A{row}', 'Imposta sostitutiva', font=f(10, True, 'FFFFFF'), bg=TEAL, align=CENTER, border=BORDER)
    row += 1
    imposta_labels = [f'{it_pct(F["impostaOrdinaria"])} (aliquota ordinaria)',
                      f'{it_pct(F["impostaStartup"])} (nuova attività, primi 5 anni)']
    i_first = row
    for lab in imposta_labels:
        put(ws, f'A{row}', lab, font=f(), border=BORDER)
        row += 1
    R['impostaList'] = f'{PAR}!$A${i_first}:$A${row - 1}'
    R['impostaStartupLabel'] = f'{PAR}!$A${i_first + 1}'
    row += 1
    put(ws, f'A{row}', 'Categorie di spesa', font=f(10, True, 'FFFFFF'), bg=TEAL, align=CENTER, border=BORDER)
    row += 1
    s_first = row
    for lab in CATEGORIE_SPESE:
        put(ws, f'A{row}', lab, font=f(), border=BORDER)
        row += 1
    R['categorie'] = f'{PAR}!$A${s_first}:$A${row - 1}'
    stampa(ws, landscape=False)
    proteggi(ws)

    # ---------------------------------------------------------------- Impostazioni
    ws = ws_imp
    larghezze(ws, [46, 56, 70])
    titolo(ws, 'Impostazioni', 'Compila le celle gialle. Tutto il resto del file si aggiorna da solo.', 'C')
    intestazioni(ws, 4, ['Voce', 'La tua scelta', 'Spiegazione'], height=22)
    voci = [
        (5, 'Nome o attività (facoltativo)', None,
         'Compare in alto nella Dashboard.'),
        (6, 'Attività (gruppo di codici ATECO)', None,
         f'Scegli il gruppo del tuo codice ATECO: determina il coefficiente di redditività. '
         f'Elenco completo nel foglio {par_name}.'),
        (7, 'Previdenza', None,
         'Gestione Separata se sei un professionista senza cassa. Artigiani o Commercianti se sei iscritto '
         'a una di queste gestioni INPS. Cassa professionale: inserisci l\'aliquota nella riga sotto.'),
        (8, 'Aliquota della cassa professionale (solo se hai una cassa)', PCT2,
         'Aliquota del contributo soggettivo della tua cassa (es. 14,5%). Il foglio non calcola minimi '
         'e contributo integrativo della cassa.'),
        (9, 'Imposta sostitutiva', None,
         '5% per i primi 5 anni di una nuova attività, se hai i requisiti di legge. Altrimenti 15%.'),
        (10, 'Riduzione contributiva del 35%', None,
         'Solo artigiani e commercianti che l\'hanno chiesta all\'INPS. Non si applica alla quota maternità.'),
        (11, 'Il bollo addebitato al cliente conta nei ricavi', None,
         'Sì (prudente): secondo l\'Agenzia delle Entrate (risposta 428/2022) il bollo riaddebitato fa parte '
         'del compenso. Dopo il D.Lgs. 192/2024 alcune interpretazioni lo escludono: chiedi al tuo '
         'commercialista. Vale per tasse e soglie.'),
        (12, f'Contributi previdenziali versati nel {anno} (facoltativo)', EUR,
         'Lascia vuoto per la stima "a regime" (si deducono i contributi calcolati sull\'anno). Se li conosci, '
         f'inserisci i contributi che versi davvero nel {anno} (saldo {anno - 1}, acconti {anno}, rate fisse '
         f'pagate nel {anno}): l\'imposta si calcola su quelli. Primo anno di attività: scrivi 0.'),
        (13, f'Acconti d\'imposta sostitutiva versati per il {anno} (facoltativo)', EUR,
         f'Somma dei due acconti {anno} (luglio e novembre). Serve al foglio Scadenze per stimare il saldo '
         f'da versare nel {anno + 1}.'),
    ]
    for r_, lab, fmt, spiega in voci:
        put(ws, f'A{r_}', lab, font=f(bold=True), align=LEFT_C, border=BORDER)
        c = ws[f'B{r_}']
        input_cell(c, fmt)
        c.font = f()
        c.alignment = LEFT_C
        put(ws, f'C{r_}', spiega, font=f(9, color=MUTED), align=WRAP_C, border=BORDER)
        ws.row_dimensions[r_].height = 44 if len(spiega) > 140 else (32 if len(spiega) > 70 else 22)
    ws.row_dimensions[11].height = 56
    ws.row_dimensions[12].height = 56

    dv = DataValidation(type='list', formula1=f'={R["coeffLabels"]}', allow_blank=True)
    dv.error, dv.errorTitle = 'Scegli una voce dell\'elenco.', 'Attività'
    ws.add_data_validation(dv)
    dv.add('B6')
    dv = DataValidation(type='list', formula1=f'={R["gestioni"]}', allow_blank=True)
    dv.error, dv.errorTitle = 'Scegli una voce dell\'elenco.', 'Previdenza'
    ws.add_data_validation(dv)
    dv.add('B7')
    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='1', allow_blank=True)
    dv.error, dv.errorTitle = 'Scrivi l\'aliquota con il simbolo %, ad esempio 14,5%.', 'Aliquota'
    ws.add_data_validation(dv)
    dv.add('B8')
    dv = DataValidation(type='list', formula1=f'={R["impostaList"]}', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('B9')
    dv = DataValidation(type='list', formula1='"Sì,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('B10')
    dv.add('B11')
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    dv.error, dv.errorTitle = 'Inserisci un importo in euro, zero o positivo.', 'Importo'
    ws.add_data_validation(dv)
    dv.add('B12')
    dv.add('B13')
    R.update(nome=f'{IMP}!$B$5', attivita=f'{IMP}!$B$6', previdenza=f'{IMP}!$B$7', aliqCassa=f'{IMP}!$B$8',
             imposta=f'{IMP}!$B$9', riduzione=f'{IMP}!$B$10', bolloScelta=f'{IMP}!$B$11',
             versati=f'{IMP}!$B$12', acconti=f'{IMP}!$B$13')

    section(ws, 'A15:C15', 'Valori usati nei calcoli (si aggiornano da soli, non modificarli)')
    calc = [
        ('coeff', 'Coefficiente di redditività',
         f'=IFERROR(INDEX({R["coeffValues"]},MATCH({R["attivita"]},{R["coeffLabels"]},0)),0)', PCT0,
         'Dal gruppo di attività scelto.'),
        ('code', 'Codice della previdenza (1-5)',
         f'=IFERROR(MATCH({R["previdenza"]},{R["gestioni"]},0),0)', '0',
         f'1-2 Gestione Separata, 3 Artigiani, 4 Commercianti, 5 Cassa; 0 = non scelta. Elenco in {par_name}.'),
        ('aliqImp', 'Aliquota dell\'imposta sostitutiva',
         f'=IF({R["imposta"]}={R["impostaStartupLabel"]},{R["imp5"]},{R["imp15"]})', PCT0, ''),
        ('aliqC', 'Aliquota contributiva',
         None, PCT2, 'Per artigiani e commercianti è l\'aliquota fino alla soglia.'),
        ('ridSi', 'Riduzione del 35% applicata', None, None, 'Solo artigiani e commercianti.'),
        ('fattore', 'Quota dei contributi dovuta dopo la riduzione', None, PCT0, ''),
        ('minimale', 'Minimale di reddito (artigiani e commercianti)', None, EUR0, ''),
        ('fissoPieno', 'Contributo fisso annuo pieno', None, EUR, ''),
        ('maternita', 'Quota maternità (non ridotta)', None, EUR, ''),
        ('fissi', 'Contributi fissi annui dovuti', None, EUR, '(fisso − maternità) × quota dovuta + maternità.'),
        ('sogliaAgg', 'Soglia per l\'aliquota aggiuntiva', None, EUR0, ''),
        ('aliqOltre', 'Aliquota oltre la soglia', None, PCT2, ''),
        ('massimale', 'Massimale di reddito', None, EUR0, 'Le casse professionali non hanno massimale qui.'),
        ('bolloRicavo', 'Bollo addebitato incluso nei ricavi', None, None, 'Vuoto vale Sì.'),
    ]
    rr = 16
    for key, lab, formula, fmt, nota in calc:
        R[key] = f'{IMP}!$B${rr}'
        put(ws, f'A{rr}', lab, font=f(), align=LEFT_C, border=BORDER)
        put(ws, f'B{rr}', formula, font=f(bold=True), fmt=fmt, bg=TEAL_BG, align=LEFT_C, border=BORDER)
        put(ws, f'C{rr}', nota, font=f(9, color=MUTED), align=LEFT_C, border=BORDER)
        rr += 1
    c = R['code']
    form = {
        'aliqC': f'=CHOOSE({c}+1,0,{R["gs"]},{R["gsPens"]},{R["artAliq"]},{R["comAliq"]},N({R["aliqCassa"]}))',
        'ridSi': f'=IF(AND({R["riduzione"]}="Sì",OR({c}=3,{c}=4)),"Sì","No")',
        'fattore': f'=IF({R["ridSi"]}="Sì",1-{R["rid"]},1)',
        'minimale': f'=IF({c}=4,{R["comMin"]},{R["artMin"]})',
        'fissoPieno': f'=IF({c}=3,{R["artFisso"]},IF({c}=4,{R["comFisso"]},0))',
        'maternita': f'=IF({c}=3,{R["artMat"]},IF({c}=4,{R["comMat"]},0))',
        'fissi': f'=({R["fissoPieno"]}-{R["maternita"]})*{R["fattore"]}+{R["maternita"]}',
        'sogliaAgg': f'=IF({c}=4,{R["comSoglia"]},{R["artSoglia"]})',
        'aliqOltre': f'=IF({c}=4,{R["comOltre"]},{R["artOltre"]})',
        'massimale': f'=IF({c}<=2,{R["gsMax"]},IF({c}=4,{R["comMax"]},{R["artMax"]}))',
        'bolloRicavo': f'=IF({R["bolloScelta"]}="No","No","Sì")',
    }
    for key, formula in form.items():
        ws[R[key].split('!')[1].replace('$', '')].value = formula
    stampa(ws, landscape=True)
    proteggi(ws)

    def variabili(red):
        """Contributi a percentuale sul reddito `red` (stessa logica di contributiForfettario)."""
        c, mx, al, sg, mn, ol, fr = (R[k] for k in ('code', 'massimale', 'aliqC', 'sogliaAgg', 'minimale',
                                                     'aliqOltre', 'fattore'))
        base = f'MIN(MAX({red},0),{mx})'
        return (f'IF(OR({c}=1,{c}=2),{base}*{al},IF({c}=5,MAX({red},0)*{al},IF(OR({c}=3,{c}=4),'
                f'(MAX(0,MIN({base},{sg})-{mn})*{al}+MAX(0,{base}-{sg})*{ol})*{fr},0)))')

    # ---------------------------------------------------------------- Fatture
    ws = ws_fat
    last = R0 + N_FATTURE - 1
    larghezze(ws, [9, 12, 12, 28, 34, 13, 10, 11, 13, 9, 11, 13, 14, 21, 22])
    titolo(ws, 'Fatture', 'Una riga per fattura. Celle gialle da compilare; quando il cliente paga, scrivi la data di '
           'incasso: conta quella (principio di cassa).', 'O')
    put(ws, 'L3', f'Incassato nel {anno}', font=f(9, True, MUTED), align=RIGHT)
    put(ws, 'M3', f'={RIE}!$B$17', font=f(10, True, TEAL), fmt=EUR, align=RIGHT)
    put(ws, 'N3', '=COUNTIF($N$5:$N$%d,"Da incassare")&" da incassare"' % last, font=f(9, color=MUTED),
        align=RIGHT)
    intestazioni(ws, 4, ['N.', 'Data emissione', 'Data incasso', 'Cliente', 'Descrizione', 'Compenso (€)',
                         'Rivalsa INPS 4%', 'Rivalsa (€)', 'Imponibile (€)', 'Bollo (€)',
                         'Bollo al cliente', 'Totale fattura (€)', 'Ricavo che conta (€)', 'Stato', 'Note'],
                 height=34)
    ws.freeze_panes = 'B5'
    ws.auto_filter.ref = f'A4:O{last}'
    for r_ in range(R0, last + 1):
        for col, fmt in (('A', None), ('B', DATA), ('C', DATA), ('D', None), ('E', None), ('F', EUR),
                         ('G', None), ('K', None), ('O', None)):
            cell = ws[f'{col}{r_}']
            input_cell(cell, fmt)
            cell.font = f()
            if col in 'GK':
                cell.alignment = Alignment(horizontal='center')
        # Compenso scritto come testo (es. incollato da un altro file): la riga resta vuota nei calcoli e
        # lo Stato lo segnala, così un errore non si propaga a tutti i totali. Lo stesso per le date.
        ws[f'H{r_}'] = (f'=IF(ISNUMBER(F{r_}),IF(OR(G{r_}="Sì",G{r_}="Si"),ROUND(F{r_}*{R["rivalsa"]},2),0),"")')
        ws[f'I{r_}'] = f'=IF(ISNUMBER(F{r_}),F{r_}+H{r_},"")'
        ws[f'J{r_}'] = f'=IF(ISNUMBER(F{r_}),IF(I{r_}>{R["sogliaBollo"]},{R["bollo"]},0),"")'
        ws[f'L{r_}'] = f'=IF(ISNUMBER(F{r_}),I{r_}+IF(K{r_}="No",0,J{r_}),"")'
        ws[f'M{r_}'] = (f'=IF(ISNUMBER(F{r_}),I{r_}+IF(AND(K{r_}<>"No",{R["bolloRicavo"]}="Sì"),J{r_},0),"")')
        ws[f'N{r_}'] = (f'=IF(F{r_}="","",IF(NOT(ISNUMBER(F{r_})),"Compenso non valido",IF(C{r_}="","Da incassare",'
                        f'IF(NOT(ISNUMBER(C{r_})),"Data non valida",IF(YEAR(C{r_})={R["anno"]},"Incassata",'
                        f'"Incassata nel "&YEAR(C{r_}))))))')
        for col in 'HIJLMN':
            cell = ws[f'{col}{r_}']
            cell.font = f(bold=(col in 'LM'))
            cell.border = BORDER
            if col != 'N':
                cell.number_format = EUR
        ws[f'N{r_}'].alignment = Alignment(horizontal='center')
    dv = DataValidation(type='date', operator='between', formula1=str(serial(dt.date(2015, 1, 1))),
                        formula2=str(serial(dt.date(2040, 12, 31))), allow_blank=True)
    dv.error, dv.errorTitle = 'Inserisci una data nel formato gg/mm/aaaa.', 'Data'
    dv.prompt, dv.promptTitle = 'Formato gg/mm/aaaa', 'Data'
    ws.add_data_validation(dv)
    dv.add(f'B{R0}:C{last}')
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    dv.error, dv.errorTitle = 'Inserisci il compenso in euro (senza rivalsa e bollo).', 'Compenso'
    ws.add_data_validation(dv)
    dv.add(f'F{R0}:F{last}')
    dv = DataValidation(type='list', formula1='"Sì,No"', allow_blank=True)
    dv.prompt, dv.promptTitle = 'Vuoto = No per la rivalsa, Sì per il bollo', 'Sì / No'
    ws.add_data_validation(dv)
    dv.add(f'G{R0}:G{last}')
    dv.add(f'K{R0}:K{last}')
    rng = f'N{R0}:N{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{R0}="Da incassare"'], font=Font(color=AMBER, bold=True),
                                                   fill=fill(AMBER_BG)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'N{R0}="Incassata"'], font=Font(color=GREEN)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(N{R0},12)="Incassata ne"'],
                                                   font=Font(color=MUTED, italic=True)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'RIGHT(N{R0},10)="non valido"'],
                                                   font=Font(color=RED, bold=True), fill=fill(RED_BG)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'RIGHT(N{R0},10)="non valida"'],
                                                   font=Font(color=RED, bold=True), fill=fill(RED_BG)))
    ws.conditional_formatting.add(f'C{R0}:C{last}', FormulaRule(
        formula=[f'AND(C{R0}<>"",B{R0}<>"",C{R0}<B{R0})'], font=Font(color=RED, bold=True), fill=fill(RED_BG)))
    stampa(ws, titoli='4:4', area=f'A1:O{last}')
    proteggi(ws, filtri=True)
    FR = {k: f'{FAT}!${k}${R0}:${k}${last}' for k in 'ABCFGHIJKLMN'}

    # ---------------------------------------------------------------- Spese
    ws = ws_spe
    slast = R0 + N_SPESE - 1
    larghezze(ws, [14, 40, 24, 14, 40])
    titolo(ws, 'Spese (facoltativo)', 'Nel forfettario i costi non si deducono: questo foglio serve solo a vedere '
           'il netto vero. Non inserire qui tasse e contributi, li calcola il file.', 'E')
    put(ws, 'C3', f'Spese nel {anno}', font=f(9, True, MUTED), align=RIGHT)
    put(ws, 'D3', f'={RIE}!$L$17', font=f(10, True, TEAL), fmt=EUR, align=RIGHT)
    intestazioni(ws, 4, ['Data pagamento', 'Descrizione', 'Categoria', 'Importo (€)', 'Note'], height=28)
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f'A4:E{slast}'
    for r_ in range(R0, slast + 1):
        for col, fmt in (('A', DATA), ('B', None), ('C', None), ('D', EUR), ('E', None)):
            cell = ws[f'{col}{r_}']
            input_cell(cell, fmt)
            cell.font = f()
    dv = DataValidation(type='date', operator='between', formula1=str(serial(dt.date(2015, 1, 1))),
                        formula2=str(serial(dt.date(2040, 12, 31))), allow_blank=True)
    dv.error, dv.errorTitle = 'Inserisci una data nel formato gg/mm/aaaa.', 'Data'
    ws.add_data_validation(dv)
    dv.add(f'A{R0}:A{slast}')
    dv = DataValidation(type='list', formula1=f'={R["categorie"]}', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'C{R0}:C{slast}')
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
    dv.error, dv.errorTitle = 'Inserisci un importo in euro.', 'Importo'
    ws.add_data_validation(dv)
    dv.add(f'D{R0}:D{slast}')
    stampa(ws, landscape=False, titoli='4:4', area=f'A1:E{slast}')
    proteggi(ws, filtri=True)
    SR = {k: f'{SPE}!${k}${R0}:${k}${slast}' for k in 'AD'}

    # ---------------------------------------------------------------- Riepilogo mensile
    ws = ws_rie
    larghezze(ws, [12, 14, 10, 15, 10, 10, 15, 15, 15, 15, 15, 13, 15, 15])
    titolo(ws, f'Riepilogo mensile {anno}',
           'Incassi per data di incasso. Le colonne "da inizio anno" stimano tasse e contributi come se l\'anno '
           'finisse quel mese; "Da accantonare nel mese" è la differenza con il mese prima.', 'N')
    intestazioni(ws, 4, ['Mese', 'Incassato nel mese', 'Fatture incassate', 'Incassato da inizio anno',
                         f'% soglia {it_num(F["sogliaRicavi"])} €', f'% soglia {it_num(F["sogliaUscitaImmediata"])} €',
                         'Reddito da inizio anno', 'Contributi stimati da inizio anno',
                         'Imposta stimata da inizio anno', 'Tasse e contributi da inizio anno',
                         'Da accantonare nel mese', 'Spese del mese', 'Netto stimato del mese',
                         'Netto da inizio anno'], height=46)
    ws.freeze_panes = 'B5'
    an = R['anno']
    for m in range(1, 13):
        r_ = 4 + m
        da, a = f'DATE({an},{m},1)', f'DATE({an},{m + 1},1)'
        put(ws, f'A{r_}', MESI[m - 1], font=f(bold=True), border=BORDER)
        ws[f'B{r_}'] = f'=SUMIFS({FR["M"]},{FR["C"]},">="&{da},{FR["C"]},"<"&{a})'
        ws[f'C{r_}'] = f'=COUNTIFS({FR["C"]},">="&{da},{FR["C"]},"<"&{a},{FR["N"]},"Incassata")'
        ws[f'D{r_}'] = f'=B{r_}' if m == 1 else f'=D{r_ - 1}+B{r_}'
        ws[f'E{r_}'] = f'=D{r_}/{R["soglia"]}'
        ws[f'F{r_}'] = f'=D{r_}/{R["uscita"]}'
        ws[f'G{r_}'] = f'=D{r_}*{R["coeff"]}'
        ws[f'H{r_}'] = f'={R["fissi"]}+{variabili(f"G{r_}")}'
        ws[f'I{r_}'] = f'=MAX(0,G{r_}-IF(ISNUMBER({R["versati"]}),{R["versati"]},H{r_}))*{R["aliqImp"]}'
        ws[f'J{r_}'] = f'=H{r_}+I{r_}'
        ws[f'K{r_}'] = (f'=J{r_}-{R["fissi"]}+{R["fissi"]}/12' if m == 1
                        else f'=J{r_}-J{r_ - 1}+{R["fissi"]}/12')
        ws[f'L{r_}'] = f'=SUMIFS({SR["D"]},{SR["A"]},">="&{da},{SR["A"]},"<"&{a})'
        ws[f'M{r_}'] = f'=B{r_}-K{r_}-L{r_}'
        ws[f'N{r_}'] = f'=M{r_}' if m == 1 else f'=N{r_ - 1}+M{r_}'
        for col in 'BCDEFGHIJKLMN':
            cell = ws[f'{col}{r_}']
            cell.font = f(bold=(col in 'BK'))
            cell.border = BORDER
            cell.number_format = '0' if col == 'C' else (PCT1 if col in 'EF' else EUR)
            if col in 'DGHIJN':
                cell.fill = fill('F9FAFB')
        ws[f'K{r_}'].fill = fill(TEAL_BG)
    put(ws, 'A17', 'Totale', font=f(bold=True, color='FFFFFF'), bg=TEAL_DARK, border=BORDER)
    tot = {'B': '=SUM(B5:B16)', 'C': '=SUM(C5:C16)', 'D': '=D16', 'E': '=E16', 'F': '=F16', 'G': '=G16',
           'H': '=H16', 'I': '=I16', 'J': '=J16', 'K': '=SUM(K5:K16)', 'L': '=SUM(L5:L16)',
           'M': '=SUM(M5:M16)', 'N': '=N16'}
    for col, formula in tot.items():
        put(ws, f'{col}17', formula, font=f(bold=True, color='FFFFFF'), bg=TEAL_DARK, border=BORDER,
            fmt='0' if col == 'C' else (PCT1 if col in 'EF' else EUR))
    note = [
        'Come leggere il riepilogo',
        '• Incassato: somma della colonna "Ricavo che conta" delle fatture con data di incasso nel mese.',
        '• Tasse e contributi da inizio anno: stessa formula della Dashboard applicata all\'incassato fino a fine mese. '
        'A dicembre coincide con il totale dell\'anno.',
        '• Da accantonare nel mese: quanto mettere da parte di quello che hai incassato nel mese. I contributi fissi '
        'di artigiani e commercianti sono divisi in 12 quote uguali.',
        '• Netto stimato: incassato meno quota da accantonare e spese del mese.',
    ]
    for i, t in enumerate(note):
        r_ = 19 + i
        ws.merge_cells(f'A{r_}:N{r_}')
        put(ws, f'A{r_}', t, font=f(10, True, TEAL_DARK) if i == 0 else f(9, color=MUTED), align=LEFT_C)
    # Solo riferimenti allo stesso foglio: Google Sheets non accetta altri fogli nella formattazione condizionale.
    ws.conditional_formatting.add('E5:E16', FormulaRule(formula=['E5>1'],
                                                        font=Font(color=AMBER, bold=True), fill=fill(AMBER_BG)))
    ws.conditional_formatting.add('F5:F16', FormulaRule(formula=['F5>1'],
                                                        font=Font(color=RED, bold=True), fill=fill(RED_BG)))
    stampa(ws, area='A1:N23')
    proteggi(ws)

    # ---------------------------------------------------------------- Dashboard
    ws = ws_das
    larghezze(ws, [2] + [13] * 12 + [3, 11, 14, 14, 14, 14])
    ws.merge_cells('B1:M1')
    put(ws, 'B1', f'=IF({R["nome"]}="","Dashboard {anno}","Dashboard {anno} · "&{R["nome"]})',
        font=f(18, True, TEAL), align=Alignment(vertical='center'))
    ws.row_dimensions[1].height = 34
    ws.merge_cells('B2:M2')
    put(ws, 'B2', (f'=IF({R["coeff"]}=0,"Scegli attività e previdenza in Impostazioni.",'
                   f'"Coefficiente "&ROUND({R["coeff"]}*100,0)&"% · "&{R["previdenza"]}&" · imposta "&'
                   f'ROUND({R["aliqImp"]}*100,0)&"%")'),
        font=f(10, color=MUTED), align=Alignment(vertical='center'))
    RT = f'{RIE}!'
    inc = f'{RT}$B$17'
    tasse = f'{RT}$J$16'
    tiles = [
        ('B', 'D', f'Incassato nel {anno}', f'={inc}', EUR0, f'={RT}$C$17&" fatture incassate"'),
        ('E', 'G', 'Tasse e contributi stimati', f'={tasse}', EUR0,
         f'=IF({inc}>0,"il "&FIXED({tasse}/{inc}*100,1)&"% di quanto incassi","")'),
        ('H', 'J', 'Netto stimato', f'={inc}-{tasse}', EUR0, f'="circa "&FIXED(({inc}-{tasse})/12,0)&" € al mese"'),
        ('K', 'M', f'Soglia di {it_num(F["sogliaRicavi"])} €', f'={inc}/{R["soglia"]}', PCT0,
         f'=IF({inc}>{R["soglia"]},"soglia superata","mancano "&FIXED({R["soglia"]}-{inc},0)&" €")'),
    ]
    for c1, c2, lab, val, fmt, nota in tiles:
        for r_ in (4, 5, 6):
            ws.merge_cells(f'{c1}{r_}:{c2}{r_}')
            for col in range(ord(c1), ord(c2) + 1):
                ws[f'{chr(col)}{r_}'].fill = fill(TEAL_BG)
        put(ws, f'{c1}4', lab, font=f(10, True, TEAL_DARK), align=Alignment(horizontal='left', vertical='bottom',
                                                                              indent=1))
        put(ws, f'{c1}5', val, font=f(22, True, TEXT), fmt=fmt, align=Alignment(horizontal='left', vertical='center',
                                                                                  indent=1))
        put(ws, f'{c1}6', nota, font=f(9, color=MUTED), align=Alignment(horizontal='left', vertical='top', indent=1))
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 36
    ws.row_dimensions[6].height = 20

    ws.merge_cells('B8:M8')
    s85, s100 = it_num(F['sogliaRicavi']), it_num(F['sogliaUscitaImmediata'])
    msg = (f'=IF({inc}>{R["uscita"]},"Oltre {s100} € di ricavi esci dal forfettario già quest\'anno: l\'IVA si applica '
           f'dall\'operazione che fa superare la soglia e il reddito dell\'intero anno si tassa con IRPEF ordinaria e '
           f'costi reali. Imposta e contributi di questo file sono solo indicativi: l\'imposta sostitutiva non è '
           f'dovuta.",IF({inc}>{R["soglia"]},"Oltre {s85} € di ricavi resti forfettario per quest\'anno, ma dal '
           f'prossimo passi al regime ordinario.","Sotto {s85} € di ricavi: nessun limite superato."))')
    put(ws, 'B8', msg, font=f(10, True), align=Alignment(wrap_text=True, vertical='center', indent=1))
    ws.row_dimensions[8].height = 44
    # L11 e L14 (sotto) = incassato / soglia: la regola usa solo celle di questo foglio.
    for rule_formula, colr, bgc in (('$L$14>1', RED, RED_BG), ('$L$11>1', AMBER, AMBER_BG),
                                    ('$L$11<=1', GREEN, GREEN_BG)):
        ws.conditional_formatting.add('B8:M8', FormulaRule(formula=[rule_formula], font=Font(color=colr, bold=True),
                                                           fill=fill(bgc), stopIfTrue=True))

    # Dettaglio del calcolo (B..G)
    section(ws, 'B10:G10', 'Dettaglio del calcolo (stessa logica del calcolatore NettoChiaro)')
    ded = f'IF(ISNUMBER({R["versati"]}),{R["versati"]},{RT}$H$16)'
    dettaglio = [
        ('Incassato nell\'anno (per data di incasso)', f'={inc}', EUR),
        ('Coefficiente di redditività', f'={R["coeff"]}', PCT0),
        ('Reddito imponibile lordo', f'={RT}$G$16', EUR),
        ('Contributi fissi (artigiani e commercianti)', f'={R["fissi"]}', EUR),
        ('Contributi a percentuale', f'={RT}$H$16-{R["fissi"]}', EUR),
        ('Contributi previdenziali totali', f'={RT}$H$16', EUR),
        ('Contributi dedotti dal reddito', f'={ded}', EUR),
        ('Reddito su cui si paga l\'imposta', f'=MAX(0,{RT}$G$16-{ded})', EUR),
        ('Aliquota dell\'imposta sostitutiva', f'={R["aliqImp"]}', PCT0),
        ('Imposta sostitutiva stimata', f'={RT}$I$16', EUR),
        ('Totale tasse e contributi', f'={tasse}', EUR),
        ('Incidenza sull\'incassato', f'=IF({inc}>0,{tasse}/{inc},0)', PCT1),
        ('Netto stimato (incassato − tasse e contributi)', f'={inc}-{tasse}', EUR),
        ('Spese registrate (facoltative)', f'={RT}$L$17', EUR),
        ('Netto dopo le spese', f'={inc}-{tasse}-{RT}$L$17', EUR),
        ('Netto medio mensile dopo le spese', f'=({inc}-{tasse}-{RT}$L$17)/12', EUR),
    ]
    DET = {}
    for i, (lab, formula, fmt) in enumerate(dettaglio):
        r_ = 11 + i
        ws.merge_cells(f'B{r_}:E{r_}')
        ws.merge_cells(f'F{r_}:G{r_}')
        strong = lab.startswith(('Totale', 'Netto stimato', 'Imposta', 'Contributi previdenziali'))
        put(ws, f'B{r_}', lab, font=f(10, strong), align=LEFT_C)
        put(ws, f'F{r_}', formula, font=f(10, True, TEAL_DARK if strong else TEXT), fmt=fmt, align=RIGHT)
        for col in 'BCDEFG':
            ws[f'{col}{r_}'].border = Border(bottom=thin)
        DET[lab] = f'F{r_}'
    r_ = 11 + len(dettaglio)
    ws.merge_cells(f'B{r_}:G{r_ + 1}')
    put(ws, f'B{r_}', f'=IF(NOT(ISNUMBER({R["versati"]})),"Contributi dedotti: stima a regime (i contributi calcolati sull\'anno). '
        f'Se conosci quelli versati nel {anno}, inseriscili in Impostazioni.","Contributi dedotti: quelli versati nel '
        f'{anno} indicati in Impostazioni.")', font=f(9, color=MUTED), align=WRAP)
    cella = DET['Imposta sostitutiva stimata']
    R['dashImposta'] = f'{DAS}!${cella[0]}${cella[1:]}'

    # Soglie (I..M)
    section(ws, 'I10:M10', 'Soglie del regime')
    barra = lambda s: (f'=REPT("█",ROUND(MIN(1,MAX(0,{inc}/{s}))*24,0))&'
                       f'REPT("░",24-ROUND(MIN(1,MAX(0,{inc}/{s}))*24,0))')
    soglie = [
        (11, f'Verso {s85} €', f'={inc}/{R["soglia"]}', PCT0),
        (12, None, barra(R['soglia']), None),
        (13, f'Mancano a {s85} €', f'=MAX(0,{R["soglia"]}-{inc})', EUR0),
        (14, f'Verso {s100} €', f'={inc}/{R["uscita"]}', PCT0),
        (15, None, barra(R['uscita']), None),
        (16, f'Mancano a {s100} €', f'=MAX(0,{R["uscita"]}-{inc})', EUR0),
    ]
    for r_, lab, formula, fmt in soglie:
        if lab is None:
            ws.merge_cells(f'I{r_}:M{r_}')
            put(ws, f'I{r_}', formula, font=Font(name=FONT, size=11, color=TEAL), align=Alignment(vertical='center'))
        else:
            ws.merge_cells(f'I{r_}:K{r_}')
            ws.merge_cells(f'L{r_}:M{r_}')
            put(ws, f'I{r_}', lab, font=f(10), align=LEFT_C)
            put(ws, f'L{r_}', formula, font=f(10, True), fmt=fmt, align=RIGHT)
    section(ws, 'I18:M18', f'Fatture {anno}')
    fat_rows = [
        (19, f'Emesse nel {anno}', f'=COUNTIFS({FR["B"]},">="&DATE({an},1,1),{FR["B"]},"<"&DATE({an}+1,1,1),{FR["F"]},"<>")', '0'),
        (20, f'Totale emesso nel {anno}', f'=SUMIFS({FR["L"]},{FR["B"]},">="&DATE({an},1,1),{FR["B"]},"<"&DATE({an}+1,1,1))', EUR),
        (21, f'Incassate nel {anno}', f'={RT}$C$17', '0'),
        (22, 'Da incassare (numero)', f'=COUNTIF({FR["N"]},"Da incassare")', '0'),
        (23, 'Da incassare (totale)', f'=SUMIFS({FR["L"]},{FR["N"]},"Da incassare")', EUR),
        (24, f'Bollo sulle fatture emesse nel {anno}', f'=SUMIFS({FR["J"]},{FR["B"]},">="&DATE({an},1,1),{FR["B"]},"<"&DATE({an}+1,1,1))', EUR),
    ]
    for r_, lab, formula, fmt in fat_rows:
        ws.merge_cells(f'I{r_}:K{r_}')
        ws.merge_cells(f'L{r_}:M{r_}')
        put(ws, f'I{r_}', lab, font=f(10), align=LEFT_C)
        put(ws, f'L{r_}', formula, font=f(10, True), fmt=fmt, align=RIGHT)
        for col in 'IJKLM':
            ws[f'{col}{r_}'].border = Border(bottom=thin)

    # Controlli
    section(ws, 'B30:M30', 'Controlli')
    c = R['code']
    controlli = [
        f'=IF({R["attivita"]}="","Scegli l\'attività in Impostazioni: senza coefficiente il reddito risulta zero.","")',
        f'=IF({c}=0,"Scegli la previdenza in Impostazioni: senza, i contributi risultano zero.","")',
        f'=IF(AND({c}=5,{R["aliqC"]}=0),"Hai scelto una cassa professionale: inserisci la sua aliquota in Impostazioni.","")',
        f'=IF(AND({R["riduzione"]}="Sì",{c}<>3,{c}<>4),"La riduzione del 35% vale solo per artigiani e commercianti: qui non è applicata.","")',
        (f'=IF(AND({c}>2,COUNTIF({FR["G"]},"Sì")+COUNTIF({FR["G"]},"Si")>0),"Fatture con rivalsa INPS 4%: "&'
         f'(COUNTIF({FR["G"]},"Sì")+COUNTIF({FR["G"]},"Si"))&". La rivalsa spetta solo agli iscritti alla Gestione '
         f'Separata.","")'),
        (f'=IF(COUNTIF({FR["N"]},"Compenso non valido")>0,"Fatture con il compenso scritto come testo: "&'
         f'COUNTIF({FR["N"]},"Compenso non valido")&". Scrivi solo il numero (es. 1500,50): finché non lo correggi '
         f'la riga non entra nei calcoli.","")'),
        (f'=IF(SUMPRODUCT(({FR["B"]}<>"")*(1-ISNUMBER({FR["B"]})))+SUMPRODUCT(({FR["C"]}<>"")*(1-ISNUMBER({FR["C"]})))>0,'
         f'"Date delle fatture scritte come testo: "&(SUMPRODUCT(({FR["B"]}<>"")*(1-ISNUMBER({FR["B"]})))+'
         f'SUMPRODUCT(({FR["C"]}<>"")*(1-ISNUMBER({FR["C"]}))))&". Scrivile come gg/mm/aaaa: un incasso con la data '
         f'non valida non entra nei calcoli.","")'),
        (f'=IF(SUMPRODUCT(ISNUMBER({FR["C"]})*ISNUMBER({FR["B"]})*({FR["C"]}<{FR["B"]}))>0,"Fatture con la data di '
         f'incasso prima di quella di emissione: "&SUMPRODUCT(ISNUMBER({FR["C"]})*ISNUMBER({FR["B"]})*'
         f'({FR["C"]}<{FR["B"]}))&". Controlla le date.","")'),
        f'=IF(COUNTIF({FR["N"]},"Incassata nel *")>0,"Fatture incassate in un altro anno, che non contano nel {anno}: "&COUNTIF({FR["N"]},"Incassata nel *")&".","")',
        (f'=IF(SUMPRODUCT(({SR["D"]}<>"")*(1-ISNUMBER({SR["D"]})*ISNUMBER({SR["A"]})))>0,"Spese senza una data valida o '
         f'con l\'importo scritto come testo: "&SUMPRODUCT(({SR["D"]}<>"")*(1-ISNUMBER({SR["D"]})*ISNUMBER({SR["A"]})))&'
         f'". Non entrano nel totale delle spese.","")'),
        (f'=IF(OR(AND({R["versati"]}<>"",NOT(ISNUMBER({R["versati"]}))),AND({R["acconti"]}<>"",NOT(ISNUMBER('
         f'{R["acconti"]}))),AND({R["aliqCassa"]}<>"",NOT(ISNUMBER({R["aliqCassa"]})))),"In Impostazioni un importo o '
         f'l\'aliquota della cassa è scritto come testo: scrivi solo il numero, altrimenti il foglio lo ignora.","")'),
        (f'=IF(AND(OR({c}=3,{c}=4),{inc}>0,{RT}$G$16<{R["minimale"]}),"Il reddito è sotto il minimale INPS: i contributi '
         f'fissi sono dovuti comunque"&IF({R["ridSi"]}="Sì"," (ridotti del 35%)."," per intero."),"")'),
    ]
    # I singoli controlli stanno nella colonna nascosta U; B31 li mostra uno per riga, senza righe vuote.
    n_c = len(controlli)
    for i, formula in enumerate(controlli):
        put(ws, f'U{31 + i}', formula, font=f(9))
    ws.column_dimensions['U'].hidden = True
    rng_u = f'$U$31:$U${30 + n_c}'
    unione = '&'.join(f'U{31 + i}&IF(U{31 + i}="","",CHAR(10))' for i in range(n_c))
    # Riquadro fisso di 8 righe (sopra i grafici): basta per i controlli che possono comparire insieme.
    ws.merge_cells('B31:M38')
    put(ws, 'B31', f'=IF(SUMPRODUCT((LEN({rng_u})>0)*1)=0,"Nessun problema rilevato.",{unione})',
        font=f(10, True, GREEN), align=Alignment(wrap_text=True, vertical='top', indent=1))
    ws.conditional_formatting.add('B31', FormulaRule(formula=[f'SUMPRODUCT((LEN({rng_u})>0)*1)>0'],
                                                     font=Font(color=AMBER, bold=True)))

    # Dati per i grafici (O..S)
    put(ws, 'O3', 'Dati per i grafici', font=f(9, True, MUTED))
    for col, lab in zip('OPQRS', ['Mese', 'Incassato', 'Da accantonare', 'Da inizio anno', f'Soglia {s85} €']):
        put(ws, f'{col}4', lab, font=f(9, True, MUTED), align=CENTER, border=BORDER)
    for m in range(1, 13):
        r_ = 4 + m
        put(ws, f'O{r_}', MESI[m - 1][:3], font=f(9, color=MUTED), border=BORDER)
        put(ws, f'P{r_}', f'={RT}B{r_}', font=f(9, color=MUTED), fmt=EUR0, border=BORDER)
        put(ws, f'Q{r_}', f'={RT}K{r_}', font=f(9, color=MUTED), fmt=EUR0, border=BORDER)
        put(ws, f'R{r_}', f'={RT}D{r_}', font=f(9, color=MUTED), fmt=EUR0, border=BORDER)
        put(ws, f'S{r_}', f'={R["soglia"]}', font=f(9, color=MUTED), fmt=EUR0, border=BORDER)

    ch = BarChart()
    ch.type = 'col'
    ch.grouping = 'clustered'
    ch.title = 'Incassi del mese e quota da accantonare'
    ch.y_axis.title = None
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E5E7EB'))
    data = Reference(ws, min_col=16, max_col=17, min_row=4, max_row=16)
    cats = Reference(ws, min_col=15, min_row=5, max_row=16)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for s, colr in zip(ch.series, (TEAL, 'F59E0B')):
        s.graphicalProperties.solidFill = colr
        s.graphicalProperties.line.solidFill = colr
    ch.gapWidth = 60
    ch.legend.position = 'b'
    ch.height, ch.width = 8.5, 15.5
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws.add_chart(ch, 'B41')

    lc = LineChart()
    lc.title = f'Incassato da inizio anno e soglia di {s85} €'
    lc.y_axis.numFmt = '#,##0'
    lc.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E5E7EB'))
    data = Reference(ws, min_col=18, max_col=19, min_row=4, max_row=16)
    lc.add_data(data, titles_from_data=True)
    lc.set_categories(cats)
    s1, s2 = lc.series
    s1.graphicalProperties.line.solidFill = TEAL
    s1.graphicalProperties.line.width = 32000
    s1.smooth = False
    s2.graphicalProperties.line.solidFill = RED
    s2.graphicalProperties.line.dashStyle = 'dash'
    s2.graphicalProperties.line.width = 19000
    s2.smooth = False
    lc.legend.position = 'b'
    lc.height, lc.width = 8.5, 15.5
    lc.x_axis.delete = False
    lc.y_axis.delete = False
    ws.add_chart(lc, 'H41')
    put(ws, 'B40', 'Grafici', font=f(11, True, TEAL_DARK))
    stampa(ws, area='A1:M58')
    ws.page_setup.fitToHeight = 1
    proteggi(ws)

    # ---------------------------------------------------------------- Scadenze
    ws = ws_sca
    larghezze(ws, [12, 44, 30, 16, 20, 58, 15, 9])
    titolo(ws, f'Scadenze {anno}-{anno + 1}',
           f'Versamenti legati all\'anno {anno}. Se una scadenza cade di sabato o festivo slitta al primo giorno '
           'lavorativo: le date qui sotto sono già spostate. Le proroghe vengono decise di anno in anno: controlla '
           'prima di pagare.', 'H')
    intestazioni(ws, 4, ['Data', 'Versamento', 'Chi lo paga', 'Codice F24', 'Importo stimato (€)', 'Note', 'Stato',
                         'Fatto'], height=30)
    ws.freeze_panes = 'A5'
    # Bollo per trimestre (blocco sotto la tabella)
    q0 = 5 + len(SCADENZE) + 3
    put(ws, f'A{q0 - 1}', 'Bollo per trimestre (fatture emesse)', font=f(11, True, TEAL_DARK))
    ws.merge_cells(f'A{q0 - 1}:C{q0 - 1}')
    Q = {}
    for t in range(1, 5):
        r_ = q0 + t - 1
        da, a = f'DATE({an},{3 * t - 2},1)', f'DATE({an},{3 * t + 1},1)'
        put(ws, f'A{r_}', f'{t}° trimestre', font=f(), border=BORDER)
        put(ws, f'B{r_}', f'=SUMIFS({FR["J"]},{FR["B"]},">="&{da},{FR["B"]},"<"&{a})', font=f(bold=True), fmt=EUR,
            border=BORDER)
        put(ws, f'C{r_}', f'=COUNTIFS({FR["J"]},">0",{FR["B"]},">="&{da},{FR["B"]},"<"&{a})&" fatture con bollo"',
            font=f(9, color=MUTED), border=BORDER)
        Q[t] = f'$B${r_}'
    rv = R['rinvioBollo']
    imp = R['dashImposta']
    importi = {
        'fissi': f'=IF(OR({c}=3,{c}=4),ROUND({R["fissi"]}/4,2),"Non dovuto")',
        'bollo1': f'=IF({Q[1]}>{rv},{Q[1]},0)',
        'bollo2': f'=IF({Q[1]}+{Q[2]}>{rv},{Q[2]}+IF({Q[1]}<={rv},{Q[1]},0),0)',
        'bollo3': f'={Q[3]}+IF({Q[1]}+{Q[2]}<={rv},{Q[1]}+{Q[2]},0)',
        'bollo4': f'={Q[4]}',
        'dich2025': 'Dalla dichiarazione',
        'dichInps': f'=IF(OR({c}=1,{c}=2,{c}=3,{c}=4),"Dalla dichiarazione","Non dovuto")',
        'dichInps2': f'=IF(OR({c}=1,{c}=2,{c}=3,{c}=4),"Dalla dichiarazione","Non dovuto")',
        # Oltre 100.000 € l'imposta sostitutiva dell'anno non è dovuta; oltre 85.000 € dal 2027 si passa al
        # regime ordinario, quindi niente acconti d'imposta sostitutiva per il 2027.
        'saldo2026': (f'=IF({inc}>{R["uscita"]},"Non dovuto (oltre {s100} €)",IF(ISNUMBER({R["acconti"]}),'
                      f'{imp}-{R["acconti"]},"Inserisci gli acconti in Impostazioni"))'),
        'acconto1': (f'=IF({inc}>{R["soglia"]},"Non dovuto (oltre {s85} €)",IF({imp}<={R["accMin"]},0,'
                     f'IF({imp}*{R["accRata"]}<{R["accPrima"]},0,{imp}*{R["accRata"]})))'),
        'acconto2': (f'=IF({inc}>{R["soglia"]},"Non dovuto (oltre {s85} €)",IF({imp}<={R["accMin"]},0,'
                     f'IF({imp}*{R["accRata"]}<{R["accPrima"]},{imp},{imp}-{imp}*{R["accRata"]})))'),
    }
    for i, (data_s, cosa, chi, codice, key, nota) in enumerate(SCADENZE):
        r_ = 5 + i
        put(ws, f'A{r_}', data_s, font=f(bold=True), fmt=DATA, border=BORDER, align=Alignment(vertical='center'))
        put(ws, f'B{r_}', cosa, font=f(), border=BORDER, align=LEFT_C)
        put(ws, f'C{r_}', chi, font=f(9), border=BORDER, align=LEFT_C)
        put(ws, f'D{r_}', codice, font=f(9), border=BORDER, align=LEFT_C)
        put(ws, f'E{r_}', importi[key], font=f(bold=True), fmt=EUR, border=BORDER,
            align=Alignment(horizontal='right', vertical='center', wrap_text=True))
        put(ws, f'F{r_}', nota, font=f(9, color=MUTED), border=BORDER, align=WRAP_C)
        put(ws, f'G{r_}', f'=IF(H{r_}="Sì","Fatto",IF(A{r_}<TODAY(),"Passata",IF(A{r_}-TODAY()<=30,"Entro 30 giorni","")))',
            font=f(9), border=BORDER, align=CENTER)
        cell = ws[f'H{r_}']
        input_cell(cell)
        cell.alignment = CENTER
        ws.row_dimensions[r_].height = (52 if len(nota) > 170 else 40 if len(nota) > 100 else
                                        30 if len(nota) > 45 or len(chi) > 30 or key.startswith(('saldo', 'acconto'))
                                        else 20)
    s_last = 4 + len(SCADENZE)
    dv = DataValidation(type='list', formula1='"Sì,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'H5:H{s_last}')
    ws.conditional_formatting.add(f'G5:G{s_last}', FormulaRule(formula=['G5="Entro 30 giorni"'],
                                                               font=Font(color=AMBER, bold=True), fill=fill(AMBER_BG)))
    ws.conditional_formatting.add(f'G5:G{s_last}', FormulaRule(formula=['G5="Fatto"'], font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add(f'A5:F{s_last}', FormulaRule(formula=['$H5="Sì"'], font=Font(color='9CA3AF')))
    nr = q0 + 5
    note_sca = [
        'Note',
        f'• Le date del {anno} sono verificate a ottobre {anno}; quelle del {anno + 1} sono i termini ordinari, salvo proroghe.',
        '• Imposta sostitutiva: gli acconti si calcolano con il metodo storico (imposta dell\'anno prima). Gli importi '
        'del 2027 sono stime basate sull\'imposta calcolata nella Dashboard.',
        '• Contributi fissi di artigiani e commercianti: rata = contributi fissi annui / 4 (Impostazioni). Gli importi '
        'esatti sono nel Cassetto previdenziale INPS.',
        '• Bollo: l\'importo è quello delle fatture emesse nel trimestre, anche se il bollo è a tuo carico. L\'Agenzia '
        'delle Entrate lo calcola e lo mostra nel portale Fatture e Corrispettivi.',
    ]
    for i, t in enumerate(note_sca):
        ws.merge_cells(f'A{nr + i}:H{nr + i}')
        put(ws, f'A{nr + i}', t, font=f(10, True, TEAL_DARK) if i == 0 else f(9, color=MUTED), align=LEFT_C)
        ws.row_dimensions[nr + i].height = 16 if i == 0 else 26
    stampa(ws, area=f'A1:H{nr + len(note_sca) - 1}')
    ws.page_setup.fitToHeight = 1  # una pagina sola: tabella, bollo per trimestre e note insieme
    proteggi(ws)

    # ---------------------------------------------------------------- Istruzioni
    crea_istruzioni(ws_istr, P, par_name)

    # Convalide: messaggi attivi e formule senza "=" iniziale (come le scrive Excel).
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            dv.showErrorMessage = True
            dv.showInputMessage = bool(dv.prompt)
            if dv.formula1 and dv.formula1.startswith('='):
                dv.formula1 = dv.formula1[1:]
            if not dv.error:
                dv.error, dv.errorTitle = 'Scegli una voce del menu a tendina.', 'Valore non valido'

    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    return wb


def crea_istruzioni(ws, P, par_name):
    anno = P['ANNO']
    F = P['FORFETTARIO']
    G = P['INPS']
    larghezze(ws, [3, 26, 92])
    ws.merge_cells('B1:C1')
    put(ws, 'B1', f'Gestionale Partita IVA forfettaria {anno}', font=f(18, True, TEAL),
        align=Alignment(vertical='center'))
    ws.row_dimensions[1].height = 36
    ws.merge_cells('B2:C2')
    put(ws, 'B2', f'NettoChiaro · parametri aggiornati al {it_data(dt.date.fromisoformat(P["AGGIORNATO"]))} · '
        'Excel, Google Sheets e LibreOffice', font=f(10, color=MUTED))
    s85, s100 = it_num(F['sogliaRicavi']), it_num(F['sogliaUscitaImmediata'])
    gs = G['gestioneSeparata']
    blocchi = [
        ('Come si usa', [
            ('1. Impostazioni', 'Scegli attività (coefficiente), previdenza e aliquota dell\'imposta.'),
            ('2. Fatture', 'Registra ogni fattura: numero, date, cliente, compenso. Rivalsa e bollo si calcolano da soli. '
             'Rivalsa INPS 4%: scegli Sì solo se la applichi (vuoto = No). Bollo al cliente: vuoto = Sì, scegli No '
             'se il bollo resta a tuo carico. Quando il cliente paga, scrivi la data di incasso.'),
            ('3. Spese', 'Facoltativo: registra i costi dell\'attività per vedere il netto vero.'),
            ('4. Dashboard', 'Incassato, imposta, contributi, netto e avanzamento verso le soglie, con i grafici.'),
            ('5. Riepilogo mensile', 'Mese per mese: incassato, quota da accantonare, netto e percentuale delle soglie.'),
            ('6. Scadenze', 'Date degli F24 dell\'anno con gli importi che il file può stimare. Segna "Fatto" quando paghi.'),
        ]),
        ('Legenda dei colori', [
            ('Giallo', 'Celle da compilare (le uniche modificabili).'),
            ('Verde chiaro', 'Risultati principali e valori calcolati da Impostazioni.'),
            ('Verde scuro', 'Intestazioni e totali.'),
            ('Bianco', 'Formule ed etichette: non modificarle.'),
        ]),
        ('Principio di cassa', [
            ('Contano gli incassi', 'Nel forfettario ricavi, tasse e soglie si calcolano su quello che incassi '
             'nell\'anno, non su quello che fatturi. Una fattura di dicembre pagata a gennaio conta nell\'anno '
             'nuovo. Il file usa sempre la data di incasso.'),
            ('Ricavo che conta', 'Compenso + rivalsa INPS + bollo addebitato al cliente (puoi escludere il bollo in '
             'Impostazioni).'),
        ]),
        ('Come calcola', [
            ('Reddito', 'Incassato × coefficiente di redditività del tuo gruppo ATECO.'),
            ('Gestione Separata', f'{it_pct(gs["aliquota"])} del reddito ({it_pct(gs["aliquotaPensionati"])} se sei '
             f'pensionato o iscritto ad altra previdenza), fino al massimale di {it_num(gs["massimale"])} €.'),
            ('Artigiani e commercianti', 'Contributi fissi sul minimale, dovuti anche con reddito più basso, più '
             'l\'aliquota sul reddito oltre il minimale (un punto in più oltre la soglia). La riduzione del 35% non '
             'tocca la quota maternità.'),
            ('Imposta sostitutiva', '(Reddito − contributi) × 15% o 5%. Stima "a regime": si deducono i contributi '
             'calcolati sull\'anno. In realtà si deducono quelli versati nell\'anno: se li conosci, inseriscili in '
             'Impostazioni.'),
            ('Bollo e rivalsa', f'Bollo di {it_num(P["BOLLO"]["importo"], 2)} € sulle fatture oltre '
             f'{it_num(P["BOLLO"]["soglia"], 2)} €. Rivalsa {it_pct(gs["rivalsa"])} solo per la Gestione Separata, '
             'se scegli Sì.'),
            ('Da accantonare', 'Ogni mese il file stima quanto mettere da parte di quello che hai incassato. Somma '
             'dei 12 mesi = tasse e contributi dell\'anno.'),
        ]),
        ('Soglie', [
            (f'{s85} €', f'Oltre {s85} € di ricavi resti forfettario quest\'anno, ma dal prossimo passi al regime ordinario.'),
            (f'{s100} €', f'Oltre {s100} € esci dal forfettario già nell\'anno: l\'IVA si applica dall\'operazione che fa '
             'superare la soglia e il reddito dell\'intero anno si tassa con IRPEF e costi reali. La Dashboard mostra '
             'un avviso e i calcoli diventano solo indicativi.'),
        ]),
        ('Avvertenze', [
            ('È una stima', 'Il file aiuta a pianificare: non sostituisce il commercialista né la dichiarazione dei '
             'redditi. Controlla sempre gli importi prima di pagare.'),
            ('Cosa non gestisce', 'Minimi e contributo integrativo delle casse professionali, altre riduzioni INPS '
             '(es. pensionati oltre 65 anni), IVA, crediti d\'imposta, redditi '
             'diversi dal forfettario.'),
            ('Fatture pagate a rate', 'Registra ogni pagamento su una riga con la sua data di incasso. Sulle righe '
             'dei pagamenti successivi al primo lascia vuota la data di emissione e scegli "No" in "Bollo al '
             'cliente": così il bollo e la fattura si contano una volta sola.'),
            ('Requisiti', 'Il file non verifica i requisiti per entrare o restare nel regime (es. redditi da lavoro '
             'dipendente, partecipazioni).'),
        ]),
        ('Uso pratico', [
            ('Protezione', 'I fogli sono protetti senza password, per non cancellare le formule per sbaglio. Per '
             'sbloccarli: Excel, Revisione > Rimuovi protezione foglio; LibreOffice, Strumenti > Proteggi foglio.'),
            ('Excel', 'Excel apre i file scaricati da internet in "Visualizzazione protetta", dove i calcoli non si '
             'aggiornano e le celle possono sembrare vuote: fai clic su "Abilita modifica".'),
            ('Google Sheets', 'Carica il file su Google Drive e aprilo con Fogli Google (File > Salva come Fogli Google).'),
            ('Date e importi', 'Scrivi le date come gg/mm/aaaa (es. 15/03/2026) e gli importi come numeri, senza '
             'testo. Se incolli dati da un altro file, il riquadro Controlli della Dashboard segnala le celle da '
             'correggere.'),
            ('Righe', f'{N_FATTURE} righe di fatture e {N_SPESE} di spese. Per l\'anno nuovo usa una copia vuota del file.'),
            ('Aggiornamento', f'Tutti i valori fiscali sono nel foglio {par_name}, con la fonte. Ogni anno esce la '
             'versione aggiornata.'),
        ]),
    ]
    r_ = 4
    for titolo_b, righe in blocchi:
        ws.merge_cells(f'B{r_}:C{r_}')
        put(ws, f'B{r_}', titolo_b, font=f(12, True, TEAL_DARK), bg=TEAL_LIGHT, align=LEFT_C)
        ws.row_dimensions[r_].height = 22
        r_ += 1
        for lab, testo in righe:
            cb = put(ws, f'B{r_}', lab, font=f(10, True), align=WRAP_C, border=Border(bottom=thin))
            put(ws, f'C{r_}', testo, font=f(10), align=WRAP_C, border=Border(bottom=thin))
            if titolo_b == 'Legenda dei colori':
                cb.fill = fill({'Giallo': INPUT, 'Verde chiaro': TEAL_BG, 'Verde scuro': TEAL_DARK,
                                'Bianco': 'FFFFFF'}[lab])
                if lab == 'Verde scuro':
                    cb.font = f(10, True, 'FFFFFF')
                if lab == 'Giallo':
                    cb.border = BORDER_INPUT
            linee = max(1, -(-len(testo) // 95))
            ws.row_dimensions[r_].height = 15 * linee + 6
            r_ += 1
        r_ += 1
    ws.merge_cells(f'B{r_}:C{r_}')
    put(ws, f'B{r_}', 'Progettato da NettoChiaro con l\'aiuto di strumenti di intelligenza artificiale; formule '
        'verificate con test. nettochiaro.com', font=f(9, color=MUTED), align=WRAP_C)
    stampa(ws, landscape=False)
    proteggi(ws)


# --------------------------------------------------------------------------------------
# Dati di esempio
# --------------------------------------------------------------------------------------
D = dt.date
ESEMPIO = {
    'impostazioni': {
        'B5': 'Laura Conti, grafica freelance',
        'B6': 'professionisti',  # id del gruppo, sostituito con l'etichetta
        'B7': 'separata',  # id della gestione, sostituito con l'etichetta
        'B9': 0,  # indice dell'opzione imposta (0 = 15%)
        'B10': 'No',
        'B11': 'Sì',
        'B13': 3000,
    },
    # n., emissione, incasso, cliente, descrizione, compenso, rivalsa, bollo al cliente
    'fatture': [
        ('18/2025', D(2025, 12, 15), D(2026, 1, 12), 'Agenzia Delta Comunicazione srl', 'Restyling logo e linee guida', 1800, 'Sì', 'Sì'),
        ('1/2026', D(2026, 1, 9), D(2026, 1, 30), 'Panificio Moretti', 'Menu e volantini', 450, 'Sì', 'Sì'),
        ('2/2026', D(2026, 1, 20), D(2026, 2, 19), 'Studio Associato Bianchi', 'Grafica pagine sito web', 2200, 'Sì', 'Sì'),
        ('3/2026', D(2026, 2, 2), D(2026, 3, 3), 'Hotel Belvedere', 'Brochure e cartella stampa', 1350, 'Sì', 'Sì'),
        ('4/2026', D(2026, 2, 16), D(2026, 2, 16), 'Giulia Ferri', 'Biglietti da visita', 70, 'No', 'Sì'),
        ('5/2026', D(2026, 2, 27), D(2026, 3, 31), 'Agenzia Delta Comunicazione srl', 'Campagna social marzo', 1600, 'Sì', 'Sì'),
        ('6/2026', D(2026, 3, 10), D(2026, 4, 9), 'Cantina Colli Alti', 'Etichette vino, 3 referenze', 2400, 'Sì', 'Sì'),
        ('7/2026', D(2026, 3, 25), D(2026, 4, 24), 'Associazione culturale Il Ponte', 'Locandine rassegna', 380, 'Sì', 'Sì'),
        ('8/2026', D(2026, 4, 8), D(2026, 5, 8), 'Studio Associato Bianchi', 'Manutenzione grafica sito', 600, 'Sì', 'Sì'),
        ('9/2026', D(2026, 4, 22), D(2026, 5, 29), 'Ferramenta Galli', 'Catalogo prodotti', 2900, 'Sì', 'Sì'),
        ('10/2026', D(2026, 5, 6), D(2026, 6, 5), 'Agenzia Delta Comunicazione srl', 'Campagna social maggio', 1600, 'Sì', 'Sì'),
        ('11/2026', D(2026, 5, 20), D(2026, 6, 19), 'Hotel Belvedere', 'Grafica sito e servizio foto', 3100, 'Sì', 'Sì'),
        ('12/2026', D(2026, 6, 3), D(2026, 7, 3), 'Palestra Linea Uno', 'Restyling logo', 950, 'Sì', 'No'),
        ('13/2026', D(2026, 6, 17), D(2026, 7, 17), 'Cantina Colli Alti', 'Packaging confezioni regalo', 2700, 'Sì', 'Sì'),
        ('14/2026', D(2026, 7, 1), D(2026, 7, 31), 'Agenzia Delta Comunicazione srl', 'Campagna social luglio', 1600, 'Sì', 'Sì'),
        ('15/2026', D(2026, 7, 15), D(2026, 8, 28), 'Studio dentistico Sorriso', 'Pieghevoli informativi', 520, 'Sì', 'Sì'),
        ('16/2026', D(2026, 7, 28), D(2026, 9, 10), 'Ferramenta Galli', 'Volantini promozione autunno', 750, 'Sì', 'Sì'),
        ('17/2026', D(2026, 9, 4), D(2026, 10, 2), 'Agenzia Delta Comunicazione srl', 'Campagna social settembre', 1600, 'Sì', 'Sì'),
        ('18/2026', D(2026, 9, 15), D(2026, 10, 15), 'Hotel Belvedere', 'Menu ristorante', 680, 'Sì', 'Sì'),
        ('19/2026', D(2026, 9, 29), D(2026, 10, 29), 'Cantina Colli Alti', 'Grafica stand fiera', 3200, 'Sì', 'Sì'),
        ('20/2026', D(2026, 10, 12), D(2026, 11, 11), 'Associazione culturale Il Ponte', 'Programma stagione', 420, 'Sì', 'Sì'),
        ('21/2026', D(2026, 10, 26), D(2026, 11, 25), 'Studio Associato Bianchi', 'Presentazione aziendale', 1300, 'Sì', 'Sì'),
        ('22/2026', D(2026, 11, 6), D(2026, 12, 4), 'Agenzia Delta Comunicazione srl', 'Campagna social novembre', 1600, 'Sì', 'Sì'),
        ('23/2026', D(2026, 11, 18), D(2026, 12, 18), 'Panificio Moretti', 'Grafica per Natale', 520, 'Sì', 'Sì'),
        ('24/2026', D(2026, 12, 2), D(2026, 12, 22), 'Ferramenta Galli', 'Calendario 2027', 1450, 'Sì', 'Sì'),
        ('25/2026', D(2026, 12, 15), None, 'Hotel Belvedere', 'Auguri e campagna gennaio', 1200, 'Sì', 'Sì'),
        ('26/2026', D(2026, 12, 21), None, 'Cantina Colli Alti', 'Etichette nuova linea', 1900, 'Sì', 'Sì'),
    ],
    # data, descrizione, categoria, importo
    'spese': [
        (D(2026, 1, 5), 'Coworking, 1° trimestre', 'Affitto e coworking', 420),
        (D(2026, 1, 15), 'Software di grafica, abbonamento annuale', 'Software e abbonamenti', 720),
        (D(2026, 1, 31), 'Hosting e dominio del portfolio', 'Software e abbonamenti', 95),
        (D(2026, 2, 10), 'Commercialista, prima rata', 'Commercialista', 400),
        (D(2026, 3, 5), 'Monitor', 'Attrezzatura', 389),
        (D(2026, 4, 7), 'Coworking, 2° trimestre', 'Affitto e coworking', 420),
        (D(2026, 4, 20), 'Corso online di tipografia', 'Formazione', 180),
        (D(2026, 6, 12), 'Stampa campioni etichette', 'Materiali', 85),
        (D(2026, 7, 6), 'Coworking, 3° trimestre', 'Affitto e coworking', 420),
        (D(2026, 9, 10), 'Commercialista, seconda rata', 'Commercialista', 400),
        (D(2026, 10, 5), 'Coworking, 4° trimestre', 'Affitto e coworking', 420),
        (D(2026, 11, 18), 'Trasferta fiera Cantina Colli Alti', 'Trasporti e trasferte', 210),
        (D(2026, 12, 1), 'Telefono e internet, anno', 'Telefono e internet', 300),
    ],
}


def compila_esempio(wb, P, E):
    ws = wb['Impostazioni']
    coeff = {c['id']: c['label'] for c in P['FORFETTARIO']['coefficienti']}
    gest = {g['id']: g['label'] for g in P['GESTIONI']}
    imposte = [f'{it_pct(P["FORFETTARIO"]["impostaOrdinaria"])} (aliquota ordinaria)',
               f'{it_pct(P["FORFETTARIO"]["impostaStartup"])} (nuova attività, primi 5 anni)']
    for coord, v in E['impostazioni'].items():
        if coord == 'B6':
            v = coeff[v]
        elif coord == 'B7':
            v = gest[v]
        elif coord == 'B9':
            v = imposte[v]
        ws[coord].value = v
    ws = wb['Fatture']
    for i, (n, em, inc, cli, desc, comp, riv, bol) in enumerate(E['fatture']):
        r_ = R0 + i
        for col, v in zip('ABCDEFGK', (n, em, inc, cli, desc, comp, riv, bol)):
            ws[f'{col}{r_}'].value = v
    ws = wb['Spese']
    for i, (d, desc, cat, imp) in enumerate(E['spese']):
        r_ = R0 + i
        for col, v in zip('ABCD', (d, desc, cat, imp)):
            ws[f'{col}{r_}'].value = v


def imposta_default(wb, P):
    """Valori predefiniti del file vuoto: 15%, nessuna riduzione, bollo nei ricavi."""
    ws = wb['Impostazioni']
    ws['B9'].value = f'{it_pct(P["FORFETTARIO"]["impostaOrdinaria"])} (aliquota ordinaria)'
    ws['B10'].value = 'No'
    ws['B11'].value = 'Sì'


def main():
    P = carica_parametri()
    out = HERE / 'files'
    out.mkdir(exist_ok=True)
    vuoto = crea_workbook(P)
    imposta_default(vuoto, P)
    p1 = out / f'{SLUG}.xlsx'
    vuoto.save(p1)
    wb = crea_workbook(P)
    imposta_default(wb, P)
    compila_esempio(wb, P, ESEMPIO)
    p2 = out / f'{SLUG}-esempio.xlsx'
    wb.save(p2)
    print(f'Scritti {p1.relative_to(HERE)} e {p2.relative_to(HERE)}')


if __name__ == '__main__':
    main()
